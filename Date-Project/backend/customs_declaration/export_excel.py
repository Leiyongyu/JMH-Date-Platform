"""
Excel 导出模块 — 基于报关单模板生成完整4-Sheet Excel文件

核心逻辑:
1. 加载原始报关单.xlsx 作为模板
2. 清除数据行，填充新数据
3. 动态调整行数（支持任意商品数量）
4. 更新全部公式和汇总
"""

import os
import copy
import datetime
from io import BytesIO
import openpyxl
from openpyxl.utils import get_column_letter
from backend.customs_declaration.config import (
    COMPANY_INFO,
    MARKUP_RATE,
    TEMPLATE_PATH,
    USD_EXCHANGE_RATE,
    VAT_RATE,
)


def calc_unit_price_usd(rmb_price):
    """计算美元单价 = ROUND(rmb/1.13*1.15*0.14, 2)"""
    return round(float(rmb_price) / (1 + VAT_RATE) * (1 + MARKUP_RATE) * USD_EXCHANGE_RATE, 2)


def calc_total_amount(rmb_price, quantity):
    """计算美元总价"""
    return round(calc_unit_price_usd(rmb_price) * int(quantity), 2)


class CustomsDeclarationExporter:
    """
    报关单 Excel 导出器
    基于模板自定义行列，支持动态商品数量
    """

    # 各 Sheet 的行结构定义
    SHEET_CONFIG = {
        'contract': {
            'sheet_index': 0,
            'data_start': 16,      # 数据起始行(1-based)
            'data_end': 40,        # 模板数据结束行
            'template_rows': 25,   # 模板中的数据行数
            'sum_qty_row': 41,     # SUM数量行 (会随插入而偏移)
            'sum_amt_row': 41,     # SUM金额行
            'sum_qty_col': 'F',    # 数量求和列
            'sum_amt_col': 'H',    # 金额求和列
            'total_label_row': 11, # TOTAL AMOUNT显示行
            'total_label_col': 'G',
            'footer_rows': [42, 43, 44, 45, 46, 47],  # 尾部固定行(需随插入偏移)
            'data_cols': {
                'A': 'line',       'B': 'description_cn',
                'C': 'sku',        'D': 'model',
                'E': 'unit',       'F': 'quantity',
                'G': 'unit_price', 'H': 'total_amount',
                'I': 'currency',
            },
        },
        'invoice': {
            'sheet_index': 1,
            'data_start': 11,
            'data_end': 35,
            'template_rows': 25,
            'sum_qty_row': 36,
            'sum_amt_row': 36,
            'sum_qty_col': 'F',
            'sum_amt_col': 'H',
            'footer_rows': [],
            'data_cols': {
                'A': 'line',       'B': 'description_cn',
                'C': 'sku',        'D': 'model',
                'E': 'unit',       'F': 'quantity',
                'G': 'unit_price', 'H': 'total_amount',
            },
        },
        'packing': {
            'sheet_index': 2,
            'data_start': 10,
            'data_end': 34,
            'template_rows': 25,
            'sum_qty_row': None,
            'sum_nw_row': 35,
            'sum_gw_row': 35,
            'sum_nw_col': 'G',
            'sum_gw_col': 'H',
            'footer_rows': [],
            'data_cols': {
                'B': 'description_cn', 'C': 'sku',
                'D': 'model',          'E': 'unit',
                'F': 'quantity',       'G': 'net_weight',
                'H': 'gross_weight',   'I': 'cbm',
                'J': 'dim_length',     'K': 'dim_width',
                'L': 'dim_height',
            },
        },
        'customs': {
            'sheet_index': 3,
            'data_start': 11,
            'data_end': 35,
            'template_rows': 25,
            'sum_qty_row': None,
            'footer_rows': [36, 37, 38],
            'data_cols': {
                'C': 'description_cn',  'D': 'sku',
                'E': 'model',
            },
        },
    }

    def __init__(self, template_path):
        self.template_path = template_path
        self.wb = None
        self.row_offset = 0  # 累计偏移量（插入/删除行导致）

    def load_template(self):
        """加载模板"""
        self.wb = openpyxl.load_workbook(self.template_path)
        self.row_offset = 0

    def export(self, order_items, po_number, inv_number, order_date=None):
        """
        主入口：生成报关单Excel

        参数:
            order_items: list of dict, 每个dict包含产品全部字段 + quantity
            po_number: PO编号
            inv_number: 发票编号
            order_date: 日期(datetime或str)
        返回:
            BytesIO 文件流
        """
        self.load_template()
        n = len(order_items)

        # 计算行偏移 (n - 25)
        row_diff = n - self.SHEET_CONFIG['contract']['template_rows']
        self.row_offset = row_diff

        # 按 Sheet 顺序处理
        self._fill_contract(order_items, po_number, order_date)
        self._fill_invoice(order_items, po_number, inv_number, order_date)
        self._fill_packing(order_items, inv_number, order_date)
        self._fill_customs(order_items, po_number, order_date)

        # 保存到 BytesIO
        output = BytesIO()
        self.wb.save(output)
        output.seek(0)
        return output

    def _adjust_rows(self, ws, data_start, data_end, footer_rows, n):
        """
        调整数据行数：插入或删除行以适应 n 个商品
        返回调整后的: (新的data_end, 新的footer_rows)
        """
        template_rows = data_end - data_start + 1
        diff = n - template_rows

        if diff > 0:
            # 需要插入行（在 data_end 之后, 即 footer 之前）
            insert_at = data_end + 1
            for _ in range(diff):
                ws.insert_rows(insert_at)
        elif diff < 0:
            # 需要删除多余行
            delete_start = data_start + n
            delete_count = abs(diff)
            ws.delete_rows(delete_start, delete_count)

        # 更新位置
        new_data_end = data_start + n - 1
        new_footer = [r + diff for r in footer_rows] if footer_rows else []
        return new_data_end, new_footer

    def _clear_data_rows(self, ws, data_start, data_end, cols):
        """清除指定数据区域的内容（保留格式），跳过合并单元格的非左上角"""
        for row in range(data_start, data_end + 1):
            for col_letter in cols:
                cell = ws[f'{col_letter}{row}']
                # 跳过 MergedCell（只读，无法写入）
                if isinstance(cell, openpyxl.cell.cell.MergedCell):
                    continue
                cell.value = None

    def _fill_contract(self, items, po_number, order_date):
        """填充合同 Sheet"""
        ws = self.wb.worksheets[self.SHEET_CONFIG['contract']['sheet_index']]
        cfg = self.SHEET_CONFIG['contract']
        n = len(items)

        # 调整行数
        new_data_end, new_footer = self._adjust_rows(
            ws, cfg['data_start'], cfg['data_end'], cfg['footer_rows'], n
        )

        # 清除原数据
        self._clear_data_rows(ws, cfg['data_start'], new_data_end,
                              ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I'])

        # 填充商品数据
        for i, item in enumerate(items):
            row = cfg['data_start'] + i
            usd_price = calc_unit_price_usd(item['rmb_price'])
            qty = int(item['quantity'])
            total = round(usd_price * qty, 2)

            ws.cell(row=row, column=1, value=i + 1)           # LINE
            ws.cell(row=row, column=2, value=item['description_cn'])
            ws.cell(row=row, column=3, value=item['sku'])
            ws.cell(row=row, column=4, value=item.get('model', '通用型'))
            ws.cell(row=row, column=5, value=item.get('unit', 'PIECE'))
            ws.cell(row=row, column=6, value=qty)
            # G列: 写入计算后的USD单价（值，非公式）
            ws.cell(row=row, column=7, value=usd_price)
            # H列: 写入计算后的总价
            ws.cell(row=row, column=8, value=total)
            ws.cell(row=row, column=9, value=item.get('currency', 'USD'))

        # 更新汇总行 (偏移后)
        sum_row = cfg['sum_qty_row'] + self.row_offset
        ws.cell(row=sum_row, column=6,
                value=f'=SUM(F{cfg["data_start"]}:F{new_data_end})')
        ws.cell(row=sum_row, column=8,
                value=f'=SUM(H{cfg["data_start"]}:H{new_data_end})')

        # 更新 TOTAL AMOUNT 引用
        ws.cell(row=cfg['total_label_row'], column=7,
                value=f'=H{sum_row}')

        # 更新 PO 编号
        ws.cell(row=5, column=7, value=po_number)
        if order_date:
            ws.cell(row=6, column=7, value=order_date)

    def _fill_invoice(self, items, po_number, inv_number, order_date):
        """填充 INVOICE Sheet"""
        ws = self.wb.worksheets[self.SHEET_CONFIG['invoice']['sheet_index']]
        cfg = self.SHEET_CONFIG['invoice']
        n = len(items)

        new_data_end, _ = self._adjust_rows(
            ws, cfg['data_start'], cfg['data_end'], cfg['footer_rows'], n
        )

        self._clear_data_rows(ws, cfg['data_start'], new_data_end,
                              ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H'])

        for i, item in enumerate(items):
            row = cfg['data_start'] + i
            usd_price = calc_unit_price_usd(item['rmb_price'])
            qty = int(item['quantity'])
            total = round(usd_price * qty, 2)

            ws.cell(row=row, column=1, value=i + 1)
            ws.cell(row=row, column=2, value=item['description_cn'])
            ws.cell(row=row, column=3, value=item['sku'])
            ws.cell(row=row, column=4, value=item.get('model', '通用型'))
            ws.cell(row=row, column=5, value=item.get('unit', 'PIECE'))
            ws.cell(row=row, column=6, value=qty)
            ws.cell(row=row, column=7, value=usd_price)
            ws.cell(row=row, column=8, value=total)

        # 更新汇总行
        sum_row = cfg['sum_qty_row'] + self.row_offset
        ws.cell(row=sum_row, column=6,
                value=f'=SUM(F{cfg["data_start"]}:F{new_data_end})')
        ws.cell(row=sum_row, column=8,
                value=f'=SUM(H{cfg["data_start"]}:H{new_data_end})')

        # 更新发票号
        ws.cell(row=7, column=7, value=inv_number)
        if order_date:
            ws.cell(row=9, column=7, value=order_date)

    def _fill_packing(self, items, inv_number, order_date):
        """填充 Packing List Sheet"""
        ws = self.wb.worksheets[self.SHEET_CONFIG['packing']['sheet_index']]
        cfg = self.SHEET_CONFIG['packing']
        n = len(items)

        new_data_end, _ = self._adjust_rows(
            ws, cfg['data_start'], cfg['data_end'], cfg['footer_rows'], n
        )

        all_cols = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']
        self._clear_data_rows(ws, cfg['data_start'], new_data_end, all_cols)

        total_nw = 0
        total_gw = 0
        total_cbm = 0

        for i, item in enumerate(items):
            row = cfg['data_start'] + i
            ws.cell(row=row, column=1, value=i + 1)
            ws.cell(row=row, column=2, value=item['description_cn'])
            ws.cell(row=row, column=3, value=item['sku'])
            ws.cell(row=row, column=4, value=item.get('model', '通用型'))
            ws.cell(row=row, column=5, value=item.get('unit', 'PIECE'))
            ws.cell(row=row, column=6, value=int(item['quantity']))
            ws.cell(row=row, column=7, value=item.get('net_weight'))
            ws.cell(row=row, column=8, value=item.get('gross_weight'))
            ws.cell(row=row, column=9, value=item.get('cbm'))
            ws.cell(row=row, column=10, value=item.get('dim_length'))
            ws.cell(row=row, column=11, value=item.get('dim_width'))
            ws.cell(row=row, column=12, value=item.get('dim_height'))

            if item.get('net_weight'):
                total_nw += float(item['net_weight'])
            if item.get('gross_weight'):
                total_gw += float(item['gross_weight'])
            if item.get('cbm'):
                total_cbm += float(item['cbm'])

        # 更新汇总
        sum_row = cfg['sum_nw_row'] + self.row_offset
        ws.cell(row=sum_row, column=7,
                value=f'=SUM(G{cfg["data_start"]}:G{new_data_end})')
        ws.cell(row=sum_row, column=8,
                value=f'=SUM(H{cfg["data_start"]}:H{new_data_end})')

        # 更新总箱数估算（简单按商品数估算，每商品放不同箱）
        ws.cell(row=sum_row - 1, column=1, value=f'CTN NO:1-{n}')

        # 更新发票号和日期（写入合并单元格的左上角）
        ws.cell(row=7, column=8, value=inv_number)   # H7 = INV.NO (H7:L7合并)
        if order_date:
            ws.cell(row=8, column=8, value=order_date)  # H8 = DATE (H8:L8合并)

    def _safe_write(self, ws, row, col, value):
        """安全写入单元格，跳过合并单元格的非左上角"""
        cell = ws.cell(row=row, column=col)
        if isinstance(cell, openpyxl.cell.cell.MergedCell):
            return  # 合并单元格的非左上角，只读
        cell.value = value

    def _fill_customs(self, items, po_number, order_date):
        """填充报关单 Sheet"""
        ws = self.wb.worksheets[self.SHEET_CONFIG['customs']['sheet_index']]
        cfg = self.SHEET_CONFIG['customs']
        n = len(items)

        new_data_end, new_footer = self._adjust_rows(
            ws, cfg['data_start'], cfg['data_end'], cfg['footer_rows'], n
        )

        # 清除数据区域
        all_cols = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']
        self._clear_data_rows(ws, cfg['data_start'], new_data_end, all_cols)

        for i, item in enumerate(items):
            row = cfg['data_start'] + i
            qty = int(item['quantity'])
            usd_price = calc_unit_price_usd(item['rmb_price'])
            total = round(usd_price * qty, 2)
            nw = item.get('net_weight', '')

            # G列: 单价/总价/币制
            price_str = f'{usd_price}/{total}/USD'

            # F列: 数量/净重
            if nw:
                qty_nw = f'{qty}个/{nw}千克'
            else:
                qty_nw = f'{qty}个'

            self._safe_write(ws, row, 1, i + 1)
            # B列: HS编码 + 申报要素
            hs_code = item.get('hs_code', '')
            hs_desc = item.get('hs_description', '')
            self._safe_write(ws, row, 2, f'{hs_code} {hs_desc}')
            self._safe_write(ws, row, 3, item['description_cn'])
            self._safe_write(ws, row, 4, item['sku'])
            self._safe_write(ws, row, 5, item.get('model', '通用型'))
            self._safe_write(ws, row, 6, qty_nw)
            self._safe_write(ws, row, 7, price_str)
            self._safe_write(ws, row, 8, item.get('origin_country', '中国'))
            self._safe_write(ws, row, 9, item.get('destination_country', '美国'))
            self._safe_write(ws, row, 10, '')              # J列(常与I合并)
            self._safe_write(ws, row, 11, item.get('source_location', ''))
            self._safe_write(ws, row, 12, '')              # L列(常与K合并)
            self._safe_write(ws, row, 13, '美元')           # M列

        # 更新合同协议号
        self._safe_write(ws, 6, 2, po_number)

        # 更新件数
        self._safe_write(ws, 7, 3, f'件数  {n}')
        # 毛重和净重在 row 7 处为文本，保留模板格式，后续可手动调整


def generate_export(items, po_number=None, inv_number=None, order_date=None):
    """
    供Flask调用的便捷函数

    参数:
        items: list of dict, 商品信息(含quantity)
        po_number: str, PO编号(可选, 自动生成)
        inv_number: str, 发票编号(可选, 自动生成)
        order_date: datetime, 日期(可选, 默认今天)
    返回:
        BytesIO 文件流
    """
    if order_date is None:
        order_date = datetime.date.today()

    # 自动生成编号
    if po_number is None:
        date_str = order_date.strftime('%y%m%d') if hasattr(order_date, 'strftime') else '260604'
        po_number = f'RVG{date_str}-0001'

    if inv_number is None:
        date_str = order_date.strftime('%y%m%d') if hasattr(order_date, 'strftime') else '260604'
        inv_number = f'INV{date_str}-0001'

    exporter = CustomsDeclarationExporter(TEMPLATE_PATH)
    return exporter.export(items, po_number, inv_number, order_date)
