"""
数据导入脚本 — 从报关单.xlsx 提取25条商品信息，导入MySQL
运行方式: python seed_data.py
"""

import sys
import os
sys.path.insert(0, os.path.dirname(__file__))

import openpyxl
from backend.customs_declaration import models
from backend.customs_declaration.config import (
    MARKUP_RATE,
    USD_EXCHANGE_RATE,
    VAT_RATE,
)


def extract_rmb_price_from_formula(formula_str):
    """
    从单价公式反推人民币含税价
    公式格式: =ROUND(RMB价格/1.13*1.15*0.14,2)
    反推: RMB = USD单价 / 0.14 / 1.15 * 1.13
    """
    if not formula_str or not str(formula_str).startswith('='):
        return None
    formula = str(formula_str)
    try:
        # 提取公式中的数字: =ROUND(638/1.13*1.15*0.14,2)
        inner = formula.replace('=ROUND(', '').replace(',2)', '').strip()
        # inner = "638/1.13*1.15*0.14"
        # 取第一个数字（分子）
        parts = inner.split('/')
        if parts:
            rmb = float(parts[0].strip())
            return round(rmb, 2)
    except (ValueError, IndexError):
        pass
    return None


def extract_hs_info(hs_cell_value):
    """
    从报关单B列解析 HS编码 和 申报要素
    格式: "8414803000 涡轮增压器，单级压缩，压缩比2.14..."
    """
    if not hs_cell_value:
        return None, None
    text = str(hs_cell_value).strip()
    # HS编码是前10位数字
    hs_code = ''
    for i, ch in enumerate(text):
        if ch.isdigit():
            hs_code += ch
            if len(hs_code) == 10:
                break
        elif hs_code:
            break
    # 申报要素 = 剩余部分
    hs_desc = text[len(hs_code):].strip() if len(hs_code) == 10 else text
    return hs_code if len(hs_code) == 10 else '', hs_desc


def seed():
    """主函数：从Excel导入数据到MySQL"""
    excel_path = os.path.join(os.path.dirname(__file__), '..', '报关单.xlsx')
    if not os.path.exists(excel_path):
        print(f'错误: 找不到Excel文件 {excel_path}')
        sys.exit(1)

    # 初始化数据库
    print('初始化数据库...')
    models.init_database()

    # 读取Excel - 需要公式和计算值
    wb_formula = openpyxl.load_workbook(excel_path, data_only=False)
    wb_data = openpyxl.load_workbook(excel_path, data_only=True)

    sn_f = wb_formula.sheetnames
    sn_d = wb_data.sheetnames

    ws_contract_f = wb_formula[sn_f[0]]   # 合同(公式)
    ws_contract_d = wb_data[sn_d[0]]       # 合同(值)
    ws_packing = wb_data[sn_d[2]]          # Packing List(值)
    ws_customs = wb_data[sn_d[3]]          # 报关单(值)

    print('提取商品数据...')
    products = []

    for row in range(16, 41):  # 合同 Row 16-40 = 25行商品
        # --- 合同 Sheet ---
        sku = str(ws_contract_d.cell(row=row, column=3).value or '').strip()
        if not sku:
            continue

        description = str(ws_contract_d.cell(row=row, column=2).value or '').strip()
        model = str(ws_contract_d.cell(row=row, column=4).value or '').strip()
        unit = str(ws_contract_d.cell(row=row, column=5).value or '').strip()
        usd_price = ws_contract_d.cell(row=row, column=7).value  # 计算后的美元单价

        # 从公式反推人民币价格
        formula = ws_contract_f.cell(row=row, column=7).value
        rmb_price = extract_rmb_price_from_formula(formula)

        # 如果无法从公式反推，则从美元单价反推
        if rmb_price is None and usd_price:
            try:
                rmb_price = round(float(usd_price) * (1 + VAT_RATE) / (1 + MARKUP_RATE) / USD_EXCHANGE_RATE, 2)
            except (ValueError, TypeError):
                rmb_price = 0

        # --- Packing List Sheet ---
        # 对应行: Packing List Row 10-34 (row - 6)
        pack_row = row - 6
        nw = ws_packing.cell(row=pack_row, column=7).value   # 净重
        gw = ws_packing.cell(row=pack_row, column=8).value   # 毛重
        cbm = ws_packing.cell(row=pack_row, column=9).value  # 体积
        dim_l = ws_packing.cell(row=pack_row, column=10).value  # 长
        dim_w = ws_packing.cell(row=pack_row, column=11).value  # 宽
        dim_h = ws_packing.cell(row=pack_row, column=12).value  # 高

        # --- 报关单 Sheet ---
        # 对应行: 报关单 Row 11-35 (row - 5)
        customs_row = row - 5
        hs_cell = ws_customs.cell(row=customs_row, column=2).value  # B列
        hs_code, hs_desc = extract_hs_info(hs_cell)
        origin = str(ws_customs.cell(row=customs_row, column=8).value or '').strip()
        dest = str(ws_customs.cell(row=customs_row, column=9).value or '').strip()
        source = str(ws_customs.cell(row=customs_row, column=11).value or '').strip()

        product = {
            'sku': sku,
            'description_cn': description,
            'model': model if model else '通用型',
            'unit': unit if unit else 'PIECE',
            'rmb_price': rmb_price or 0,
            'currency': 'USD',
            'net_weight': float(nw) if nw else None,
            'gross_weight': float(gw) if gw else None,
            'cbm': float(cbm) if cbm else None,
            'dim_length': float(dim_l) if dim_l else None,
            'dim_width': float(dim_w) if dim_w else None,
            'dim_height': float(dim_h) if dim_h else None,
            'hs_code': hs_code or '',
            'hs_description': hs_desc or '',
            'origin_country': origin if origin else '中国',
            'destination_country': dest if dest else '美国',
            'source_location': source.replace('\n', ' ') if source else '',
        }
        products.append(product)
        print(f'  {sku}: {description} | RMB {rmb_price} | USD {usd_price}')

    # 批量插入MySQL
    print(f'\n导入 {len(products)} 条商品到 MySQL...')
    for p in products:
        models.upsert_product(p)

    print('完成！')
    # 验证
    all_products = models.get_all_products()
    print(f'数据库中共有 {len(all_products)} 条商品记录')


if __name__ == '__main__':
    seed()
