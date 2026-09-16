"""Export persisted pivot values, not recomputed live inventory."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.services.ebay_inventory_detail_service import CHINA
from backend.services.ebay_inventory_detail_export_service import _ILLEGAL_XML
from backend.services.ebay_inventory_pivot_service import list_pivot

COLUMNS = (
    ("owner", "负责人", "text"), ("site", "站点", "text"), ("stat_date", "统计时间", "text"),
    ("sku_count", "SKU数", "qty"), ("overseas_sellable_quantity", "海外可售", "qty"),
    ("overseas_total_quantity", "海外总库存", "qty"), ("sales_qty_30d", "近30天销量", "qty"),
    ("in_stock_sales_ratio", "可售库销比", "percent"), ("total_stock_sales_ratio", "总库销比", "percent"),
    ("overseas_sellable_value", "海外可售货值（人民币）", "money"),
    ("overseas_total_value", "海外总货值（人民币）", "money"),
    ("warehouse_rent_30d_cny", "30天谷仓仓租（人民币）", "money"),
)


def export_pivot(**filters):
    data = list_pivot(**filters, paginate=False)
    if not data["items"]:
        raise ValueError("当前筛选条件下没有可导出的历史透视数据")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Ebay库存历史透视")
    sheet.freeze_panes = "D2"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.row_dimensions[1].height = 36
    header = []
    for index, (_, label, _) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 23 if index > 9 else 18
        cell = WriteOnlyCell(sheet, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24486D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.append(cell)
    sheet.append(header)
    for row_no, row in enumerate(data["items"], 2):
        is_owner_total = row.get("row_type") == "OWNER_TOTAL"
        cells = []
        for key, _, kind in COLUMNS:
            value = row.get(key)
            cell = WriteOnlyCell(sheet)
            if value is None:
                cell.value = "--"
            elif kind == "text":
                cell.value = _ILLEGAL_XML.sub("", str(value))
                cell.data_type = "s"
            else:
                number = Decimal(str(value))
                cell.value = number
                cell.number_format = ("0.00%" if kind == "percent" else "#,##0.00" if kind == "money"
                                      else "#,##0" if number == number.to_integral_value() else "#,##0.######")
            cell.alignment = Alignment(vertical="center", horizontal="left" if kind == "text" else "right")
            if row_no % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F7FB")
            if is_owner_total:
                cell.font = Font(bold=True, color="DC2626")
                cell.fill = PatternFill("solid", fgColor="FEF2F2")
            cells.append(cell)
        sheet.append(cells)
    sheet.auto_filter.ref = f"A1:L{len(data['items']) + 1}"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return f"Ebay库存历史透视-{datetime.now(CHINA):%Y%m%d%H%M%S}.xlsx", output.getvalue()
