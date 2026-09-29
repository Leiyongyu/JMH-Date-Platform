"""三个库存模块统一导出，只读已保存快照，不受页面行筛选影响。"""
from io import BytesIO
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import DataBarRule

from backend.services.ebay_inventory_detail_service import list_inventory
from backend.services.ebay_inventory_pivot_service import list_pivot
from backend.services.ebay_inventory_age_ratio_service import read_range
from backend.services.ebay_inventory_detail_export_service import append_inventory_sheet, _ILLEGAL_XML
from backend.services.ebay_inventory_pivot_export_service import append_pivot_sheet
from backend.repositories.performance_repository import named_lock

OWNER_COLUMNS = [('stat_date', '统计日期', 'text')] + [
    (key + '_value', label + '货值', 'money') for key, label in
    [('under_90', '<90'), ('days_90_120', '90-120'), ('days_120_180', '120-180'), ('over_180', '>180')]]
OWNER_COLUMNS += [('total_value', '总货值', 'money')] + [
    (key + '_ratio', label + '占比', 'percent') for key, label in
    [('under_90', '<90'), ('days_90_120', '90-120'), ('days_120_180', '120-180'), ('over_180', '>180')]]
OWNER_COLUMNS += [('name', '负责人', 'text')]
SITE_COLUMNS = [('stat_date', '统计日期', 'text'), ('name', '站点', 'text'),
                ('total_quantity', '总库存', 'quantity'), ('over_180_quantity', '>180天库存', 'quantity'),
                ('over_180_quantity_ratio', '库存占比', 'percent'), ('total_value', '总货值', 'money'),
                ('over_180_value', '>180天货值', 'money'), ('over_180_ratio', '货值占比', 'percent')]


def append_age_sheet(workbook, report, *, title='海外仓库龄占比'):
    from openpyxl.utils import get_column_letter
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = 'C3'
    sheet.sheet_view.showGridLines = False
    for col in range(1, 12):
        sheet.column_dimensions[get_column_letter(col)].width = 20
    line = 1
    for title, columns, rows in [('个人维度', OWNER_COLUMNS, report['owners']),
                                  ('站点维度', SITE_COLUMNS, report['sites'])]:
        sheet.append([title])
        headers = []
        for _, label, _ in columns:
            cell = WriteOnlyCell(sheet, label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24486D')
            headers.append(cell)
        sheet.append(headers)
        for row in rows:
            cells = []
            for key, _, kind in columns:
                value = row.get(key)
                cell = WriteOnlyCell(sheet)
                if value is None:
                    cell.value = '--'
                elif kind == 'text':
                    cell.value = _ILLEGAL_XML.sub('', str(value))
                    cell.data_type = 's'
                else:
                    cell.value = Decimal(str(value))
                    cell.number_format = '0.00%' if kind == 'percent' else '#,##0.00' if kind == 'money' else '#,##0.######'
                cells.append(cell)
            sheet.append(cells)
        if rows:
            for index, (_, _, kind) in enumerate(columns, 1):
                if kind == 'percent':
                    col = get_column_letter(index)
                    sheet.conditional_formatting.add(f'{col}{line + 2}:{col}{line + 1 + len(rows)}',
                        DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='F5A0AC'))
        sheet.append([])
        line += len(rows) + 3


def export_inventory(**filters):
    # Same lock order as unified refresh. No source fetch/recalculation on export.
    with named_lock('inventory:ebay-age-ratio') as age_lock:
        if not age_lock:
            raise ValueError('库龄快照正在更新，请稍后导出')
        with named_lock('inventory:ebay-pivot') as detail_lock:
            if not detail_lock:
                raise ValueError('库存快照正在更新，请稍后导出')
            return _export_saved(**filters)


def _export_saved(**filters):
    options = {key: filters[key] for key in ('stat_date', 'start_date', 'end_date', 'sort_field', 'sort_order') if key in filters}
    detail = list_inventory(**options, paginate=False)
    start, end = filters.get('start_date'), filters.get('end_date')
    if not start:
        start = end = detail.get('metadata', {}).get('stat_date')
    if not start or not end:
        raise ValueError('请选择统计日期范围后导出三个模块')
    pivot = list_pivot(start_date=start, end_date=end, paginate=False)
    age = read_range(start, end)
    if not detail['items'] and not pivot['items'] and not age['owners'] and not age['sites']:
        raise ValueError('所选统计日期范围没有可导出的库存数据')
    if sum(map(len, [detail['items'], pivot['items'], age['owners'], age['sites']])) > 200000:
        raise ValueError('三个模块合计超过20万行，请缩小日期范围后分次导出')
    detail['items'].sort(key=lambda row: row.get('stat_date', ''))
    workbook = Workbook(write_only=True)
    date_suffix = start if start == end else f"{start.replace('-', '')}-{end.replace('-', '')}"
    append_inventory_sheet(workbook, detail['items'], title=f'Ebay库存明细-{date_suffix}')
    append_pivot_sheet(workbook, pivot['items'], title=f'Ebay库存历史透视-{date_suffix}')
    append_age_sheet(workbook, age, title=f'海外仓库龄占比-{date_suffix}')
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    label = start if start == end else f'{start}_至_{end}'
    return f'库存明细持续更新-ebay-{label}.xlsx', output.getvalue()
