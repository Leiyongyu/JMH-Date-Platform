"""eBay库存采购价导入：保留完整SKU的所有不同价格，供中间码聚合使用。"""
from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation, localcontext
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from backend.services.ebay_inventory_grade_parser import (
    ERROR_VALUES, MAX_EXPANDED_BYTES, MAX_FILE_BYTES, MAX_ROWS, WORKBOOK_ERRORS,
)

MAX_COLUMNS = 64
MAX_HEADER_ROWS = 20
MAX_ARCHIVE_ENTRIES = 2000
MAX_SKU_LENGTH = 255
MAX_MIDDLE_CODE_LENGTH = 64
MAX_WARNINGS = 200
SKU_HEADERS = {"SKU", "产品代码"}
PRICE_HEADERS = {"单价(默认采购价)", "单价（默认采购价）", "单价（含税）", "单价", "采购价"}
MIDDLE_HEADERS = {"中间码"}
PRICE_SCALE = Decimal("0.000001")
MAX_PRICE = Decimal("999999999999999999.999999")
NUMBER_PATTERN = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?\Z")


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _middle_code(sku: str) -> str | None:
    # 与库存页面第二段规则一致，不把非纯数字段强制转换成有效分组。
    parts = sku.split("-")
    middle = parts[1].strip() if len(parts) > 1 else ""
    if (middle and len(middle) <= MAX_MIDDLE_CODE_LENGTH
            and all("0" <= char <= "9" for char in middle)):
        return middle
    return None


def _unit_price(value) -> Decimal:
    text = _text(value)
    if not text:
        raise ValueError("单价不能为空；公式必须有已计算的缓存结果")
    if isinstance(value, bool) or not NUMBER_PATTERN.fullmatch(text):
        raise ValueError("单价必须是有效的非负有限数值")
    try:
        price = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("单价必须是有效的非负有限数值") from exc
    if not price.is_finite() or price < 0:
        raise ValueError("单价必须是有效的非负有限数值")
    if price > MAX_PRICE:
        raise ValueError("单价超出DECIMAL(24,6)范围，整数部分最多18位")
    # 检查尾部非零位，不能四舍五入后再去重，否则会改变上传价格集合。
    _, digits, exponent = price.as_tuple()
    if exponent < -6 and any(digits[max(0, len(digits) + exponent + 6):]):
        raise ValueError("单价最多支持6位有效小数，不能自动舍入")
    with localcontext() as context:
        context.prec = 24
        return price.quantize(PRICE_SCALE)


def _header_indices(values, sheet_title: str):
    names = [_text(value).upper() for value in values]
    sku = [index for index, name in enumerate(names) if name in SKU_HEADERS]
    price = [index for index, name in enumerate(names) if name in PRICE_HEADERS]
    middle = [index for index, name in enumerate(names) if name in MIDDLE_HEADERS]
    if not sku or not price:
        return None
    if len(sku) != 1 or len(price) != 1 or len(middle) > 1:
        raise ValueError(f"{sheet_title} 表头有重复的SKU、单价或中间码列，整份未导入")
    return sku[0], price[0], middle[0] if middle else None


def parse_prices(content: bytes, filename: str) -> dict:
    """单价坏行整份拒绝；无数字中间码的SKU保留并警告，不参与中间码匹配。"""
    if Path(filename).suffix.lower() != ".xlsx":
        raise ValueError("价格文件只支持 .xlsx，表头必须包含SKU/产品代码、单价")
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("价格文件不能为空，且不能超过10MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if (len(entries) > MAX_ARCHIVE_ENTRIES
                    or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES):
                raise ValueError("价格文件解压后过大，请拆分文件后上传")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except WORKBOOK_ERRORS as exc:
        raise ValueError("价格文件不是有效的Excel工作簿") from exc

    records, warnings = {}, []
    source_rows = duplicate_rows = unmatched_middle_rows = matched_sheets = 0
    try:
        for sheet in workbook:
            # 不信任工作表dimension；按真实单元格宽度检查，避免漏掉第65列之后的数据。
            sheet.reset_dimensions()
            header = None
            has_content = False
            for row_no, cells in enumerate(sheet.iter_rows(), 1):
                if row_no > MAX_ROWS + MAX_HEADER_ROWS:
                    raise ValueError(f"单张价格表不能超过{MAX_ROWS}行")
                values = [cell.value for cell in cells]
                if not any(_text(value) for value in values):
                    continue
                has_content = True
                location = f"{sheet.title} 第{row_no}行"
                if any(_text(value) for value in values[MAX_COLUMNS:]):
                    raise ValueError(f"{location}：列数过多，价格文件最多支持{MAX_COLUMNS}列")
                if header is None:
                    if row_no > MAX_HEADER_ROWS:
                        raise ValueError(f"{sheet.title} 前{MAX_HEADER_ROWS}行未找到SKU/产品代码、单价表头")
                    header = _header_indices(values, sheet.title)
                    if header is not None:
                        matched_sheets += 1
                    continue

                source_rows += 1
                if source_rows > MAX_ROWS:
                    raise ValueError(f"一次最多导入{MAX_ROWS}行价格")
                selected = [cells[index] if index is not None and index < len(cells) else None
                            for index in header]
                selected_values = [cell.value if cell is not None else None for cell in selected]
                if any(cell is not None and cell.data_type == "e" for cell in selected) or any(
                    _text(value).upper() in ERROR_VALUES for value in selected_values
                ):
                    raise ValueError(f"{location}：SKU、单价或中间码含Excel错误值，整份未导入")
                sku = _text(selected_values[0]).upper()
                if not sku:
                    raise ValueError(f"{location}：SKU不能为空，整份未导入")
                if len(sku) > MAX_SKU_LENGTH:
                    raise ValueError(f"{location}：SKU超过{MAX_SKU_LENGTH}字符，整份未导入")
                try:
                    price = _unit_price(selected_values[1])
                except ValueError as exc:
                    raise ValueError(f"{location}：{exc}，整份未导入") from exc
                middle = _middle_code(sku)
                if header[2] is not None and _text(selected_values[2]) != (middle or ""):
                    raise ValueError(f"{location}：中间码与SKU第二段数字不一致，整份未导入")
                if middle is None:
                    unmatched_middle_rows += 1
                    if len(warnings) < MAX_WARNINGS:
                        warnings.append(f"{location}：{sku}没有有效数字中间码，已保留完整SKU价格但不参与中间码匹配")
                key = sku, price
                if key in records:
                    duplicate_rows += 1
                    continue
                records[key] = {"sku": sku, "middle_code": middle, "unit_price": price}
            if has_content and header is None:
                raise ValueError(f"{sheet.title}未找到SKU/产品代码、单价表头，整份未导入")
        if not matched_sheets or not records:
            raise ValueError("文件中没有有效价格数据，请检查SKU/产品代码及单价")
        rows = list(records.values())
        return {"rows": rows, "source_rows": source_rows, "duplicate_rows": duplicate_rows,
                "skipped_rows": 0, "warnings": warnings, "warning_count": unmatched_middle_rows,
                "unmatched_middle_rows": unmatched_middle_rows,
                "sku_count": len({row["sku"] for row in rows}),
                "middle_code_count": len({row["middle_code"] for row in rows if row["middle_code"] is not None})}
    except WORKBOOK_ERRORS as exc:
        raise ValueError("价格工作簿结构或XML内容损坏，整份未导入") from exc
    finally:
        workbook.close()
