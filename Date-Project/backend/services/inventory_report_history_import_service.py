"""One-off, non-overwriting import of historical GROUP report snapshots."""

from __future__ import annotations

import hashlib
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from backend.repositories import inventory_report_etl_repository as repo
from backend.services.inventory_report_group_rows import with_group_display_rows


MONTH_RE = re.compile(r"^(20\d{2})年(1[0-2]|[1-9])月1日")
BASE_HEADERS = {
    "total_goods_value": "总货值",
    "local_end_in_transit_qty": "成都仓在途数量",
    "local_end_in_transit_total_cost": "成都仓在途金额",
    "local_end_inventory_qty": "成都仓在库数量",
    "local_end_inventory_total_cost": "成都仓在库金额",
    "overseas_end_in_transit_qty": "海外仓/FBA在途数量",
    "overseas_end_in_transit_total_cost": "海外仓/FBA在途金额",
    "overseas_end_inventory_qty": "海外仓/FBA在库数量",
    "overseas_end_inventory_total_cost": "海外仓/FBA在库金额",
    "fba_transit_inventory_amount": "FBA在途金额+FBA在库金额",
}
OPTIONAL_HEADERS = {
    "inventory_health_rate": "库存健康度",
    "sales_target_usd": "销售目标",
    "actual_achievement_amount_usd": "实际达成",
    "target_achievement_rate": "目标达成率",
    "turnover_days_by_value": "周转天数（货值）",
    "next_month_opening_inventory_qty": "月初库存数量",
    "monthly_sales_qty": "当月销量",
    "opening_inventory_sales_ratio": "月初库销比",
    "turnover_days_by_sku": "周转天数（SKU）",
    "ctu_over_30_cost": "成都仓30天以上货值",
    "inventory_age_90_180_cost": "90-180库龄成本",
    "inventory_age_180_plus_cost": "180+库龄成本",
}


def _header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").replace("（", "(").replace("）", ")")).upper()


def _previous_month(month: str) -> str:
    year, number = map(int, month.split("-"))
    return f"{year - 1:04d}-12" if number == 1 else f"{year:04d}-{number - 1:02d}"


def _number(cell, *, percent: bool = False) -> str | None:
    if cell is None or cell.data_type == "e" or cell.value in (None, ""):
        return None
    raw = str(cell.value).strip().replace(",", "")
    percentage = raw.endswith("%")
    if percentage:
        raw = raw[:-1]
    try:
        value = Decimal(raw)
    except (InvalidOperation, ValueError):
        return None
    if not value.is_finite():
        return None
    if percent and percentage:
        value /= 100
    return str(value)


def _columns(headers: list[str]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for field, title in BASE_HEADERS.items():
        normalized = _header(title)
        if normalized in headers:
            columns[field] = headers.index(normalized)
    dynamic = {
        "inventory_health_rate": lambda h: h.startswith("库存健康度"),
        "sales_target_usd": lambda h: "销售" in h and "目标" in h and "GMV" not in h,
        "actual_achievement_amount_usd": lambda h: "实际达成" in h,
        "target_achievement_rate": lambda h: "目标达成率" in h or "目标完成率" in h,
        "turnover_days_by_value": lambda h: "周转天数" in h and "货值" in h and not h.endswith("-NEW"),
        "next_month_opening_inventory_qty": lambda h: "月初库存数量" in h,
        "monthly_sales_qty": lambda h: bool(re.search(r"\d+月销量", h)),
        "opening_inventory_sales_ratio": lambda h: "月初库销比" in h,
        "turnover_days_by_sku": lambda h: "周转天数" in h and ("SKU数量" in h or h.endswith("-数量")),
        "ctu_over_30_cost": lambda h: h.startswith("成都仓30+") or h.startswith("成都仓30天以上"),
        "inventory_age_90_180_cost": lambda h: h == "90-180库龄成本" or h == "海外仓90-180货值",
        "inventory_age_180_plus_cost": lambda h: h == "180+库龄成本" or h == "海外仓180+货值",
    }
    for field, matches in dynamic.items():
        candidates = [index for index, h in enumerate(headers) if matches(h)]
        # Prefer an exact library-cost column over the older warehouse-value column.
        if field.startswith("inventory_age_"):
            exact = "90-180库龄成本" if field == "inventory_age_90_180_cost" else "180+库龄成本"
            candidates.sort(key=lambda index: headers[index] != exact)
        if candidates:
            columns[field] = candidates[0]
    return columns


def parse_history_workbook(content: bytes, file_name: str) -> list[dict[str, Any]]:
    if not content or len(content) > 50 * 1024 * 1024:
        raise ValueError("历史文件为空或超过50MB")
    if not file_name.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("历史文件仅支持.xlsx或.xlsm")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取历史Excel文件") from exc
    parsed: list[dict[str, Any]] = []
    seen_months: set[str] = set()
    try:
        for sheet in workbook:
            match = MONTH_RE.match(sheet.title.strip())
            if not match:
                raise ValueError(f"Sheet {sheet.title} 缺少可识别的年月")
            report_month = f"{match.group(1)}-{int(match.group(2)):02d}"
            if report_month in seen_months:
                raise ValueError(f"展示月份 {report_month} 有重复Sheet，请先保留一张")
            if report_month >= date.today().strftime("%Y-%m"):
                raise ValueError(f"{report_month} 不是已结束的历史月份")
            seen_months.add(report_month)
            iterator = sheet.iter_rows()
            header_cells = next(iterator, None)
            if not header_cells:
                raise ValueError(f"Sheet {sheet.title} 没有表头")
            headers = [_header(cell.value) for cell in header_cells]
            department_index = next((i for i, h in enumerate(headers) if h.startswith("部门")), None)
            columns = _columns(headers)
            if department_index is None or "total_goods_value" not in columns:
                raise ValueError(f"Sheet {sheet.title} 缺少部门或总货值表头")
            items: list[dict[str, Any]] = []
            seen_groups: set[str] = set()
            error_count = 0
            for cells in iterator:
                group = str(cells[department_index].value or "").strip().upper()
                if not group:
                    continue
                if not re.fullmatch(r"(?:AMZ|EBAY)-[A-Z0-9-]+", group):
                    continue
                if group in seen_groups:
                    raise ValueError(f"Sheet {sheet.title} 中组别 {group} 重复")
                seen_groups.add(group)
                item: dict[str, Any] = {
                    "department_code": group,
                    "department_name": group,
                    "display_order": len(items) + 1,
                    "is_total": 0,
                    "historical_import": True,
                }
                for field, index in columns.items():
                    cell = cells[index] if index < len(cells) else None
                    if cell is not None and cell.data_type == "e":
                        error_count += 1
                    item[field] = _number(cell, percent=field.endswith("_rate") or field == "opening_inventory_sales_ratio")
                items.append(item)
            if not items:
                raise ValueError(f"Sheet {sheet.title} 没有组别数据")
            if "AMZ-US3" not in seen_groups:
                raise ValueError(f"Sheet {sheet.title} 缺少 AMZ-US3 组别")
            source_row_count = len(items)
            items = with_group_display_rows({
                "historical_import": True, "items": items,
            })["items"]
            parsed.append({
                "sheet": sheet.title,
                "report_month": report_month,
                "stat_month": _previous_month(report_month),
                "items": items,
                "row_count": len(items),
                "source_row_count": source_row_count,
                "error_cells": error_count,
                "missing_fields": [label for field, label in {**BASE_HEADERS, **OPTIONAL_HEADERS}.items() if field not in columns],
                "ignored_headers": [cell.value for i, cell in enumerate(header_cells) if cell.value is not None and i not in set(columns.values()) | {department_index}],
            })
    finally:
        workbook.close()
    return parsed


def history_preview(content: bytes, file_name: str) -> dict[str, Any]:
    parsed = parse_history_workbook(content, file_name)
    status = repo.history_month_statuses([sheet["stat_month"] for sheet in parsed])
    return {
        "file_sha256": hashlib.sha256(content).hexdigest(),
        "sheets": [{key: value for key, value in sheet.items() if key != "items"} | {"status": status.get(sheet["stat_month"], "available")}
                   for sheet in parsed],
    }


def import_history(content: bytes, file_name: str, expected_sha256: str) -> dict[str, Any]:
    digest = hashlib.sha256(content).hexdigest()
    if expected_sha256 != digest:
        raise ValueError("文件与预览时不一致，请重新预览")
    parsed = parse_history_workbook(content, file_name)
    payloads = [(sheet["stat_month"], {
        "stat_month": sheet["stat_month"],
        "source_stat_month": sheet["stat_month"],
        "report_month": sheet["report_month"],
        "historical_import": True,
        "source_sheet": sheet["sheet"],
        "source_file_sha256": digest,
        "items": sheet["items"],
    }) for sheet in parsed]
    imported, skipped = repo.insert_history_group_snapshots(payloads)
    return {"imported_months": imported, "skipped_months": skipped, "file_sha256": digest}
