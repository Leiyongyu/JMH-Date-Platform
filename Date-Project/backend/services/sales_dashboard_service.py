"""Uploaded finance history. Replace only complete platform/month slices."""
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import re
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from backend.database import db_connection

TABLES = {'ebay': 'ods_sales_dashboard_ebay_monthly', 'amz': 'ods_sales_dashboard_amz_monthly'}
COMMON = {
    '平台收入-总(RMB)': 'platform_income', '平台支出(RMB）-总': 'platform_expense',
    '毛利(RMB)': 'gross_profit', '毛利率': 'gross_margin',
}
FIELDS = {
    'ebay': {**COMMON, '净额（人民币）': 'net_amount', '销售总额': 'gross_sales',
             '退款': 'refund_amount', '赔偿': 'compensation', '成交费': 'transaction_fee',
             '罚款': 'penalty', '广告费': 'advertising_fee', '刊登/副标题': 'listing_fee',
             '其它费用': 'other_fee'},
    'amz': {**COMMON, '总收款（人民币）': 'total_receipts', '总销售额（人民币）': 'gross_sales',
            '退款金额（人民币）': 'refund_amount', '平台佣金（人民币）': 'commission',
            '发货费（人民币）': 'shipping_fee', '广告和推广费（人民币）': 'advertising_fee'},
}
METRICS = {'ebay': 'net_amount', 'amz': 'total_receipts'}
METRIC_LABELS = {'ebay': '净额（人民币）', 'amz': '总收款（人民币）'}
MAX_FILE_SIZE = 10 * 1024 * 1024


def month_key(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m')
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value or '').strip()
    match = re.fullmatch(r'(20\d{2})(?:[-/年]?)(\d{1,2})月?', text)
    if not match or not 1 <= int(match[2]) <= 12:
        raise ValueError(f'无效年月：{text}，请使用 YYYYMM 或 YYYY-MM')
    return f'{match[1]}-{int(match[2]):02}'


def normalize_header(value):
    return re.sub(r'\s+', '', str(value or '')).replace('（', '(').replace('）', ')')


def number(value, location):
    if value is None or str(value).strip() == '':
        return None
    try:
        raw = str(value).strip().replace(',', '')
        result = Decimal(raw.rstrip('%')) / (100 if raw.endswith('%') else 1)
        if not result.is_finite() or abs(result) >= Decimal('1e16'):
            raise InvalidOperation
        return result.quantize(Decimal('0.00000001'))
    except (InvalidOperation, ValueError):
        raise ValueError(f'{location} 不是有效数值：{value}') from None


def parse_workbook(content, filename):
    if not filename.lower().endswith('.xlsx'):
        raise ValueError('仅支持 .xlsx 文件，请使用 Excel 另存为此格式')
    if not content or len(content) > MAX_FILE_SIZE:
        raise ValueError('请选择非空且不超过 10MB 的文件')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError('Excel 解压后过大，请拆分文件')
        values = load_workbook(BytesIO(content), data_only=True, read_only=True)
        formulas = load_workbook(BytesIO(content), data_only=False, read_only=True)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError('无法读取 Excel，请检查文件是否损坏') from exc
    rows, warnings, seen, platforms = [], [], set(), set()
    try:
        for sheet in values:
            platform = sheet.title.strip().lower()
            if platform not in TABLES:
                continue
            if platform in platforms:
                raise ValueError(f'{platform} 存在重复 Sheet')
            platforms.add(platform)
            if sheet.max_row > 50001 or sheet.max_column > 100:
                raise ValueError(f'{sheet.title} 超过 50000 行或 100 列限制')
            header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            names = [normalize_header(v) for v in header]
            nonempty = [v for v in names if v]
            if len(set(nonempty)) != len(nonempty):
                raise ValueError(f'{sheet.title} 存在重复表头')
            required = ['日期', '平台', '站点', *FIELDS[platform]]
            missing = [h for h in required if normalize_header(h) not in names]
            if missing:
                raise ValueError(f'{sheet.title} 缺少表头：' + '、'.join(missing))
            indices = {h: names.index(normalize_header(h)) for h in required}
            count = 0
            for row_no, (cells, fcells) in enumerate(zip(
                    sheet.iter_rows(min_row=2), formulas[sheet.title].iter_rows(min_row=2)), 2):
                vals = [c.value for c in cells]
                if all(v is None or str(v).strip() == '' for v in vals):
                    if not any(c.data_type == 'f' for c in fcells):
                        continue
                location = f'{sheet.title} 第{row_no}行'
                month = month_key(vals[indices['日期']])
                if str(vals[indices['平台']] or '').strip().lower() != platform:
                    raise ValueError(f'{location} 平台与 Sheet 不一致')
                group = str(vals[indices['站点']] or '').strip().upper()
                if not group or len(group) > 64 or any(s in group for s in ('合计', '小计', '总计')):
                    raise ValueError(f'{location} 站点/组别为空、过长或为汇总行，请仅上传明细组别')
                key = platform, month, group
                if key in seen:
                    raise ValueError(f'{location} 平台＋年月＋组别重复：{key}')
                seen.add(key)
                record = {'platform': platform, 'stat_month': month, 'group_code': group,
                          'source_sheet': sheet.title, 'source_row': row_no}
                raw = []
                for idx, cell in enumerate(cells):
                    original = fcells[idx]
                    if original.data_type == 'f' and cell.value is None:
                        raise ValueError(f'{sheet.title}!{cell.coordinate} 公式无缓存值，请先用 Excel 计算并保存')
                    if cell.data_type == 'e':
                        raise ValueError(f'{sheet.title}!{cell.coordinate} 公式错误：{cell.value}')
                    raw.append({'column': get_column_letter(idx + 1), 'header': header[idx],
                                'value': cell.value, 'formula': original.value if original.data_type == 'f' else None})
                for label, field in FIELDS[platform].items():
                    record[field] = number(vals[indices[label]], f'{location} {label}')
                if record[METRICS[platform]] is None:
                    warnings.append(f'{location} {METRIC_LABELS[platform]}为空，汇总时跳过该空值，不补0')
                if platform == 'amz' and record['total_receipts'] == 0 and (
                        record['platform_income'] is None or record['platform_expense'] is None):
                    warnings.append(f'{location} {month}/{group} 总收款为0，但收入或支出为空；按原表保留0，请核对')
                record['raw_data'] = json.dumps(raw, ensure_ascii=False, default=str)
                rows.append(record)
                count += 1
            if not count:
                raise ValueError(f'{sheet.title} 没有有效数据；空表不会清空已保存月份')
    finally:
        values.close()
        formulas.close()
    if not rows:
        raise ValueError('没有找到有效的 amz 或 ebay Sheet')
    scopes = defaultdict(list)
    for row in rows:
        scopes[row['platform'], row['stat_month']].append(row)
    summary = []
    for (platform, month), records in sorted(scopes.items()):
        amounts = [r[METRICS[platform]] for r in records]
        summary.append({'platform': platform, 'month': month, 'rows': len(records),
                        'groups': [r['group_code'] for r in records],
                        'amount': str(sum(amounts, Decimal(0))) if all(v is not None for v in amounts) else None})
    return {'sha256': sha256(content).hexdigest(), 'rows': rows, 'scopes': summary,
            'warnings': warnings, 'row_count': len(rows)}


def preview(content, filename):
    result = parse_workbook(content, filename)
    return {k: v for k, v in result.items() if k != 'rows'}


def import_history(content, filename, expected_sha, operator):
    parsed = parse_workbook(content, filename)
    if parsed['sha256'] != expected_sha:
        raise ValueError('文件与预览不一致，请重新预览后确认')
    # One connection-level lock serializes imports; one transaction covers both platforms.
    with db_connection() as conn, conn.cursor() as cur:
        cur.execute("SELECT GET_LOCK('jmh:sales-dashboard-import',10) acquired")
        if cur.fetchone()['acquired'] != 1:
            raise ValueError('其他销售历史数据正在导入，请稍后重试')
        try:
            conn.begin()
            for scope in parsed['scopes']:
                cur.execute(f"DELETE FROM {TABLES[scope['platform']]} WHERE stat_month=%s", (scope['month'],))
            for platform, table in TABLES.items():
                fields = ['stat_month', 'group_code', *FIELDS[platform].values(),
                          'source_file', 'source_sheet', 'source_row', 'source_sha256', 'raw_data', 'imported_by']
                records = []
                for row in parsed['rows']:
                    if row['platform'] != platform:
                        continue
                    row.update(source_file=filename[:255], source_sha256=parsed['sha256'], imported_by=operator[:64])
                    records.append(tuple(row[f] for f in fields))
                if records:
                    cur.executemany(f"INSERT INTO {table} ({','.join(fields)}) VALUES ({','.join(['%s'] * len(fields))})", records)
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.execute("SELECT RELEASE_LOCK('jmh:sales-dashboard-import')")
    return {k: v for k, v in parsed.items() if k != 'rows'}


def options():
    result = {}
    with db_connection() as conn, conn.cursor() as cur:
        for platform, table in TABLES.items():
            cur.execute(f'SELECT DISTINCT stat_month,group_code FROM {table} ORDER BY stat_month,group_code')
            rows = cur.fetchall()
            result[platform] = {'months': sorted({r['stat_month'] for r in rows}),
                                'groups': sorted({r['group_code'] for r in rows})}
    return result


def report(platform, start_month, end_month, groups):
    if platform not in TABLES:
        raise ValueError('平台仅支持 ebay 或 amz')
    start, end = month_key(start_month), month_key(end_month)
    if start > end:
        raise ValueError('开始月份不能晚于结束月份')
    months = []
    year, month = map(int, start.split('-'))
    while f'{year:04}-{month:02}' <= end:
        months.append(f'{year:04}-{month:02}')
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    if len(months) > 120:
        raise ValueError('一次最多查询120个月')
    selected = list(dict.fromkeys(str(g).strip().upper() for g in groups if str(g).strip()))
    if len(selected) > 100 or any(len(g) > 64 for g in selected):
        raise ValueError('组别筛选过多或过长')
    params = [start, end]
    condition = ''
    if selected:
        condition = f" AND group_code IN ({','.join(['%s'] * len(selected))})"
        params.extend(selected)
    with db_connection() as conn, conn.cursor() as cur:
        cur.execute(f"SELECT stat_month,group_code,{METRICS[platform]} amount,{','.join(FIELDS[platform].values())},raw_data FROM {TABLES[platform]} "
                    f'WHERE stat_month BETWEEN %s AND %s{condition} ORDER BY stat_month,group_code', params)
        rows = cur.fetchall()
    keys = selected or sorted({r['group_code'] for r in rows})
    lookup = {(r['stat_month'], r['group_code']): r for r in rows}
    def amount(value):
        return str(value) if value is not None else None

    def available_sum(values):
        available = [Decimal(v) for v in values if v is not None]
        return sum(available, Decimal(0)) if available else None

    def money_metric(field, label):
        series = [{'name': g, 'values': [amount(lookup.get((m, g), {}).get(field)) for m in months]} for g in keys]
        totals, missing = [], []
        for m in months:
            # Sum available values; absent months/groups never invalidate other uploaded data.
            values = ([lookup.get((m, g), {}).get(field) for g in selected] if selected
                      else [r.get(field) for r in rows if r['stat_month'] == m])
            totals.append(available_sum(values))
            if not values or any(value is None for value in values):
                missing.append(m)
        summary_value = available_sum(totals)
        return {'label': label, 'unit': 'CNY', 'series': series, 'total': list(map(amount, totals)),
                'summary': {'value': amount(summary_value), 'missing_months': missing,
                            'complete': not missing, 'reason': 'missing_data' if summary_value is None else None}}

    metrics = {
        'net_sales': money_metric('amount', '净销售额'),
        'gross_profit': money_metric('gross_profit', '毛利润'),
        'platform_income': money_metric('platform_income', '平台收入'),
        'refund_amount': money_metric('refund_amount', '退款金额'),
    }
    def ratio(numerator, denominator):
        if numerator is None or denominator is None or Decimal(denominator) == 0:
            return None
        return amount(Decimal(numerator) / Decimal(denominator))
    # Ratio of sums for the entire range and for each month/group; never average source percentages.
    def ratio_metric(numerator, denominator, label):
        missing = sorted(set(numerator['summary']['missing_months'] + denominator['summary']['missing_months']))
        numerator_value = numerator['summary']['value']
        denominator_value = denominator['summary']['value']
        return {
            'label': label, 'unit': 'ratio',
            'total': [ratio(n, d) for n, d in zip(numerator['total'], denominator['total'])],
            'series': [{'name': n['name'], 'values': [ratio(nv, dv) for nv, dv in zip(n['values'], d['values'])]}
                       for n, d in zip(numerator['series'], denominator['series'])],
            'summary': {'value': ratio(numerator_value, denominator_value),
                        'missing_months': missing, 'complete': not missing,
                        'reason': 'missing_data' if numerator_value is None or denominator_value is None else
                                  'zero_denominator' if denominator_value is not None and Decimal(denominator_value) == 0 else None},
        }
    metrics['gross_margin'] = ratio_metric(metrics['gross_profit'], metrics['platform_income'], '毛利率')
    metrics['refund_rate'] = ratio_metric(metrics['refund_amount'], money_metric('gross_sales', '销售总额'), '退款率')
    detail_columns = [{'key': field, 'label': label, 'unit': 'ratio' if field == 'gross_margin' else 'CNY'}
                      for label, field in FIELDS[platform].items()]
    detail_rows = []
    known_headers = {normalize_header(h) for h in ['日期', '平台', '站点', *FIELDS[platform]]}
    extra_headers = {}
    for row in rows:
        detail = {'stat_month': row['stat_month'], 'platform': platform, 'group_code': row['group_code'],
                  **{field: amount(row.get(field)) for field in FIELDS[platform].values()}}
        # Extra uploaded business columns remain visible, without import metadata or formula source.
        raw = row.get('raw_data') or []
        if isinstance(raw, str):
            raw = json.loads(raw)
        for cell in raw:
            label = str(cell.get('header') or '').strip()
            normalized = normalize_header(label)
            if not label or normalized in known_headers:
                continue
            if normalized not in extra_headers:
                key = f'extra_{len(extra_headers)}'
                extra_headers[normalized] = key
                detail_columns.append({'key': key, 'label': label, 'unit': 'text'})
            detail[extra_headers[normalized]] = cell.get('value')
        detail_rows.append(detail)
    detail_summary = {field: amount(available_sum([row.get(field) for row in rows]))
                      for field in FIELDS[platform].values() if field != 'gross_margin'}
    detail_summary['gross_margin'] = metrics['gross_margin']['summary']['value']
    return {'platform': platform, 'metric': METRIC_LABELS[platform], 'currency': 'CNY',
            'months': months, 'series': metrics['net_sales']['series'], 'total': metrics['net_sales']['total'],
            'metrics': metrics, 'details': {'columns': detail_columns, 'rows': detail_rows, 'summary': detail_summary}}
