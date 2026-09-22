"""库存明细Excel导出，使用页面同一数据服务，不接受客户端金额。"""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.services.ebay_inventory_detail_service import CHINA, list_inventory

EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COLUMNS = (
    ("site", "站点", "text"), ("sku", "SKU", "text"),
    ("sku_middle_site_code", "中间码+站点", "text"),
    ("sku_middle_code", "中间码", "text"), ("brand", "品牌", "text"),
    ("product_name", "产品名称", "text"), ("grade", "等级", "text"),
    ("overseas_in_transit_quantity", "海外在途", "qty"),
    ("overseas_sellable_quantity", "海外可售", "qty"),
    ("overseas_total_quantity", "海外总库存", "qty"),
    ("chengdu_in_transit_quantity", "成都在途", "qty"),
    ("chengdu_sellable_quantity", "成都可售", "qty"),
    ("procurement_plan_quantity", "采购计划", "qty"),
    ("pending_outbound_quantity", "待出库", "qty"),
    ("cycle_total_quantity", "周期总库存", "qty"),
    ("overseas_max_age_days", "海外最高库龄", "qty"),
    ("sales_qty_30d", "近30天销量", "qty"),
    ("average_monthly_sales_3m", "近3个月均销量", "decimal2"),
    ("profit_rate", "利润率", "percent"),
    ("max_monthly_sales", "历史最大月销", "qty"),
    ("in_stock_sales_ratio", "在库库销比", "percent"),
    ("total_stock_sales_ratio", "总库销比", "percent"),
    ("unit_price_tax", "单价（含税）", "money"),
    ("overseas_sellable_value", "海外可售货值", "money"),
    ("overseas_total_value", "海外总货值", "money"),
    ("owner", "负责人", "text"),
    ("warehouse_rent_30d_cny", "30天谷仓仓租（人民币）", "money"),
    ("total_duration_months", "总时长（月）", "decimal2"),
    ("total_stock_sales_ratio_months", "总库销比（月）", "percent"),
    ("purchase_quantity", "申购量", "qty"),
    ("last_sold_at", "最后售出时间", "text"),
    ("stat_date", "统计日期", "text"),
)
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def export_inventory(**filters) -> tuple[str, bytes]:
    data = list_inventory(**filters, paginate=False)
    if not data["items"]:
        raise ValueError("当前筛选条件下没有可导出的库存数据")
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Ebay库存明细")
    sheet.freeze_panes = "C2"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.row_dimensions[1].height = 36
    sheet.sheet_format.defaultRowHeight = 22
    body_font = Font(name="Arial", size=10, color="243247")
    header_font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
    body_alignment = Alignment(vertical="center", horizontal="right")
    text_alignment = Alignment(vertical="center", horizontal="left")
    stripe = PatternFill("solid", fgColor="F4F7FB")
    header = []
    for index, (key, label, _) in enumerate(COLUMNS, 1):
        width = 42 if key == "product_name" else 26 if key == "sku" else 21 if key == "warehouse_rent_30d_cny" else 16
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = WriteOnlyCell(sheet, label)
        cell.font = header_font
        color = "24486D" if index <= 7 else "346A70" if index <= 16 else "806031" if index <= 20 else "625B83"
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.append(cell)
    sheet.append(header)
    for row_no, item in enumerate(data["items"], 2):
        cells = []
        for key, _, kind in COLUMNS:
            value = item.get(key)
            cell = WriteOnlyCell(sheet)
            if value is None and item.get("history_origin") == "EXCEL_IMPORT":
                cell.value = "--"
                cell.data_type = "s"
            if value is not None:
                if kind == "text":
                    cell.value = _ILLEGAL_XML.sub("", str(value))
                    cell.data_type = "s"  # SKU/品名以=开头时仍是文本，不执行公式。
                else:
                    number = Decimal(str(value))
                    cell.value = number
                    # Excel scales percentages for display; keep the original numeric ratio.
                    cell.number_format = (
                        "0.00%" if kind == "percent" else "#,##0.00" if kind == "money"
                        else "0.00" if kind == "decimal2" else "0.000000" if kind == "ratio"
                        else "#,##0" if number == number.to_integral_value() else "#,##0.00"
                    )
            cell.font = body_font
            cell.alignment = text_alignment if kind == "text" else body_alignment
            if row_no % 2 == 0:
                cell.fill = stripe
            cells.append(cell)
        sheet.append(cells)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(data['items']) + 1}"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    stat_date = data.get("metadata", {}).get("stat_date")
    prefix = f"Ebay库存明细-{stat_date}" if stat_date else "Ebay库存明细"
    filename = f"{prefix}-{datetime.now(CHINA):%Y%m%d%H%M%S}.xlsx"
    return filename, output.getvalue()
