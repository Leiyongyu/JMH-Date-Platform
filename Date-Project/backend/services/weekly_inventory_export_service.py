from __future__ import annotations

import os
import tempfile
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from backend.config import settings
from backend.services.weekly_inventory_sync_service import number

# Export scope only. Keep extraction and ODS snapshots complete; warehouse names
# are display labels and must not determine inclusion (names can change).
EXPORT_WAREHOUSE_IDS = frozenset({18677, 19561})  # CTUAMZ-EU / CTUAMZ-UK 中转仓

# Contiguous provenance blocks. Age buckets, purchase price and inventory value
# belong to inventoryDetails, NOT the product-management purchase-price field.
BASE_COLUMNS = [("仓库名称", "warehouse_name"), ("SKU", "sku"), ("本地产品id", "product_id"),
                ("实际库存总量", "product_total"), ("可用量", "product_valid_num"),
                ("次品量", "product_bad_num"), ("待检待上架量", "product_qc_num"), ("锁定量", "product_lock_num")]
BIN_COLUMNS = [("仓位名称", "whb_name"), ("总量(仓位)", "total"),
               ("锁定量(仓位)", "lock_num"), ("未锁定量(仓位)", "valid_num")]
SPEC_COLUMNS = [(f"采购-{label}规格-{direction}(CM)", f"cg_{key}_{axis}")
                for label, key in (("包装", "package"), ("外箱", "box"))
                for direction, axis in (("长", "length"), ("宽", "width"), ("高", "height"))]
SPEC_COLUMNS += [("采购-产品净重(G)", "cg_product_net_weight"), ("采购-产品毛重(KG)", "cg_product_gross_weight"),
                 ("采购-外箱实重(KG)", "cg_box_weight")]
PRODUCT_COLUMNS = [("产品名称", "product_name"), *SPEC_COLUMNS]
SOURCE_STYLES = {
    "inventory": {"header": "285E91", "body": "EAF2FA"},
    "bin": {"header": "286B52", "body": "E9F4EE"},
    "product": {"header": "9B5B17", "body": "FFF2E1"},
}


def export_root():
    return (Path(settings.export_output_dir) / "weekly_inventory").resolve()


def download_path(record):
    if not record or record["status"] != "SUCCESS":
        raise FileNotFoundError("周报不存在或尚未生成成功")
    path = Path(record["file_path"]).resolve()
    if not path.is_relative_to(export_root()) or path.name != record["file_name"] or path.suffix.lower() != ".xlsx":
        raise ValueError("周报文件路径不合法")
    if not path.is_file():
        raise FileNotFoundError("留底文件缺失，请联系管理员检查磁盘")
    return path


def build_rows(groups, names):
    # Select an entire representative row, not independent MAXs. Its seller's
    # buckets must follow the same representative. Third-party data stays in ODS,
    # but is deliberately excluded from this export.
    inventory = {}
    def rank(row):
        qty = number(row.get("product_total"))
        return (qty is not None, qty if qty is not None else Decimal(0), str(row.get("seller_id") or "0"))
    for row in groups["inventory"]:
        key = (row["wid"], row["product_id"])
        previous = inventory.get(key)
        if previous and previous.get("sku") != row.get("sku"):
            raise ValueError("同仓同产品ID对应多个SKU，拒绝不确定的合并")
        if previous is None or rank(row) > rank(previous):
            inventory[key] = row
    products = {row["product_id"]: row for row in groups["products"]}
    bins = defaultdict(list)
    for row in groups["bins"]:
        bins[(row["wid"], row["product_id"])].append(row)
    ages, label_order = {}, {}
    for row in groups["age"]:
        key = (row["wid"], row["product_id"], str(row["seller_id"]))
        ages.setdefault(key, {})[row["bucket_name"]] = row["qty"]
        label_order[row["bucket_name"]] = min(label_order.get(row["bucket_name"], row["bucket_index"]), row["bucket_index"])
    labels = sorted(label_order, key=lambda name: (label_order[name], name))
    headers = [label for label, _ in BASE_COLUMNS] + labels + ["采购单价", "库存金额"]
    headers += [label for label, _ in BIN_COLUMNS] + [label for label, _ in PRODUCT_COLUMNS]
    output = []
    for key, source in sorted(inventory.items(), key=lambda pair: (str(names.get(pair[0][0]) or ""), str(pair[1].get("sku") or ""), pair[0])):
        product = products.get(source["product_id"], {})
        bin_rows = bins[key]
        # Export named bins only; the original ODS snapshot remains complete.
        # Unnamed quantities must not be added to a different, named bin.
        bins_by_name = defaultdict(list)
        for bin_row in bin_rows:
            bin_name = str(bin_row.get("whb_name") or "").strip()
            if bin_name:
                bins_by_name[bin_name].append(bin_row)
        if not bins_by_name:
            continue
        row = {**source, "warehouse_name": names.get(source["wid"]) or next((b.get("wh_name") for b in bin_rows if b.get("wh_name")), None),
               "product_name": product.get("product_name") or next((b.get("product_name") for b in bin_rows if b.get("product_name")), None)}
        if not row["warehouse_name"]:
            raise ValueError(f"仓库 {source['wid']} 缺少名称，请先同步仓库字典")
        buckets = ages.get((*key, str(source["seller_id"])), {})
        inventory_values = [row.get(field) for _, field in BASE_COLUMNS]
        inventory_values += [buckets.get(label) for label in labels]
        price, qty = number(source.get("purchase_price")), number(source.get("product_total"))
        inventory_values += [price, price * qty if price is not None and qty is not None else None]
        product_values = []
        for _, field in PRODUCT_COLUMNS:
            value = row["product_name"] if field == "product_name" else product.get(field)
            if field == "cg_product_gross_weight":
                # Source/ODS remains grams; only the exported gross weight is kg.
                value = number(value)
                value = value / Decimal("1000") if value is not None else None
            product_values.append(value)
        for bin_name in sorted(bins_by_name):
            bin_values = [bin_name]
            for _, field in BIN_COLUMNS[1:]:
                quantities = [number(b.get(field)) for b in bins_by_name[bin_name]]
                # Do not infer total from locked/valid or turn missing into zero.
                bin_values.append(sum(quantities, Decimal(0))
                                  if quantities and all(q is not None for q in quantities) else None)
            # Inventory/product values still repeat across bins; only the three
            # bin metrics may be summed across this SKU's expanded rows.
            output.append(inventory_values + bin_values + product_values)
    return headers, output


def write_export(groups, names, path: Path):
    # Both scheduled generation and manual snapshot generation use this entry.
    # Filter before aggregation, including dynamic age headers, without mutating
    # the original snapshot or mixing another warehouse's bin quantities.
    selected = {
        group: [row for row in groups[group] if row["wid"] in EXPORT_WAREHOUSE_IDS]
        for group in ("inventory", "bins", "age")
    }
    product_ids = {row["product_id"] for row in selected["inventory"]}
    selected["products"] = [row for row in groups["products"] if row["product_id"] in product_ids]
    headers, rows = build_rows(selected, names)
    if not rows:
        raise ValueError("CTUAMZ-EU中转仓(18677)、CTUAMZ-UK中转仓(19561)没有可导出的周报数据（仅导出仓位名称非空的明细）")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("仓位库存明细")
    sheet.freeze_panes = "D2"  # Keep warehouse, SKU and product ID visible.
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 48
    sheet.sheet_format.defaultRowHeight = 24
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    inventory_columns = len(headers) - len(BIN_COLUMNS) - len(PRODUCT_COLUMNS)
    sources = (["inventory"] * inventory_columns + ["bin"] * len(BIN_COLUMNS)
               + ["product"] * len(PRODUCT_COLUMNS))
    fills = {group: {kind: PatternFill(fill_type="solid", fgColor=color)
                     for kind, color in colors.items()}
             for group, colors in SOURCE_STYLES.items()}
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="253344")
    for column, label in enumerate(headers, 1):
        width = {"仓库名称": 26, "SKU": 27, "本地产品id": 16, "产品名称": 42, "仓位名称": 27}.get(label, 19)
        sheet.column_dimensions[get_column_letter(column)].width = width

    def safe(values, *, header=False):
        cells = []
        for value, source in zip(values, sources, strict=True):
            cell = WriteOnlyCell(sheet, value=value)
            cell.fill = fills[source]["header" if header else "body"]
            cell.font = header_font if header else body_font
            cell.alignment = Alignment(horizontal="center" if header else "left" if isinstance(value, str) else "right",
                                       vertical="center", wrap_text=header)
            if isinstance(value, str):
                cell.data_type = "s"  # Never execute an imported SKU/name as a formula.
            if isinstance(value, Decimal):
                # Optional decimals alone leave a trailing dot for integers in Excel.
                # Keep numeric values intact (including fractional dimensions/costs).
                cell.number_format = "0" if value == value.to_integral_value() else "0.0#####"
            cells.append(cell)
        return cells
    sheet.append(safe(headers, header=True))
    for row in rows:
        sheet.append(safe(row))
    fd, temporary = tempfile.mkstemp(prefix=".weekly-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        workbook.save(temporary)
        with open(temporary, "r+b") as stream:
            os.fsync(stream.fileno())
        # Exclusive atomic publication, on the same volume; never replace archives.
        os.link(temporary, path)
    finally:
        workbook.close()
        Path(temporary).unlink(missing_ok=True)
    return {"row_count": len(rows), "column_count": len(headers), "file_size": path.stat().st_size}
