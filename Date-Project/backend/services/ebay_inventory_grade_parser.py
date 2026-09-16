"""eBay库存明细等级导入：保留原等级文本，按站点与完整SKU更新。"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

try:
    from lxml.etree import XMLSyntaxError
except ImportError:
    XMLSyntaxError = ParseError

WORKBOOK_ERRORS = (BadZipFile, InvalidFileException, OSError, KeyError, ParseError, XMLSyntaxError)

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 80 * 1024 * 1024
MAX_ROWS = 50000
SITE_NAMES = {
    "DE": "德国", "GERMANY": "德国", "德国": "德国",
    "UK": "英国", "GB": "英国", "UNITED KINGDOM": "英国", "英国": "英国",
    "US": "美国", "USA": "美国", "UNITED STATES": "美国", "美国": "美国",
    "FR": "法国", "FRANCE": "法国", "法国": "法国",
}
ERROR_VALUES = {"#N/A", "#NAME?", "#REF!", "#VALUE!", "#DIV/0!", "#NUM!", "#NULL!"}


def normalize_site(value: str) -> str:
    text = str(value or "").strip()
    return SITE_NAMES.get(text.upper(), text)


def parse_grades(content: bytes, filename: str) -> dict:
    """坏行明确报告并跳过；同一有效键等级冲突则整份拒绝，避免后写覆盖。"""
    if Path(filename).suffix.lower() != ".xlsx":
        raise ValueError("等级文件只支持 .xlsx，表头必须包含 SKU、站点、等级")
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("等级文件不能为空，且不能超过10MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                raise ValueError("等级文件解压后过大，请拆分文件后上传")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except WORKBOOK_ERRORS as exc:
        raise ValueError("等级文件不是有效的Excel工作簿") from exc

    records, warnings = {}, []
    duplicate_rows = skipped_rows = source_rows = matched_sheets = 0
    try:
        for sheet in workbook:
            # 不依赖可能失真的 worksheet dimension；仍以计数上限控制读取。
            sheet.reset_dimensions()
            header = None
            has_content = False
            for row_no, cells in enumerate(sheet.iter_rows(max_col=65), 1):
                if row_no > MAX_ROWS + 20:
                    raise ValueError(f"单张等级表不能超过{MAX_ROWS}行")
                values = [cell.value for cell in cells]
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                has_content = True
                if len(cells) > 64 and cells[64].value is not None:
                    raise ValueError(f"{sheet.title}列数过多，等级文件最多支持64列")
                if header is None:
                    names = [str(value or "").strip().upper() for value in values]
                    if all(name in names for name in ("SKU", "站点", "等级")):
                        if any(names.count(name) != 1 for name in ("SKU", "站点", "等级")):
                            raise ValueError(f"{sheet.title} 表头有重复的SKU、站点或等级列")
                        header = [names.index(name) for name in ("SKU", "站点", "等级")]
                        matched_sheets += 1
                        continue
                    if row_no >= 20:
                        raise ValueError(f"{sheet.title} 前20行未找到SKU、站点、等级表头")
                    continue

                source_rows += 1
                if source_rows > MAX_ROWS:
                    raise ValueError(f"一次最多导入{MAX_ROWS}行等级")
                selected = [cells[index] if index < len(cells) else None for index in header]
                text = [str(cell.value).strip() if cell is not None and cell.value is not None else ""
                        for cell in selected]
                sku, site, grade = text
                problem = None
                if any(cell is not None and cell.data_type == "e" for cell in selected) or any(
                    value.upper() in ERROR_VALUES for value in text
                ):
                    problem = "SKU、站点或等级含Excel错误值"
                elif not all(text):
                    problem = "SKU、站点和等级均不能为空"
                elif len(sku) > 255 or len(grade) > 64:
                    problem = "SKU超过255字符或等级超过64字符"
                elif site.upper() not in SITE_NAMES:
                    problem = f"无法识别站点 {site[:30]}，请使用DE/UK/US/FR或中文站点"
                if problem:
                    skipped_rows += 1
                    if len(warnings) < 200:
                        warnings.append(f"{sheet.title} 第{row_no}行：{problem}")
                    continue

                key = (normalize_site(site), sku.upper())
                if key in records:
                    if records[key]["grade"] != grade:
                        raise ValueError(f"{sheet.title} 第{row_no}行：{key[0]} / {key[1]}存在不同等级，整份未导入")
                    duplicate_rows += 1
                    continue
                records[key] = {"site": key[0], "sku": key[1], "grade": grade}
            if has_content and header is None:
                raise ValueError(f"{sheet.title}未找到SKU、站点、等级表头，整份未导入")
        if not matched_sheets or not records:
            raise ValueError("文件中没有有效等级数据，请检查SKU、站点、等级及Excel错误值")
        return {"rows": list(records.values()), "source_rows": source_rows,
                "duplicate_rows": duplicate_rows, "skipped_rows": skipped_rows,
                "warnings": warnings, "warning_count": skipped_rows}
    except WORKBOOK_ERRORS as exc:
        raise ValueError("等级工作簿结构或XML内容损坏，整份未导入") from exc
    finally:
        workbook.close()
