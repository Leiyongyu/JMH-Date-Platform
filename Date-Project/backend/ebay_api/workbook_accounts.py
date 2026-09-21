"""Read local seller tokens without persisting or logging workbook secrets."""
from __future__ import annotations

import calendar
import os
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from backend.config import PROJECT_ROOT, load_dotenv
from .credentials import EbayApiError, EbayCredentials


def add_months(value: date, months: int) -> date:
    index = value.year * 12 + value.month - 1 + months
    year, month0 = divmod(index, 12)
    month = month0 + 1
    return date(year, month, min(value.day, calendar.monthrange(year, month)[1]))


def _date(value, row_no: int):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError:
        raise EbayApiError(f'账号文件第{row_no}行授权日期或到期日期不合法，须为YYYY-MM-DD') from None


def read_workbook_accounts(*, allow_duplicate_tokens: bool = False) -> dict:
    """Return credentials in repr-redacted dataclasses; caller must not log whole rows."""
    load_dotenv()
    configured_path = os.getenv('EBAY_SELLER_ACCOUNTS_FILE', '').strip()
    if not configured_path:
        raise EbayApiError('未配置EBAY_SELLER_ACCOUNTS_FILE')
    path = Path(configured_path)
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    if path.suffix.lower() != '.xlsx':
        raise EbayApiError('店铺凭证文件必须是xlsx格式')
    client_id = os.getenv('EBAY_CLIENT_ID', '').strip()
    secret = os.getenv('EBAY_CLIENT_SECRET', '').strip()
    if not client_id or not secret:
        raise EbayApiError('请配置EBAY_CLIENT_ID和EBAY_CLIENT_SECRET')
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        sheet_name = os.getenv('EBAY_SELLER_ACCOUNTS_SHEET', '紫鸟店铺').strip()
        if sheet_name not in workbook.sheetnames:
            raise EbayApiError('账号文件缺少配置的工作表')
        sheet = workbook[sheet_name]
        if sheet.max_row > 10000 or sheet.max_column > 100:
            raise EbayApiError('账号文件尺寸超过安全读取上限')
        rows = sheet.iter_rows()
        first = next(rows, ())
        headers = [str(c.value or '').strip() for c in first]
        while headers and not headers[-1]:
            headers.pop()
        required = ('店铺名称', '平台', '店铺账号', 'browserId', '秘钥')
        if not headers or headers[-1] != '秘钥' or any(headers.count(h) != 1 for h in required):
            raise EbayApiError('账号文件表头不符合预期：需店铺名称、平台、店铺账号、browserId，最后一列为秘钥')
        columns = {h: i for i, h in enumerate(headers) if h}
        for optional in ('授权日期', '到期日期'):
            if headers.count(optional) > 1:
                raise EbayApiError('账号文件日期列重复')
        accounts, missing, seen_ids, seen_tokens = [], [], set(), set()
        for row_no, cells in enumerate(rows, 2):
            def get(name):
                return cells[columns[name]].value if name in columns else None
            platform = str(get('平台') or '').strip()
            if not (platform.casefold() == 'ebay' or platform.casefold().startswith('ebay-')):
                continue
            for name in required + tuple(h for h in ('授权日期', '到期日期') if h in columns):
                if cells[columns[name]].data_type in ('f', 'e'):
                    raise EbayApiError(f'账号文件第{row_no}行含公式或错误单元格，不读取为凭证')
            name = str(get('店铺名称') or '').strip()
            browser_id = str(get('browserId') or '').strip()
            login = str(get('店铺账号') or '').strip()
            if not name or not browser_id or len(name) > 200 or len(browser_id) > 128:
                raise EbayApiError(f'账号文件第{row_no}行店铺标识不合法')
            if browser_id in seen_ids:
                raise EbayApiError(f'账号文件第{row_no}行browserId重复')
            seen_ids.add(browser_id)
            token = get('秘钥')
            if token is None or (isinstance(token, str) and not token.strip()):
                missing.append({'source_row': row_no, 'shop_name': name, 'browser_id': browser_id})
                continue
            if not isinstance(token, str) or any(c.isspace() for c in token.strip()):
                raise EbayApiError(f'账号文件第{row_no}行密钥格式不合法')
            token = token.strip()
            if token in seen_tokens and not allow_duplicate_tokens:
                raise EbayApiError(f'账号文件第{row_no}行密钥重复，须核对店铺授权')
            seen_tokens.add(token)
            issued = _date(get('授权日期'), row_no)
            expires = _date(get('到期日期'), row_no)
            if issued and expires and expires <= issued:
                raise EbayApiError(f'账号文件第{row_no}行到期日期必须晚于授权日期')
            if not expires and issued:
                expires = add_months(issued, 18)
            if not expires:
                expires = _date(os.getenv('EBAY_SELLER_TOKEN_EXPIRES_ON', '').strip(), row_no)
            accounts.append({'source_row': row_no, 'shop_name': name, 'source_login': login,
                             'browser_id': browser_id, '_credentials': EbayCredentials(client_id, secret, token),
                             'issued_on': issued, 'expires_on': expires,
                             'remind_on': add_months(expires, -1) if expires else None,
                             'source': 'workbook'})
        if not accounts:
            raise EbayApiError('账号文件中没有可用的eBay店铺密钥，禁止执行覆盖')
        return {'accounts': accounts, 'missing_credentials': missing, 'ebay_rows': len(accounts) + len(missing)}
    except EbayApiError:
        raise
    except Exception:
        raise EbayApiError('无法读取店铺凭证文件，请检查路径、权限或文件是否损坏') from None
    finally:
        if workbook is not None:
            workbook.close()


def credential_expiry_report(today: date | None = None) -> dict:
    # Read only metadata; no OAuth or listing requests and no secrets in response.
    from zoneinfo import ZoneInfo
    today = today or datetime.now(ZoneInfo('Asia/Shanghai')).date()
    source = read_workbook_accounts()
    rows = []
    for a in source['accounts']:
        expires, remind = a['expires_on'], a['remind_on']
        status = ('UNKNOWN_EXPIRY' if expires is None else 'EXPIRED' if today >= expires
                  else 'EXPIRING' if today >= remind else 'VALID')
        rows.append({'shop_name': a['shop_name'], 'browser_id': a['browser_id'], 'source_row': a['source_row'],
                     'status': status, 'expires_on': expires.isoformat() if expires else None,
                     'remind_on': remind.isoformat() if remind else None})
    return {'checked_on': today.isoformat(), 'accounts': rows,
            'missing_credentials': source['missing_credentials'], 'ebay_rows': source['ebay_rows']}
