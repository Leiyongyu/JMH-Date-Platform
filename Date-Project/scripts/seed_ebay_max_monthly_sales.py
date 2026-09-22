"""用业务方表格为"历史最大月销"高水位种入初值；默认只试算，不写库。

多数情况下用不到：deploy/ebay-inventory-detail/04_回填历史最大月销.sql 会扫描
全部订单历史滑动取最大，已能覆盖业务方表格里的数（实测中间码10756德国回填出
199，业务方表格为197）。只有当峰值早于订单表起始月、回填也够不到时，才需要
用本脚本从外部表格补。种入的值没有对应窗口，peak_window_end 置空。

用法：
    python scripts/seed_ebay_max_monthly_sales.py <xlsx路径> [--sheet Sheet4] [--apply]

默认读取表头中的「中间码+站点」与「历史最大月销」两列，位置可变、不依赖列号。
键形如 20192DE：数字部分是中间码，后缀 DE/UK/US 对应德国/英国/美国。
不带 --apply 时只打印将要写入的行数与样例，不连写库。
"""
from __future__ import annotations

import argparse
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from backend.database import db_connection
from backend.services.ebay_inventory_detail_service import MIDDLE_SITE_SUFFIXES

TABLE = "dws_ebay_inventory_max_monthly_sales"
KEY_HEADER = "中间码+站点"
VALUE_HEADER = "历史最大月销"
SUFFIX_TO_SITE = {suffix: site for site, suffix in MIDDLE_SITE_SUFFIXES.items()}


def parse_key(raw) -> tuple[str, str] | None:
    """20192DE -> ('德国', '20192')；后缀未知或中间码非纯数字则跳过。"""
    text = str(raw or "").strip().upper()
    for suffix, site in SUFFIX_TO_SITE.items():
        if text.endswith(suffix):
            middle = text[: -len(suffix)]
            # 保留前导零：中间码在库里是文本，00123 与 123 不是同一个产品。
            return (site, middle) if middle and middle.isdigit() else None
    return None


def read_rows(path: Path, sheet: str | None) -> tuple[list[dict], list[str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("工作表为空")
        titles = [str(cell or "").strip() for cell in header]
        for name in (KEY_HEADER, VALUE_HEADER):
            if name not in titles:
                raise ValueError(f"表头缺少「{name}」列，实际表头：{titles}")
        key_at, value_at = titles.index(KEY_HEADER), titles.index(VALUE_HEADER)
        best: dict[tuple[str, str], Decimal] = {}
        skipped: list[str] = []
        for row in rows:
            if key_at >= len(row) or value_at >= len(row):
                continue
            parsed = parse_key(row[key_at])
            if parsed is None:
                if str(row[key_at] or "").strip():
                    skipped.append(f"无法识别的键：{row[key_at]}")
                continue
            try:
                value = Decimal(str(row[value_at]))
            except (InvalidOperation, TypeError, ValueError):
                skipped.append(f"{row[key_at]} 的历史最大月销不是数值：{row[value_at]}")
                continue
            if not value.is_finite() or value <= 0:
                continue
            # 同一键在表里出现多次时取最大，与"高水位"语义一致。
            if parsed not in best or value > best[parsed]:
                best[parsed] = value
    finally:
        workbook.close()
    payload = [{"site": site, "product_key_type": "MIDDLE", "product_key": middle,
                "max_monthly_sales": value}
               for (site, middle), value in sorted(best.items())]
    return payload, skipped


def apply_rows(rows: list[dict]) -> int:
    """只升不降：低于或等于已存值的不覆盖，由 GREATEST 保证可重复执行。"""
    query = f"""
        INSERT INTO {TABLE}
            (site,product_key_type,product_key,max_monthly_sales,peak_window_end,value_source)
        VALUES (%(site)s,%(product_key_type)s,%(product_key)s,%(max_monthly_sales)s,NULL,'SEEDED')
        ON DUPLICATE KEY UPDATE
            value_source=IF(VALUES(max_monthly_sales)>max_monthly_sales,'SEEDED',value_source),
            -- 种入的值来自外部表格，不知道对应哪个30天窗口，峰值窗口置空。
            peak_window_end=IF(VALUES(max_monthly_sales)>max_monthly_sales,NULL,peak_window_end),
            max_monthly_sales=GREATEST(max_monthly_sales,VALUES(max_monthly_sales))
    """
    with db_connection() as connection:
        try:
            connection.begin()
            with connection.cursor() as cursor:
                for offset in range(0, len(rows), 500):
                    cursor.executemany(query, rows[offset:offset + 500])
            connection.commit()
        except Exception:
            connection.rollback()
            raise
    return len(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="为eBay库存明细历史最大月销种入初值")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", default=None, help="工作表名，默认第一个")
    parser.add_argument("--apply", action="store_true", help="真正写库；不加时只试算")
    args = parser.parse_args()

    rows, skipped = read_rows(args.workbook, args.sheet)
    print(f"可导入 {len(rows)} 行，跳过 {len(skipped)} 行")
    for note in skipped[:10]:
        print("  跳过:", note)
    for row in rows[:5]:
        print("  样例:", row)
    if not args.apply:
        print("\n试算模式，未写库。确认无误后加 --apply 重跑。")
        return 0
    print(f"\n已写入 {apply_rows(rows)} 行（只升不降，低于已存值的不覆盖）。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
