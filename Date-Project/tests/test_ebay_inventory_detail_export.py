"""库存明细导出只使用模拟服务数据，不访问数据库或真实Excel文件。"""
from io import BytesIO
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from backend.services import ebay_inventory_detail_export_service as export_service


def test_export_contains_all_30_columns_numeric_cells_and_safe_text(monkeypatch):
    data = {"items": [{
        "site": "德国", "sku": "=1+1", "brand": "FRD", "grade": "A",
        "sku_middle_code": "00123",
        "sku_middle_site_code": "00123DE",
        "stat_date": "2026-09-16",
        "product_name": "产品\x01名称", "overseas_in_transit_quantity": "0",
        "overseas_sellable_quantity": "41", "overseas_total_quantity": "41",
        "sales_qty_30d": "12", "average_daily_sales_30d": "0.4",
        "average_monthly_sales_3m": "16.67",
        "in_stock_sales_ratio": "1.666667", "total_stock_sales_ratio_months": "1.25",
        "warehouse_rent_30d_cny": "81.03", "unit_price_tax": None,
        "total_duration_months": None,
    }]}
    query = MagicMock(return_value=data)
    monkeypatch.setattr(export_service, "list_inventory", query)
    filters = {"site": "德国", "grade": "A", "selected_keys": [{"site": "德国", "sku": "=1+1"}]}
    filename, content = export_service.export_inventory(**filters)
    query.assert_called_once_with(**filters, paginate=False)
    assert filename.startswith("Ebay库存明细-") and filename.endswith(".xlsx")
    workbook = load_workbook(BytesIO(content), data_only=False)
    try:
        sheet = workbook["Ebay库存明细"]
        assert sheet.max_column == len(export_service.COLUMNS) == 30
        assert sheet.max_row == 2
        assert tuple(cell.value for cell in sheet[1]) == tuple(column[1] for column in export_service.COLUMNS)
        columns = {field: index for index, (field, _, _) in enumerate(export_service.COLUMNS, 1)}
        sku = sheet.cell(2, columns["sku"])
        assert sku.value == "=1+1" and sku.data_type == "s"
        middle = sheet.cell(2, columns["sku_middle_code"])
        assert columns["sku_middle_site_code"] == columns["sku"] + 1
        assert columns["sku_middle_code"] == columns["sku_middle_site_code"] + 1
        joined = sheet.cell(2, columns["sku_middle_site_code"])
        assert joined.value == "00123DE" and joined.data_type == "s"
        assert middle.value == "00123" and middle.data_type == "s"
        assert sheet.cell(2, columns["product_name"]).value == "产品名称"
        quantity = sheet.cell(2, columns["overseas_sellable_quantity"])
        assert quantity.value == 41 and quantity.data_type == "n"
        assert quantity.number_format == "#,##0"
        monthly = sheet.cell(2, columns["average_monthly_sales_3m"])
        assert monthly.value == 16.67 and monthly.data_type == "n"
        assert monthly.number_format == "0.00"
        assert sheet.cell(1, columns["average_monthly_sales_3m"]).value == "近3个月均销量"
        assert "average_daily_sales_30d" not in columns
        ratio = sheet.cell(2, columns["in_stock_sales_ratio"])
        assert ratio.value == 1.666667 and ratio.number_format == "0.00%"
        monthly_ratio = sheet.cell(2, columns["total_stock_sales_ratio_months"])
        assert monthly_ratio.value == 1.25 and monthly_ratio.number_format == "0.00%"
        money = sheet.cell(2, columns["warehouse_rent_30d_cny"])
        assert money.value == 81.03 and money.number_format == "#,##0.00"
        assert sheet.cell(2, columns["unit_price_tax"]).value is None
        assert sheet.cell(2, columns["total_duration_months"]).value is None
        assert sheet.freeze_panes == "C2"
        assert sheet.auto_filter.ref == "A1:AD2"
        assert sheet.cell(2, columns["stat_date"]).value == "2026-09-16"
        assert sheet.page_setup.paperSize == 8  # OOXML A3, including write-only exports.
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.page_setup.fitToHeight == 0
    finally:
        workbook.close()


@pytest.mark.parametrize("value", [None, "0", "160", "-40", "-39"])
def test_purchase_quantity_exports_backend_result_without_clamp_or_recalculation(monkeypatch, value):
    monkeypatch.setattr(export_service, "list_inventory", lambda **kwargs: {
        "items": [{"site": "德国", "sku": "SKU-1", "purchase_quantity": value}]
    })
    _, content = export_service.export_inventory()
    workbook = load_workbook(BytesIO(content), read_only=True)
    try:
        index = [column[0] for column in export_service.COLUMNS].index("purchase_quantity") + 1
        cell = workbook.active.cell(2, index)
        if value is None:
            assert cell.value is None
        else:
            from decimal import Decimal
            assert Decimal(str(cell.value)) == Decimal(value)
            assert cell.number_format == "#,##0"
            assert cell.data_type == "n"
    finally:
        workbook.close()


@pytest.mark.parametrize("day", ["2026-08-31", "2024-12-09", "2026-08", None])
def test_fixed_duration_and_last_sold_date_excel_display(monkeypatch, day):
    monkeypatch.setattr(export_service, "list_inventory", lambda **kwargs: {"items": [
        {"site": "德国", "sku": "SKU-1", "total_duration_months": "4.03", "last_sold_at": day}
    ]})
    _, content = export_service.export_inventory()
    workbook = load_workbook(BytesIO(content), read_only=True)
    try:
        columns = {key: index for index, (key, _, _) in enumerate(export_service.COLUMNS, 1)}
        duration = workbook.active.cell(2, columns["total_duration_months"])
        assert duration.value == 4.03 and duration.number_format == "0.00"
        # Old frozen YYYY-MM values are not given a fabricated day.
        assert workbook.active.cell(2, columns["last_sold_at"]).value == day
    finally:
        workbook.close()


def test_pending_outbound_and_derived_fields_export_service_values(monkeypatch):
    values = {"procurement_plan_quantity": "0", "pending_outbound_quantity": "12", "cycle_total_quantity": "52",
              "total_stock_sales_ratio_months": "0.52", "purchase_quantity": "351"}
    monkeypatch.setattr(export_service, "list_inventory", lambda **kwargs: {
        "items": [{"site": "德国", "sku": "SKU-1", **values}]
    })
    _, content = export_service.export_inventory()
    workbook = load_workbook(BytesIO(content), read_only=True)
    try:
        from decimal import Decimal
        columns = {key: index for index, (key, _, _) in enumerate(export_service.COLUMNS, 1)}
        for key, expected in values.items():
            cell = workbook.active.cell(2, columns[key])
            assert cell.data_type == "n"
            assert Decimal(str(cell.value)) == Decimal(expected)
    finally:
        workbook.close()


def test_no_selection_exports_all_filtered_rows_in_service_order(monkeypatch):
    query = MagicMock(return_value={"items": [
        {"site": "英国", "sku": "SKU-B"}, {"site": "英国", "sku": "SKU-A"},
    ]})
    monkeypatch.setattr(export_service, "list_inventory", query)
    _, content = export_service.export_inventory(site="英国", selected_keys=[], sort_field="sku", sort_order="descending")
    query.assert_called_once_with(site="英国", selected_keys=[], sort_field="sku", sort_order="descending", paginate=False)
    workbook = load_workbook(BytesIO(content), read_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
        assert [row[1] for row in rows[1:]] == ["SKU-B", "SKU-A"]
    finally:
        workbook.close()


@pytest.mark.parametrize("value", ["0", "0.125", "1.25", "2.5", None])
def test_stock_ratios_use_excel_percent_format_without_scaling_values(monkeypatch, value):
    fields = ("in_stock_sales_ratio", "total_stock_sales_ratio", "total_stock_sales_ratio_months")
    monkeypatch.setattr(export_service, "list_inventory", lambda **kwargs: {
        "items": [{"site": "德国", "sku": "SKU-1", **{field: value for field in fields}}]
    })
    _, content = export_service.export_inventory()
    workbook = load_workbook(BytesIO(content))
    try:
        columns = {field: index for index, (field, _, _) in enumerate(export_service.COLUMNS, 1)}
        for field in fields:
            cell = workbook.active.cell(2, columns[field])
            if value is None:
                assert cell.value is None
            else:
                assert cell.value == float(value)  # 1.25 stays 1.25, displayed as 125.00%.
                assert cell.data_type == "n"
                assert cell.number_format == "0.00%"
        assert workbook.active.cell(2, columns["total_duration_months"]).value is None
    finally:
        workbook.close()


def test_empty_export_is_explicit_error(monkeypatch):
    monkeypatch.setattr(export_service, "list_inventory", lambda **kwargs: {"items": []})
    with pytest.raises(ValueError, match="没有可导出"):
        export_service.export_inventory(selected_keys=[])


@pytest.mark.parametrize("value", ["123.45", "0", None])
def test_price_and_valuation_columns_keep_numeric_zero_null_and_money_format(monkeypatch, value):
    fields = ("unit_price_tax", "overseas_sellable_value", "overseas_total_value")
    monkeypatch.setattr(export_service, "list_inventory", lambda **kwargs: {
        "items": [{"site": "德国", "sku": "FRD-70618-0687", **{field: value for field in fields}}]
    })
    _, content = export_service.export_inventory()
    workbook = load_workbook(BytesIO(content), data_only=True)
    try:
        columns = {field: index for index, (field, _, _) in enumerate(export_service.COLUMNS, 1)}
        for field in fields:
            cell = workbook.active.cell(2, columns[field])
            if value is None:
                assert cell.value is None
            else:
                assert cell.value == float(value)
                assert cell.data_type == "n"
                assert cell.number_format == "#,##0.00"
    finally:
        workbook.close()


@pytest.mark.parametrize("value", ["0", "58", "100", None])
def test_overseas_highest_age_exports_integer_numeric_or_blank(monkeypatch, value):
    monkeypatch.setattr(export_service, "list_inventory", lambda **kwargs: {
        "items": [{"site": "英国", "sku": "FRD-70618-0687", "overseas_max_age_days": value}]
    })
    _, content = export_service.export_inventory()
    workbook = load_workbook(BytesIO(content), data_only=True)
    try:
        columns = {field: index for index, (field, _, _) in enumerate(export_service.COLUMNS, 1)}
        cell = workbook.active.cell(2, columns["overseas_max_age_days"])
        assert workbook.active.max_column == 30
        if value is None:
            assert cell.value is None
        else:
            assert cell.value == int(value)
            assert cell.data_type == "n"
            assert cell.number_format == "#,##0"
    finally:
        workbook.close()
