"""统一入口回归：仅内存工作簿与隔离存储，不操作真实数据库。"""
from contextlib import contextmanager
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from backend.services import ebay_inventory_bundle_export_service as export
from backend.services import ebay_inventory_pivot_service as pivot
from backend.services import ebay_inventory_age_ratio_service as age
from backend.repositories import ebay_inventory_pivot_repository as repo
from backend.repositories import ebay_inventory_age_ratio_repository as age_repo
from test_ebay_inventory_pivot_api_export import client


@contextmanager
def lock(_):
    yield True


def test_three_sheets_use_same_dates_ignore_row_filters_and_preserve_numeric_percent(monkeypatch):
    monkeypatch.setattr(export, 'named_lock', lock)
    detail = MagicMock(return_value={'items': [{'site': '德国', 'sku': '=not-a-formula',
        'stat_date': '2026-09-29', 'unit_price_tax': '12.34'}]})
    totals = MagicMock(return_value={'items': [{'owner': '张三', 'site': '德国', 'stat_date': '2026-09-29',
                                              'sku_count': 1, 'row_type': 'OWNER_TOTAL'}]})
    ratio = MagicMock(return_value={'owners': [{'name': '=owner', 'stat_date': '2026-09-29',
                         'total_value': '100', 'under_90_ratio': '0.25'}],
                        'sites': [{'name': '美国', 'stat_date': '2026-09-29', 'total_quantity': '12',
                                   'over_180_quantity_ratio': '0.5'}]})
    monkeypatch.setattr(export, 'list_inventory', detail)
    monkeypatch.setattr(export, 'list_pivot', totals)
    monkeypatch.setattr(export, 'read_range', ratio)
    name, content = export.export_inventory(start_date='2026-09-01', end_date='2026-09-29',
                                            owner='other', sku='NO', site='NO', page=999)
    assert name == '库存明细持续更新-ebay-2026-09-01_至_2026-09-29.xlsx'
    detail.assert_called_once_with(start_date='2026-09-01', end_date='2026-09-29', paginate=False)
    totals.assert_called_once_with(start_date='2026-09-01', end_date='2026-09-29', paginate=False)
    ratio.assert_called_once_with('2026-09-01', '2026-09-29')
    book = load_workbook(BytesIO(content))
    try:
        assert book.sheetnames == ['Ebay库存明细-20260901-20260929',
                                   'Ebay库存历史透视-20260901-20260929',
                                   '海外仓库龄占比-20260901-20260929']
        assert all(len(name) <= 31 for name in book.sheetnames)
        assert book.worksheets[0]['B2'].data_type == 's'
        sheet = book.worksheets[2]
        assert sheet['G3'].value == 0.25 and sheet['G3'].number_format == '0.00%'
        assert sheet['K3'].value == '=owner' and sheet['K3'].data_type == 's'
        assert sheet['A5'].value == '站点维度' and sheet['C7'].value == 12
        assert len(sheet.conditional_formatting) == 6
    finally:
        book.close()


def test_age_only_history_still_exports_three_sheets(monkeypatch):
    monkeypatch.setattr(export, 'named_lock', lock)
    monkeypatch.setattr(export, 'list_inventory', lambda **k: {'items': []})
    monkeypatch.setattr(export, 'list_pivot', lambda **k: {'items': []})
    monkeypatch.setattr(export, 'read_range', lambda *a: {'owners': [{'name': '张三'}], 'sites': []})
    _, data = export.export_inventory(start_date='2026-08-03', end_date='2026-08-03')
    book = load_workbook(BytesIO(data))
    assert len(book.sheetnames) == 3
    assert book.sheetnames == ['Ebay库存明细-2026-08-03',
                               'Ebay库存历史透视-2026-08-03',
                               '海外仓库龄占比-2026-08-03']
    assert book.worksheets[0].max_row == 1
    book.close()


def test_unified_refresh_calculates_before_publishing_and_uses_same_timestamp(monkeypatch):
    monkeypatch.setattr(pivot, 'named_lock', lock)
    build = MagicMock(return_value={'stat_date': '2026-09-29'})
    save = MagicMock(return_value={'stat_date': '2026-09-29'})
    monkeypatch.setattr(age, 'build_snapshot', build)
    monkeypatch.setattr(pivot, 'capture_snapshot', save)
    assert pivot.capture_all_snapshots()['stat_date'] == '2026-09-29'
    save.assert_called_once_with(trigger_type='PAGE_REFRESH', generated_at=build.call_args.args[0],
                                 age_report=build.return_value)
    build.side_effect = ValueError('no cost data')
    save.reset_mock()
    with pytest.raises(ValueError):
        pivot.capture_all_snapshots()
    save.assert_not_called()


def test_age_save_failure_rolls_back_detail_and_pivot_together(monkeypatch):
    connection = MagicMock()
    connection.__enter__.return_value = connection
    cursor = connection.cursor.return_value.__enter__.return_value
    cursor.fetchone.return_value = {'id': 12}
    monkeypatch.setattr(repo, 'db_connection', lambda: connection)
    monkeypatch.setattr(age_repo, 'save_snapshot_cursor', MagicMock(side_effect=RuntimeError('save failed')))
    header = dict(stat_date=date(2026, 9, 29), stat_month='2026-09', generated_at='2026-09-29 12:00:00',
                  inventory_batch_id='batch', trigger_type='PAGE_REFRESH', item_count=0, metadata={})
    with pytest.raises(RuntimeError):
        repo.replace_day(header, [], [], age_report={'stat_date': '2026-09-29'})
    connection.commit.assert_not_called()
    connection.rollback.assert_called_once()
    age_repo.save_snapshot_cursor.assert_called_once_with(cursor, {'stat_date': '2026-09-29'})


def test_list_date_range_forwarded(client, monkeypatch):
    from backend.api.v1 import ebay_inventory_detail as api
    read = MagicMock(return_value={'items': []})
    monkeypatch.setattr(api.service, 'list_inventory', read)
    result = client.get('/api/v1/finance/ebay-inventory-detail/list',
                        params={'start_date': '2026-08-01', 'end_date': '2026-08-31'})
    assert result.status_code == 200
    assert read.call_args.kwargs['start_date'] == '2026-08-01'
    assert read.call_args.kwargs['end_date'] == '2026-08-31'
