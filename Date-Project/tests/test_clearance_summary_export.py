from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from backend.services import clearance_export_service as service


def test_four_sheets_share_selected_month_and_page_data(monkeypatch, tmp_path):
    calls = []
    groups = [{"group_code": code, "total_inventory_cost": Decimal("12.34"),
               "ctu_over_30_cost": Decimal("50"), "ctu_over_30_qty": 2,
               "inventory_91_180_cost": 0, "previous_month_91_180_cost": None}
              for code in ("EBAY-1", "EU", "US1", "US2", "US2-MJ", "US1-ZXY")]

    def group_data(month):
        assert month == "2026-09"
        return {"items": groups}

    def dimensions(month, group, *, dimension=None):
        calls.append((month, group, dimension))
        return {"items": [{"dimension_name": '=HYPERLINK("bad")', "brands": ["FRD"],
                           "previous_brands": ["BMW"], "total_inventory_qty": 3,
                           "total_inventory_cost": Decimal("12.34"),
                           "inventory_91_180_cost": 0, "previous_month_91_180_cost": None}]}

    monkeypatch.setattr(service.repo, "list_groups", group_data)
    monkeypatch.setattr(service, "list_dimensions", dimensions)
    monkeypatch.setattr(service, "settings", SimpleNamespace(export_output_dir=str(tmp_path)))
    path, name = service.export_monthly_summary("2026-09")
    assert "2026-09" in name
    with Path(path).open("rb") as source:
        wb = load_workbook(source)
        assert wb.sheetnames == [f"{label}-2026-09" for label in ("组别", "店铺", "个人", "站点")]
        for sheet in wb:
            headers = [c.value for c in sheet[1]]
            assert not any("品牌" in header for header in headers)
            rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
            assert all(row["快照月份"] == "2026-09" for row in rows)
            assert all(row["海外仓/FBA合计库龄成本"] == 12.34 for row in rows)
            assert all(row["91–180天库龄成本"] == 0 for row in rows)
            assert all(row["91–180天上月库龄成本"] is None for row in rows)
            assert not any(cell.data_type == "f" for row in sheet for cell in row)
        group_sheet = wb.worksheets[0]
        assert sum(row[5] or 0 for row in group_sheet.iter_rows(min_row=2, values_only=True)) == 250
        assert group_sheet["F7"].value is None  # US3 shared value appears once.
        assert wb.worksheets[1].max_row == 6  # All five Amazon groups.
        assert wb.worksheets[2].max_row == 2
        wb.close()
    assert groups[-1]["ctu_over_30_cost"] == 50  # No mutation of source data.
    assert len(calls) == 7
    assert all(call[0] == "2026-09" for call in calls)
    assert ("2026-09", "EBAY-1", "OWNER") in calls


def test_no_data_and_invalid_month_do_not_export(monkeypatch, tmp_path):
    monkeypatch.setattr(service.repo, "list_groups", lambda month: {"items": []})
    monkeypatch.setattr(service, "settings", SimpleNamespace(export_output_dir=str(tmp_path)))
    with pytest.raises(ValueError, match="没有可导出"):
        service.export_monthly_summary("2026-01")
    with pytest.raises(ValueError, match="有效"):
        service.export_monthly_summary("2026-13")
    assert not list(tmp_path.iterdir())


def test_empty_dimensions_still_have_headers_and_single_us3_keeps_ctu(monkeypatch, tmp_path):
    monkeypatch.setattr(service.repo, "list_groups", lambda month: {"items": [
        {"group_code": "US1-ZXY", "ctu_over_30_cost": 7}]})
    monkeypatch.setattr(service, "list_dimensions", lambda *args, **kwargs: {"items": []})
    monkeypatch.setattr(service, "settings", SimpleNamespace(export_output_dir=str(tmp_path)))
    path, _ = service.export_monthly_summary("2026-09")
    wb = load_workbook(path)
    assert len(wb.worksheets) == 4
    assert wb.worksheets[0]["F2"].value == 7
    assert all(sheet.max_row == 1 for sheet in wb.worksheets[1:])
    wb.close()


def test_export_api_permission_validation_download_and_cleanup(monkeypatch, tmp_path):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from backend.api import deps
    from backend.api.v1 import finance

    monkeypatch.setattr(deps, "settings", SimpleNamespace(python_internal_api_token="test-internal"))
    output = tmp_path / "test.xlsx"
    calls = []

    def export(month):
        calls.append(month)
        if month == "2026-08":
            raise ValueError("没有可导出的滞销清货数据")
        wb = service.Workbook()
        wb.save(output)
        wb.close()
        return str(output), f"{month}-汇总.xlsx"

    monkeypatch.setattr(finance, "export_monthly_summary", export)
    app = FastAPI()
    app.include_router(finance.router)
    url = "/api/v1/finance/slow-moving-clearance/monthly-summary-exports"
    headers = {"X-Internal-Token": "test-internal"}
    with TestClient(app) as client:
        assert client.get(url, params={"pull_month": "2026-09"}).status_code == 401
        assert client.get(url, headers=headers, params={"pull_month": "2026-13"}).status_code == 422
        assert client.get(url, headers=headers, params={"pull_month": "2026-08"}).status_code == 404
        response = client.get(url, headers=headers, params={"pull_month": "2026-09"})
        assert response.status_code == 200
        assert response.content.startswith(b"PK")
        assert "2026-09" in response.headers["content-disposition"]
        assert not output.exists()
    assert calls == ["2026-08", "2026-09"]
