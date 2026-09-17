"""Offline import checks: frozen values, original row identity and atomic storage."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import Workbook, load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.services import ebay_inventory_history_import_service as importer
from backend.services import ebay_inventory_detail_service as service
from backend.services.ebay_inventory_detail_export_service import export_inventory, COLUMNS
from backend.repositories import ebay_inventory_pivot_repository as repo
from backend.api.v1.ebay_inventory_detail import router, require_internal_access

DAY = date(2026, 9, 7)


def book(overrides=None, *, duplicate=False, bad_date=False, bad_site=False, omit=None):
    wb = Workbook()
    wb.remove(wb.active)
    headers = ["统计时间", *importer.COLUMNS, "站点"]
    for name, site in importer.SHEETS.items():
        if name == omit:
            continue
        ws = wb.create_sheet(name)
        ws.append(["11月负责人" if name.endswith("UK") and h == "负责人" else h for h in headers])
        values = {h: 0 for h in headers}
        values.update({"统计时间": "oops" if bad_date else datetime(2026, 9, 7),
                       "产品代码": "DAS-10053-0121", "品牌": "DAS", "核心码": "00123",
                       "负责人": "历史负责人", "站点": "FR" if bad_site else site,
                       "最后售出时间": 7.3, "海外总库存": 999,
                       "在库库销比": 0.25, "近3月均销量/预估销量": 6.789})
        values.update(overrides or {})
        ws.append([values.get(h) for h in headers])
        if duplicate:
            values["海外总库存"] = 22
            ws.append([values.get(h) for h in headers])
    ignored = wb.create_sheet("销售日历订单")
    ignored.append(["not a valid historical schema"])
    data = BytesIO()
    wb.save(data)
    wb.close()
    return data.getvalue()


def parsed(**kwargs):
    return importer.parse_history(book(**kwargs), "历史.xlsx")


def test_original_rows_and_month_day_preserved_without_recalculation():
    result = parsed(duplicate=True, overrides={"产品代码": None, "海外可售": "#N/A", "海外在途": "=1+2"})
    rows = result["days"][DAY]
    assert result["source_rows"] == 6
    assert result["duplicate_rows_preserved"] == 3
    assert result["missing_sku_rows"] == 6
    assert len({(r["site"], r["record_key"]) for r in rows}) == 6
    for r in rows:
        assert r["sku"] is None
        assert r["sku_middle_site_code"] is None  # never derive from middle + site
        assert r["sku_middle_code"] == "00123"
        assert r["last_sold_at"] == "7.3"
        assert r["owner"] == "历史负责人"
        assert r["overseas_in_transit_quantity"] is None  # formula has no cache
        assert r["overseas_sellable_quantity"] is None
        assert r["overseas_total_quantity"] in {Decimal(999), Decimal(22)}
        assert r["in_stock_sales_ratio"] == Decimal("0.25")
        assert r["average_monthly_sales_3m"] == Decimal("6.789")


@pytest.mark.parametrize("value", [None, 0, "0", "0.00", "--", "#N/A", "#DIV/0!", "oops", "NaN", "Infinity", "0%"])
def test_missing_zero_error_or_invalid_number_becomes_null(value):
    row = parsed(overrides={"海外在途": value})["days"][DAY][0]
    assert row["overseas_in_transit_quantity"] is None


@pytest.mark.parametrize("kwargs", [{"bad_date": True}, {"bad_site": True}, {"omit": "库存明细持续更新-DE"}])
def test_key_or_sheet_errors_reject_file(kwargs):
    with pytest.raises(ValueError):
        parsed(**kwargs)


def test_non_numeric_cells_and_percentage_are_not_business_recomputed():
    row = parsed(overrides={"在库库销比": "25%", "最后售出时间": "07.30", "实际申购数": 20,
                            "对应仓库": "UK", "单价(含税)": -3})["days"][DAY][0]
    assert row["in_stock_sales_ratio"] == Decimal("0.25")
    assert row["last_sold_at"] == "07.30"
    assert row["source_actual_purchase_quantity"] == 20
    assert row["source_warehouse"] == "UK"
    assert row["unit_price_tax"] == -3


def test_cached_formula_value_is_read_and_never_evaluated():
    original = book({"海外在途": "=1+2"})
    modified = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(modified, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/sheet"):
                content = content.replace(b"<f>1+2</f><v></v>", b"<f>1+2</f><v>88</v>")
            target.writestr(info, content)
    rows = importer.parse_history(modified.getvalue(), "history.xlsx")["days"][DAY]
    assert all(row["overseas_in_transit_quantity"] == 88 for row in rows)  # not recalculated as 3


def fake_db(monkeypatch, existing=None):
    connection = MagicMock()
    connection.__enter__.return_value = connection
    cursor = connection.cursor.return_value.__enter__.return_value
    cursor.fetchall.return_value = existing or []
    cursor.lastrowid = 42
    monkeypatch.setattr(repo, "db_connection", lambda: connection)
    return connection, cursor


def test_atomic_store_keeps_duplicates_and_empty_sku_without_deleting(monkeypatch):
    days = parsed(duplicate=True, overrides={"产品代码": None})["days"]
    connection, cursor = fake_db(monkeypatch)
    result = repo.insert_import_days(days, "a" * 64, "历史.xlsx", "tester")
    assert result["imported_rows"] == 6
    params = cursor.executemany.call_args.args[1]
    assert len(params) == 6 and all(p[2] == "" for p in params)
    assert len({(p[1], p[3]) for p in params}) == 6
    assert not any("DELETE" in call.args[0] or "UPDATE " in call.args[0]
                   for call in cursor.execute.call_args_list)
    connection.commit.assert_called_once()


def test_conflicting_date_refused_before_any_write(monkeypatch):
    days = parsed()["days"]
    connection, cursor = fake_db(monkeypatch, [{"id": 1, "stat_date": DAY, "trigger_type": "PAGE_REFRESH",
                                               "item_count": 3, "metadata_json": {}}])
    with pytest.raises(ValueError, match="未覆盖"):
        repo.insert_import_days(days, "a" * 64, "历史.xlsx", "tester")
    cursor.executemany.assert_not_called()
    assert not any("INSERT" in call.args[0] for call in cursor.execute.call_args_list)
    connection.rollback.assert_called_once()
    connection.commit.assert_not_called()


def test_identical_file_is_idempotent_and_checks_stored_count(monkeypatch):
    days = parsed()["days"]
    connection, cursor = fake_db(monkeypatch, [{"id": 1, "stat_date": DAY, "trigger_type": "EXCEL_IMPORT",
        "item_count": 3, "metadata_json": {"import_file_sha256": "a" * 64}}])
    cursor.fetchone.return_value = {"total": 3}
    result = repo.insert_import_days(days, "a" * 64, "历史.xlsx", "tester")
    assert result["already_imported"] and result["skipped_existing_dates"] == 1
    cursor.executemany.assert_not_called()
    cursor.fetchone.return_value = {"total": 2}
    with pytest.raises(ValueError):
        repo.insert_import_days(days, "a" * 64, "历史.xlsx", "tester")


def test_write_error_rolls_back_entire_file(monkeypatch):
    connection, cursor = fake_db(monkeypatch)
    cursor.executemany.side_effect = RuntimeError("write failed")
    with pytest.raises(RuntimeError):
        repo.insert_import_days(parsed()["days"], "a" * 64, "历史.xlsx", "tester")
    connection.rollback.assert_called_once()
    connection.commit.assert_not_called()


def test_history_listing_selection_and_export_do_not_call_live_sources(monkeypatch):
    rows = parsed(duplicate=True, overrides={"产品代码": None})["days"][DAY]
    monkeypatch.setattr(service, "load_calculated_inventory", MagicMock(side_effect=AssertionError("live forbidden")))
    monkeypatch.setattr(repo, "read_inventory_day", lambda day: ([dict(r) for r in rows], {"stat_date": str(DAY)}, []))
    result = service.list_inventory(stat_date=str(DAY), paginate=False)
    assert result["pagination"]["total"] == 6
    assert all(r["sku_middle_site_code"] is None for r in result["items"])
    assert all(r["average_monthly_sales_3m"] == "6.789" for r in result["items"])
    chosen = [{"site": "美国", "sku": "", "record_key": "US:3"}]
    selected = service.list_inventory(stat_date=str(DAY), selected_keys=chosen)
    assert len(selected["items"]) == 1 and selected["items"][0]["overseas_total_quantity"] == "22"
    _, content = export_inventory(stat_date=str(DAY), selected_keys=chosen)
    wb = load_workbook(BytesIO(content), data_only=False)
    exported = list(wb.active.values)
    fields = [col[0] for col in COLUMNS]
    assert exported[1][fields.index("sku")] == "--"
    assert exported[1][fields.index("last_sold_at")] == "7.3"
    wb.close()


def test_api_import_requires_internal_auth_and_calls_importer(monkeypatch):
    app = FastAPI()
    app.include_router(router)
    @app.middleware("http")
    async def request_id(request, call_next):
        request.state.request_id = "history-test"
        return await call_next(request)
    processor = MagicMock(return_value={"imported_rows": 3})
    monkeypatch.setattr(importer, "import_history", processor)
    with TestClient(app) as client:
        response = client.post("/api/v1/finance/ebay-inventory-detail/history/import", files={"file": ("history.xlsx", b"PK")})
        assert response.status_code in {401, 403}
        processor.assert_not_called()
        app.dependency_overrides[require_internal_access] = lambda: None
        response = client.post("/api/v1/finance/ebay-inventory-detail/history/import?operator=tester", files={"file": ("history.xlsx", b"PK")})
        assert response.status_code == 200
        processor.assert_called_once_with(b"PK", "history.xlsx", "tester")
