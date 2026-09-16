from io import BytesIO
from decimal import Decimal

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.api.deps import require_internal_access
from backend.api.v1 import ebay_inventory_detail as api
from backend.services import ebay_inventory_pivot_export_service as exporter


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[require_internal_access] = lambda: None

    @app.middleware("http")
    async def request_id(request, call_next):
        request.state.request_id = "pivot-test"
        return await call_next(request)

    with TestClient(app) as value:
        yield value


def test_detail_stat_date_is_forwarded_for_list_and_export(client, monkeypatch):
    calls = []
    monkeypatch.setattr(api.service, "list_inventory", lambda **kw: calls.append(kw) or {"items": []})
    response = client.get("/api/v1/finance/ebay-inventory-detail/list",
                          params={"stat_date": "2026-09-15"})
    assert response.status_code == 200
    assert calls[-1]["stat_date"] == "2026-09-15"
    monkeypatch.setattr(api, "export_inventory", lambda **kw: calls.append(kw) or ("历史.xlsx", b"PKtest"))
    response = client.post("/api/v1/finance/ebay-inventory-detail/export",
                           json={"stat_date": "2026-09-15", "selected_keys": []})
    assert response.status_code == 200
    assert calls[-1]["stat_date"] == "2026-09-15"


def test_refresh_calls_capture_once_without_client_date_filters_or_external_sync(client, monkeypatch):
    from unittest.mock import MagicMock
    capture = MagicMock(return_value={"stat_date": "2026-09-16", "snapshot_id": 1, "item_count": 2670})
    monkeypatch.setattr(api.pivot_service, "capture_snapshot", capture)
    response = client.post("/api/v1/finance/ebay-inventory-detail/snapshot/recalculate",
                           params={"stat_date": "2020-01-01", "site": "英国"},
                           json={"stat_date": "2020-01-01", "page": 2, "sku": "one"})
    assert response.status_code == 200
    assert response.json()["data"]["stat_date"] == "2026-09-16"
    capture.assert_called_once_with(trigger_type="PAGE_REFRESH")
    assert client.get("/api/v1/finance/ebay-inventory-detail/snapshot/recalculate").status_code == 405


def test_refresh_failure_is_explicit_and_does_not_retry_or_report_success(client, monkeypatch):
    from unittest.mock import MagicMock
    capture = MagicMock(side_effect=ValueError("历史透视正在生成，请稍后重试"))
    monkeypatch.setattr(api.pivot_service, "capture_snapshot", capture)
    response = client.post("/api/v1/finance/ebay-inventory-detail/snapshot/recalculate")
    assert response.status_code == 400
    assert "正在生成" in response.json()["detail"]
    capture.assert_called_once()


def test_pivot_query_translates_filters(client, monkeypatch):
    called = {}
    def read(**kw):
        called.update(kw)
        return {"items": [], "options": {"dates": []}, "pagination": {"total": 0}}
    monkeypatch.setattr(api.pivot_service, "list_pivot", read)
    response = client.get("/api/v1/finance/ebay-inventory-detail/pivot",
                          params={"start_date": "2026-09-01", "end_date": "2026-09-16",
                                  "owner": "负责人", "site": "德国", "page": 2, "page_size": 20})
    assert response.status_code == 200
    assert response.json()["data"]["options"]["dates"] == []
    assert called["page"] == 2 and called["owner"] == "负责人"


def test_invalid_dates_return_400(client):
    response = client.get("/api/v1/finance/ebay-inventory-detail/pivot",
                          params={"start_date": "2026-09-16", "end_date": "2026-09-01"})
    assert response.status_code == 400
    assert "开始日期" in response.json()["detail"]


@pytest.mark.parametrize("params", [{"page": 0}, {"page_size": 201}, {"start_date": "2026-09-16-too-long"}])
def test_query_bounds(client, params):
    assert client.get("/api/v1/finance/ebay-inventory-detail/pivot", params=params).status_code == 422


def test_export_response_is_binary_with_utf8_filename(client, monkeypatch):
    monkeypatch.setattr(api, "export_pivot", lambda **kw: ("历史透视.xlsx", b"PK\x03\x04test"))
    response = client.get("/api/v1/finance/ebay-inventory-detail/pivot/export")
    assert response.status_code == 200
    assert response.content.startswith(b"PK")
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"


def test_empty_export_is_400(client, monkeypatch):
    def empty(**kw):
        raise ValueError("当前筛选条件下没有可导出的历史透视数据")
    monkeypatch.setattr(api, "export_pivot", empty)
    response = client.get("/api/v1/finance/ebay-inventory-detail/pivot/export")
    assert response.status_code == 400
    assert "没有" in response.json()["detail"]


def test_xlsx_matches_frozen_values_and_preserves_text(monkeypatch):
    row = {"owner": "=1+1", "site": "德国", "stat_date": "2026-09-16", "sku_count": 2,
           "overseas_sellable_quantity": "8", "overseas_total_quantity": "12", "sales_qty_30d": "10",
           "in_stock_sales_ratio": "0.800000", "total_stock_sales_ratio": "1.200000",
           "overseas_sellable_value": None, "overseas_total_value": "12.34",
           "warehouse_rent_30d_cny": "0.00"}
    called = {}
    def read(**kw):
        called.update(kw)
        return {"items": [row]}
    monkeypatch.setattr(exporter, "list_pivot", read)
    filename, content = exporter.export_pivot(site="德国")
    assert called == {"site": "德国", "paginate": False}
    assert filename.endswith(".xlsx")
    book = load_workbook(BytesIO(content))
    sheet = book.active
    assert sheet.max_row == 2 and sheet.max_column == 12
    assert sheet["A2"].value == "=1+1" and sheet["A2"].data_type == "s"
    assert sheet["H2"].value == .8 and sheet["H2"].number_format == "0.00%"
    assert sheet["J2"].value == "--"
    assert sheet["K2"].value == 12.34 and sheet["L2"].value == 0
    assert sheet.freeze_panes == "D2"
    book.close()


def test_partial_missing_export_uses_available_amounts(monkeypatch):
    from backend.services.ebay_inventory_pivot_service import aggregate_inventory
    base = {"owner": "负责人", "site": "德国", "overseas_sellable_quantity": 1,
            "overseas_total_quantity": 2, "sales_qty_30d": 3}
    row = aggregate_inventory([
        {**base, "sku": "A", "overseas_sellable_value": Decimal("12.345"),
         "overseas_total_value": Decimal("24.69"), "warehouse_rent_30d_cny": Decimal("1.005")},
        {**base, "sku": "B", "overseas_sellable_value": None,
         "overseas_total_value": None, "warehouse_rent_30d_cny": None},
    ])[0]
    row["stat_date"] = "2026-09-16"
    monkeypatch.setattr(exporter, "list_pivot", lambda **kw: {"items": [row]})
    _, content = exporter.export_pivot()
    book = load_workbook(BytesIO(content))
    assert book.active["J2"].value == 12.35
    assert book.active["K2"].value == 24.69
    assert book.active["L2"].value == 1.01
    assert book.active["D2"].value == 2  # missing money does not drop the SKU.
    book.close()


def test_export_preserves_owner_totals_order_values_and_red_style(monkeypatch):
    detail = {
        "row_type": "DETAIL", "owner": "负责人甲", "site": "德国", "stat_date": "2026-09-16",
        "sku_count": 2, "overseas_sellable_quantity": "8", "overseas_total_quantity": "12",
        "sales_qty_30d": "10", "in_stock_sales_ratio": "0.800000",
        "total_stock_sales_ratio": "1.200000", "overseas_sellable_value": None,
        "overseas_total_value": "12.34", "warehouse_rent_30d_cny": "0.00",
    }
    total = {
        **detail, "row_type": "OWNER_TOTAL", "site": "负责人汇总", "site_count": 2,
        "sku_count": 5, "overseas_sellable_quantity": "18", "overseas_total_quantity": "30",
        "sales_qty_30d": "20", "in_stock_sales_ratio": "0.900000",
        "total_stock_sales_ratio": "1.500000", "overseas_total_value": "32.34",
    }
    previous_detail = {
        **detail, "stat_date": "2026-09-09", "sales_qty_30d": "0",
        "in_stock_sales_ratio": None, "total_stock_sales_ratio": None,
    }
    rows = [
        detail,
        {**detail, "site": "英国", "sku_count": 3, "overseas_sellable_quantity": "10",
         "overseas_total_quantity": "18", "overseas_total_value": "20.00"},
        total,
        previous_detail,
        {**previous_detail, "row_type": "OWNER_TOTAL", "site": "负责人汇总", "site_count": 1},
    ]
    called = {}

    def read(**kw):
        called.update(kw)
        return {"items": rows}

    monkeypatch.setattr(exporter, "list_pivot", read)
    filters = {"owner": "负责人甲", "site": "德国", "start_date": "2026-09-09",
               "end_date": "2026-09-16"}
    _, content = exporter.export_pivot(**filters)
    assert called == {**filters, "paginate": False}

    book = load_workbook(BytesIO(content))
    sheet = book.active
    assert sheet.max_row == 6 and sheet.max_column == 12
    assert [(sheet.cell(index, 2).value, sheet.cell(index, 3).value) for index in range(2, 7)] == [
        ("德国", "2026-09-16"), ("英国", "2026-09-16"), ("负责人汇总", "2026-09-16"),
        ("德国", "2026-09-09"), ("负责人汇总", "2026-09-09"),
    ]
    assert sheet["D4"].value == 5 and sheet["E4"].value == 18
    assert sheet["H4"].value == .9 and sheet["H4"].number_format == "0.00%"
    assert sheet["I4"].value == 1.5 and sheet["I4"].number_format == "0.00%"
    assert sheet["J4"].value == "--" and sheet["K4"].value == 32.34
    assert sheet["L4"].value == 0 and sheet["L4"].number_format == "#,##0.00"
    assert sheet["G6"].value == 0 and sheet["H6"].value == "--" and sheet["I6"].value == "--"
    for row_number in (4, 6):
        for cell in sheet[row_number]:
            assert cell.font.bold is True
            assert cell.font.color.type == "rgb" and cell.font.color.rgb == "00DC2626"
            assert cell.fill.fgColor.rgb == "00FEF2F2"
    for row_number in (2, 3, 5):
        for cell in sheet[row_number]:
            assert cell.font.bold is False
            assert cell.font.color is None or cell.font.color.type != "rgb" or cell.font.color.rgb != "00DC2626"
    assert sheet["A1"].fill.fgColor.rgb == "0024486D"
    assert sheet["A2"].fill.fgColor.rgb == "00F4F7FB"
    assert sheet["A3"].fill.patternType is None
    assert sheet.auto_filter.ref == "A1:L6"
    assert sheet.freeze_panes == "D2"
    book.close()


def test_export_empty_does_not_create_workbook(monkeypatch):
    monkeypatch.setattr(exporter, "list_pivot", lambda **kw: {"items": []})
    with pytest.raises(ValueError, match="没有可导出"):
        exporter.export_pivot()
