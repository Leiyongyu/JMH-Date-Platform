"""Composite middle-code/site mapping and mixed-site search tests."""
from __future__ import annotations

from pathlib import Path
from shutil import copy2

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from backend.ebay_tool import app as tool


def run_immediate(coroutine):
    try:
        coroutine.send(None)
    except StopIteration as completed:
        return completed.value
    raise AssertionError("Unexpected asynchronous I/O in offline test")


def options():
    return dict(site="auto", free_shipping="false", search_strategy="category_lock",
                ebay_sort="price", delivery_country="", delivery_postal_code="",
                pickup_country="", pickup_postal_code="", pickup_radius="", pickup_unit="",
                price_min="", price_max="", us_postal_code="08505")


@pytest.fixture
def mapping(monkeypatch):
    monkeypatch.setattr(tool, "_mapping_schema", "site_key")
    monkeypatch.setattr(tool, "_site_key_oe_map", {
        "10144US": ("10144", "us", "OE-US"),
        "10144UK": ("10144", "uk", "OE-UK"),
        "50244US": ("50244", "us", ""),
        "30532US": ("30532", "us", "SAME-OE"),
        "30528US": ("30528", "us", "SAME-OE"),
    })


def test_exact_site_key_lookup_and_blank_oe(mapping):
    targets, matches, missing = tool._resolve_site_key_targets(
        ["10144us", "10144UK", "10144DE", "50244US", "bad"]
    )
    assert [(t["inputKey"], t["site"], t["oe"]) for t in targets] == [
        ("10144US", "us", "OE-US"), ("10144UK", "uk", "OE-UK")
    ]
    assert matches == {"10144US": "OE-US", "10144UK": "OE-UK"}
    assert any("10144DE" in error and "站点没有映射" in error for error in missing)
    assert any("50244US" in error and "为空" in error for error in missing)
    assert any("bad" in error and "格式" in error for error in missing)
    assert any("中间码不存在" in error for error in tool._resolve_site_key_targets(["99999US"])[2])


def test_loader_derives_site_from_key_when_formula_caches_are_blank(monkeypatch):
    class Sheet:
        def iter_rows(self, *, min_row, max_row=None, max_col=None, values_only):
            if max_row == 1:
                return iter([("中间码+站点", "中间码", "站点", "OE号")])
            return iter([("00123UK", None, None, "OE-UK")])

    class Workbook:
        active = Sheet()

        def close(self):
            pass

    monkeypatch.setattr(tool.os.path, "exists", lambda path: True)
    monkeypatch.setattr(tool, "load_workbook", lambda *a, **kw: Workbook())
    monkeypatch.setattr(tool, "_mapping_schema", tool._mapping_schema)
    monkeypatch.setattr(tool, "_site_key_oe_map", tool._site_key_oe_map)
    monkeypatch.setattr(tool, "_sku_oe_map", tool._sku_oe_map)
    monkeypatch.setattr(tool, "_middle_code_oe_map", tool._middle_code_oe_map)
    monkeypatch.setattr(tool, "_custom_middle_codes", tool._custom_middle_codes)
    tool.load_sku_oe_map()
    assert tool._site_key_oe_map["00123UK"] == ("00123", "uk", "OE-UK")


def test_mixed_sites_and_same_oe_keep_independent_groups(mapping, monkeypatch):
    calls = []

    def scrape(oe, site, **kwargs):
        calls.append((oe, site))
        return ([{"oe": oe, "itemId": "1", "pf": 1}], None, 0,
                "category_locked", None, [1])

    monkeypatch.setattr(tool, "scrape_one_oe", scrape)
    targets, _, _ = tool._resolve_site_key_targets(
        ["10144US", "10144UK", "30532US", "30528US"]
    )
    task_id = "offline-mixed-sites"
    try:
        tool.search_batch([], "auto", task_id, targets=targets)
        results = tool.tasks[task_id]["results"]
        assert [(r["inputKey"], r["site"], r["oe"]) for r in results] == [
            ("10144US", "US", "OE-US"), ("10144UK", "UK", "OE-UK"),
            ("30532US", "US", "SAME-OE"), ("30528US", "US", "SAME-OE")
        ]
        assert set(calls) == {("OE-US", "us"), ("OE-UK", "uk"), ("SAME-OE", "us")}
        assert results[2]["groupKey"] != results[3]["groupKey"]
        assert results[2]["items"][0]["inputKey"] == "30532US"
        assert results[3]["items"][0]["inputKey"] == "30528US"
    finally:
        tool.tasks.pop(task_id, None)


def test_text_endpoint_sends_targets_without_manual_site(mapping, monkeypatch):
    received = []
    monkeypatch.setattr(tool, "_start_search_task", lambda *a, **kw: (received.append((a, kw)) or {"ok": True}))
    result = run_immediate(tool.api_scrape(
        nkw="10144US, 10144UK, 50244US", input_type="site_key", **options()
    ))
    assert result == {"ok": True}
    assert [t["site"] for t in received[0][1]["targets"]] == ["us", "uk"]
    assert any("50244US" in error for error in received[0][0][5])


def test_unmatched_input_does_not_fall_back_to_raw_oe(mapping):
    with pytest.raises(HTTPException, match="没有可查询的映射"):
        run_immediate(tool.api_scrape(
            nkw="50244US, 10144DE", input_type="site_key", **options()
        ))


def test_real_mapping_workbook_is_four_column_source():
    source = Path(tool.SKU_OE_FILE)
    assert source.exists()
    tool.load_sku_oe_map()
    assert tool._mapping_schema == "site_key"
    assert tool._site_key_oe_map["10144US"] == ("10144", "us", "7L6616403B")
    assert tool._site_key_oe_map["50244US"][2] == ""
    assert len(tool._site_key_oe_map) >= 1000
    listing = run_immediate(tool.api_sku_list())
    assert listing["readOnly"] is False
    assert listing["siteKeyMode"] is True
    assert listing["total"] == len(tool._site_key_oe_map)
    with pytest.raises(HTTPException) as error:
        tool._require_writable_legacy_map()
    assert error.value.status_code == 409


@pytest.fixture
def editable_copy(tmp_path, monkeypatch):
    source = Path(tool.SKU_OE_FILE)
    copy_path = tmp_path / source.name
    copy2(source, copy_path)
    monkeypatch.setattr(tool, "SKU_OE_FILE", str(copy_path))
    for name in ("_site_key_oe_map", "_sku_oe_map", "_middle_code_oe_map", "_custom_middle_codes", "_mapping_schema"):
        monkeypatch.setattr(tool, name, getattr(tool, name))
    assert tool.load_sku_oe_map()
    return copy_path


def test_add_and_edit_preserve_four_columns_formulas_and_styles(editable_copy):
    original = load_workbook(editable_copy)
    original_sheet = original.active
    original_rows = original_sheet.max_row
    original_formula = original_sheet["B2"].value
    original_style = original_sheet.cell(original_rows, 1).style_id
    original.close()

    tool._save_site_key_entry(None, "00999US", "NEW-OE")
    assert tool._site_key_oe_map["00999US"] == ("00999", "us", "NEW-OE")
    tool._save_site_key_entry("50244US", "50244US", "FILLED-OE")
    assert tool._resolve_site_key_targets(["50244US"])[1] == {"50244US": "FILLED-OE"}
    tool._save_site_key_entry("00999US", "00999UK", "EDITED-OE")
    assert "00999US" not in tool._site_key_oe_map
    assert tool._site_key_oe_map["00999UK"] == ("00999", "uk", "EDITED-OE")

    saved = load_workbook(editable_copy)
    sheet = saved.active
    assert sheet.max_column == 4
    assert sheet.max_row == original_rows + 1
    assert sheet["B2"].value == original_formula
    assert sheet["D2"].value == "7L6616403B"
    assert sheet.cell(original_rows + 1, 1).value == "00999UK"
    assert sheet.cell(original_rows + 1, 2).value == f"=LEFT(A{original_rows + 1},LEN(A{original_rows + 1})-2)"
    assert sheet.cell(original_rows + 1, 3).value == f"=RIGHT(A{original_rows + 1},2)"
    assert sheet.cell(original_rows + 1, 1).style_id == original_style
    saved.close()


def test_duplicate_and_failed_save_leave_memory_and_file_unchanged(editable_copy, monkeypatch):
    before = editable_copy.read_bytes()
    in_memory = dict(tool._site_key_oe_map)
    with pytest.raises(HTTPException, match="已存在"):
        tool._save_site_key_entry(None, "10144US", "DUPLICATE")
    assert editable_copy.read_bytes() == before
    assert tool._site_key_oe_map == in_memory

    monkeypatch.setattr(tool.os, "replace", lambda *args: (_ for _ in ()).throw(PermissionError("locked")))
    with pytest.raises(HTTPException, match="文件可能正在被占用"):
        tool._save_site_key_entry("50244US", "50244US", "NEW-OE")
    assert editable_copy.read_bytes() == before
    assert tool._site_key_oe_map == in_memory


def test_add_update_api_and_specific_query_errors(editable_copy):
    class Request:
        def __init__(self, body):
            self.body = body

        async def json(self):
            return self.body

    result = run_immediate(tool.api_sku_add(Request({"sku": "00998us", "oe": "API-OE"})))
    assert result["sku"] == "00998US"
    result = run_immediate(tool.api_sku_update(Request({
        "originalSku": "00998US", "sku": "00998UK", "oe": "EDIT-OE",
    })))
    assert result["sku"] == "00998UK"
    assert tool._resolve_site_key_targets(["00998UK"])[1] == {"00998UK": "EDIT-OE"}

    with pytest.raises(HTTPException, match="中间码不存在"):
        run_immediate(tool.api_scrape(nkw="99999US", input_type="site_key", **options()))
    with pytest.raises(HTTPException, match="对应站点没有映射"):
        run_immediate(tool.api_scrape(nkw="10144DE", input_type="site_key", **options()))
    with pytest.raises(HTTPException, match="OE 号为空"):
        run_immediate(tool.api_scrape(nkw="50244US", input_type="site_key", **options()))

def test_excel_query_reads_first_column_as_site_keys():
    content = Path(tool.SKU_OE_FILE).read_bytes()
    keys, kind, total, _, _ = tool._parse_excel_keywords_from_bytes(
        content, "SKU主OE对照表.xlsx", site_key_mode=True
    )
    assert kind == "site_key"
    assert total >= 1000
    assert "10141US" in keys and "10141UK" in keys


def test_excel_endpoint_uses_site_mapping_without_manual_site(monkeypatch):
    class Upload:
        filename = "query.xlsx"

        async def read(self):
            return Path(tool.SKU_OE_FILE).read_bytes()

    received = []
    monkeypatch.setattr(tool, "_start_search_task", lambda *a, **kw: (received.append((a, kw)) or {"ok": True}))
    assert run_immediate(tool.api_scrape_file(
        file=Upload(), input_type="site_key", **options()
    )) == {"ok": True}
    targets = received[0][1]["targets"]
    assert any(t["inputKey"] == "10144US" and t["site"] == "us" for t in targets)
    assert any(t["inputKey"] == "10141UK" and t["site"] == "uk" for t in targets)
    assert any("50244US" in error for error in received[0][0][5])
