from contextlib import contextmanager
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook
from backend.services import sales_dashboard_service as service


def workbook(platforms=('amz', 'ebay'), duplicate=False, missing_metric=False, formula=False):
    book = Workbook()
    book.remove(book.active)
    for platform in platforms:
        sheet = book.create_sheet(platform)
        fields = list(service.FIELDS[platform].values())
        sheet.append(['日期', '平台', '站点', *service.FIELDS[platform]])
        row = [202601, platform, 'US', *[None if missing_metric and f == service.METRICS[platform] else 100 for f in fields]]
        sheet.append(row)
        if duplicate:
            sheet.append(row)
        if formula:
            sheet.cell(2, 4, '=1+1')
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_all_metrics_preserved_and_month_from_row_not_filename():
    result = service.parse_workbook(workbook(), '2025-12.xlsx')
    assert result['row_count'] == 2
    assert {r['month'] for r in result['scopes']} == {'2026-01'}
    for row in result['rows']:
        assert row['gross_profit'] == Decimal(100)
        assert row[service.METRICS[row['platform']]] == Decimal(100)
        assert '日期' in row['raw_data']


def test_duplicate_and_uncached_formula_rejected_before_writes():
    with pytest.raises(ValueError, match='重复'):
        service.parse_workbook(workbook(duplicate=True), 'input.xlsx')
    with pytest.raises(ValueError, match='缓存'):
        service.parse_workbook(workbook(formula=True), 'input.xlsx')


def test_missing_not_zero_and_single_platform_supported():
    result = service.preview(workbook(('ebay',), missing_metric=True), 'input.xlsx')
    assert result['scopes'][0]['amount'] is None
    assert len(result['warnings']) == 1
    assert 'rows' not in result


@pytest.mark.parametrize('value', ['202613', '2026-00', '9月', '', None])
def test_invalid_month(value):
    with pytest.raises(ValueError):
        service.month_key(value)


def connection(monkeypatch, rows=None):
    conn, cur = MagicMock(), MagicMock()
    conn.cursor.return_value.__enter__.return_value = cur
    cur.fetchone.return_value = {'acquired': 1}
    cur.fetchall.return_value = rows or []
    @contextmanager
    def db():
        yield conn
    monkeypatch.setattr(service, 'db_connection', db)
    return conn, cur


def test_replace_only_present_platform_month_and_repeat_is_idempotent(monkeypatch):
    conn, cur = connection(monkeypatch)
    content = workbook(('ebay',))
    digest = service.preview(content, 'input.xlsx')['sha256']
    for _ in range(2):
        service.import_history(content, 'input.xlsx', digest, 'tester')
    deletes = [c.args for c in cur.execute.call_args_list if c.args[0].startswith('DELETE')]
    assert deletes == [('DELETE FROM ods_sales_dashboard_ebay_monthly WHERE stat_month=%s', ('2026-01',))] * 2
    assert conn.commit.call_count == 2
    assert all('amz' not in c.args[0] for c in cur.executemany.call_args_list)


def test_bad_digest_and_invalid_sheet_never_open_db(monkeypatch):
    db = MagicMock()
    monkeypatch.setattr(service, 'db_connection', db)
    with pytest.raises(ValueError, match='预览'):
        service.import_history(workbook(), 'input.xlsx', 'bad', 'tester')
    with pytest.raises(ValueError):
        service.import_history(workbook(duplicate=True), 'input.xlsx', 'bad', 'tester')
    db.assert_not_called()


def test_two_platform_failure_rolls_back_whole_import(monkeypatch):
    conn, cur = connection(monkeypatch)
    cur.executemany.side_effect = RuntimeError('database failure')
    content = workbook()
    with pytest.raises(RuntimeError):
        service.import_history(content, 'input.xlsx', service.preview(content, 'input.xlsx')['sha256'], 'tester')
    conn.commit.assert_not_called()
    conn.rollback.assert_called_once()
    assert any('RELEASE_LOCK' in c.args[0] for c in cur.execute.call_args_list)


def test_report_range_order_null_zero_and_group_filter(monkeypatch):
    conn, cur = connection(monkeypatch, [
        {'stat_month': '2026-01', 'group_code': 'US', 'amount': Decimal(0)},
        {'stat_month': '2026-03', 'group_code': 'US', 'amount': Decimal('200.25')},
    ])
    result = service.report('ebay', '202601', '202603', ['US'])
    assert result['months'] == ['2026-01', '2026-02', '2026-03']
    assert result['total'] == ['0', None, '200.25']
    assert result['series'][0]['values'] == result['total']
    assert cur.execute.call_args.args[1] == ['2026-01', '2026-03', 'US']
    assert 'net_amount amount' in cur.execute.call_args.args[0]
    assert result['metric'] == '净额（人民币）'


def test_amz_metric_and_invalid_range(monkeypatch):
    _, cur = connection(monkeypatch)
    service.report('amz', '2026-01', '2026-02', [])
    assert 'total_receipts amount' in cur.execute.call_args.args[0]
    for args in [('other', '202601', '202602'), ('amz', '202602', '202601'), ('amz', '200001', '202601')]:
        with pytest.raises(ValueError):
            service.report(*args, [])


@pytest.mark.parametrize('platform', ['ebay', 'amz'])
def test_metric_cards_sum_range_and_weight_margin_not_average(monkeypatch, platform):
    conn, cur = connection(monkeypatch, [
        {'stat_month': '2026-01', 'group_code': 'US', 'amount': Decimal('90'), 'gross_profit': Decimal('20'), 'platform_income': Decimal('100')},
        {'stat_month': '2026-01', 'group_code': 'DE', 'amount': Decimal('80'), 'gross_profit': Decimal('10'), 'platform_income': Decimal('100')},
        {'stat_month': '2026-02', 'group_code': 'US', 'amount': Decimal('270'), 'gross_profit': Decimal('120'), 'platform_income': Decimal('300')},
        {'stat_month': '2026-02', 'group_code': 'DE', 'amount': Decimal('90'), 'gross_profit': Decimal('50'), 'platform_income': Decimal('100')},
    ])
    result = service.report(platform, '2026-01', '2026-02', ['US', 'DE'])
    metrics = result['metrics']
    assert metrics['net_sales']['summary']['value'] == '530'
    assert metrics['gross_profit']['total'] == ['30', '170']
    assert metrics['gross_profit']['summary']['value'] == '200'
    assert metrics['platform_income']['summary']['value'] == '600'
    margin = metrics['gross_margin']
    assert Decimal(margin['summary']['value']) == Decimal(200) / Decimal(600)
    assert margin['total'] == ['0.15', '0.425']
    assert margin['series'][0]['values'] == ['0.2', '0.4']
    assert margin['summary']['reason'] is None
    assert 'gross_margin' in cur.execute.call_args.args[0]  # Source value is only for detail rows.
    assert cur.execute.call_count == 1
    conn.commit.assert_not_called()


def test_metrics_missing_groups_gaps_and_null_are_not_zero(monkeypatch):
    connection(monkeypatch, [
        {'stat_month': '2026-01', 'group_code': 'US', 'amount': None, 'gross_profit': Decimal(0), 'platform_income': Decimal(100)},
        {'stat_month': '2026-03', 'group_code': 'US', 'amount': Decimal(30), 'gross_profit': Decimal(-10), 'platform_income': Decimal(100)},
    ])
    metrics = service.report('ebay', '2026-01', '2026-03', [])['metrics']
    assert metrics['gross_profit']['total'] == ['0', None, '-10']
    assert metrics['net_sales']['summary']['value'] == '30'
    assert metrics['gross_profit']['summary']['value'] == '-10'
    assert metrics['gross_profit']['summary']['missing_months'] == ['2026-02']
    assert metrics['gross_margin']['total'] == ['0', None, '-0.1']
    assert metrics['gross_margin']['summary']['reason'] is None
    assert metrics['gross_margin']['summary']['value'] == '-0.05'
    one = service.report('ebay', '2026-01', '2026-01', ['US', 'DE'])['metrics']
    assert one['gross_profit']['total'] == ['0']
    assert one['gross_profit']['summary']['value'] == '0'
    assert one['gross_profit']['series'][0]['values'] == ['0']
    assert one['gross_profit']['summary']['complete'] is False


@pytest.mark.parametrize('incomes,expected,reason', [
    ([0, 100], '0.2', None), ([100, -100], None, 'zero_denominator'),
    ([-100, -100], '-0.1', None),
])
def test_margin_handles_zero_negative_and_uses_range_denominator(monkeypatch, incomes, expected, reason):
    connection(monkeypatch, [
        {'stat_month': f'2026-0{i + 1}', 'group_code': 'US', 'amount': Decimal(0),
         'gross_profit': Decimal(10), 'platform_income': Decimal(income)}
        for i, income in enumerate(incomes)
    ])
    metrics = service.report('ebay', '2026-01', '2026-02', [])['metrics']
    assert metrics['gross_margin']['summary']['value'] == expected
    assert metrics['gross_margin']['summary']['reason'] == reason
    if incomes[0] == 0:
        assert metrics['gross_margin']['total'][0] is None
    assert metrics['net_sales']['summary']['value'] == '0'


def test_internal_api_auth_and_java_compatible_query_and_multipart(monkeypatch):
    from types import SimpleNamespace
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from backend.api import deps
    from backend.api.v1.sales_dashboard import router
    monkeypatch.setattr(deps, 'settings', SimpleNamespace(python_internal_api_token='test-only-token'))
    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as client:
        url = '/api/v1/finance/sales-dashboard'
        assert client.get(url + '/options').status_code == 401
        headers = {'X-Internal-Token': 'test-only-token'}
        content = workbook()
        result = client.post(url + '/preview', headers=headers, files={'file': ('history.xlsx', content)})
        assert result.status_code == 200
        assert result.json()['data']['row_count'] == 2
        call = MagicMock(return_value={'ok': True})
        monkeypatch.setattr(service, 'import_history', call)
        digest = result.json()['data']['sha256']
        response = client.post(url + '/import', headers=headers,
                               params={'file_sha256': digest, 'operator': 'actual-user'},
                               files={'file': ('history.xlsx', content)})
        assert response.status_code == 200
        call.assert_called_once_with(content, 'history.xlsx', digest, 'actual-user')
        query = MagicMock(return_value={'months': []})
        monkeypatch.setattr(service, 'report', query)
        assert client.get(url + '/report', headers=headers, params={'platform': 'ebay',
            'start_month': '2026-01', 'end_month': '2026-09', 'groups_csv': 'DE,US'}).status_code == 200
        query.assert_called_once_with('ebay', '2026-01', '2026-09', ['DE', 'US'])


@pytest.mark.parametrize('platform,refund_header,sales_header', [
    ('ebay', '退款', '销售总额'),
    ('amz', '退款金额（人民币）', '总销售额（人民币）'),
])
def test_refunds_use_platform_fields_weighted_range_and_preserve_sign(monkeypatch, platform, refund_header, sales_header):
    assert service.FIELDS[platform][refund_header] == 'refund_amount'
    assert service.FIELDS[platform][sales_header] == 'gross_sales'
    _, cur = connection(monkeypatch, [
        {'stat_month': '2026-01', 'group_code': 'US', 'refund_amount': Decimal('-10'), 'gross_sales': Decimal(100)},
        {'stat_month': '2026-01', 'group_code': 'DE', 'refund_amount': Decimal('-30'), 'gross_sales': Decimal(300)},
        {'stat_month': '2026-02', 'group_code': 'US', 'refund_amount': Decimal('-120'), 'gross_sales': Decimal(600)},
        {'stat_month': '2026-02', 'group_code': 'DE', 'refund_amount': Decimal(0), 'gross_sales': Decimal(0)},
    ])
    metrics = service.report(platform, '2026-01', '2026-02', ['US', 'DE'])['metrics']
    assert metrics['refund_amount']['total'] == ['-40', '-120']
    assert metrics['refund_amount']['summary']['value'] == '-160'
    assert metrics['refund_rate']['total'] == ['-0.1', '-0.2']
    assert metrics['refund_rate']['summary']['value'] == '-0.16'
    assert metrics['refund_rate']['series'][1]['values'] == ['-0.1', None]
    assert metrics['refund_rate']['summary']['complete'] is True
    assert 'refund_amount' in cur.execute.call_args.args[0]
    assert 'gross_sales' in cur.execute.call_args.args[0]


@pytest.mark.parametrize('refund,sales,expected,reason', [
    (0, 100, '0', None), (10, 100, '0.1', None), (-10, 100, '-0.1', None),
    (10, 0, None, 'zero_denominator'), (None, 100, None, 'missing_data'),
    (10, None, None, 'missing_data'),
])
def test_refund_rate_zero_missing_and_positive(monkeypatch, refund, sales, expected, reason):
    connection(monkeypatch, [{'stat_month': '2026-01', 'group_code': 'US',
                             'refund_amount': refund, 'gross_sales': sales}])
    metrics = service.report('ebay', '2026-01', '2026-01', [])['metrics']
    assert metrics['refund_rate']['summary']['value'] == expected
    assert metrics['refund_rate']['summary']['reason'] == reason
    assert metrics['refund_amount']['summary']['value'] == (str(refund) if refund is not None else None)


def test_refunds_missing_month_or_selected_group_sum_available_data(monkeypatch):
    connection(monkeypatch, [{'stat_month': '2026-01', 'group_code': 'US',
                             'refund_amount': -10, 'gross_sales': 100}])
    for end, groups in [('2026-02', []), ('2026-01', ['US', 'DE'])]:
        metrics = service.report('amz', '2026-01', end, groups)['metrics']
        for key, expected in [('refund_amount', '-10'), ('refund_rate', '-0.1')]:
            assert metrics[key]['summary']['value'] == expected
            assert metrics[key]['summary']['reason'] is None
            assert metrics[key]['summary']['complete'] is False


def test_amz_empty_september_does_not_erase_period_cards(monkeypatch):
    connection(monkeypatch, [
        {'stat_month': '2026-08', 'group_code': 'US', 'amount': 90, 'gross_profit': 20,
         'platform_income': 100, 'refund_amount': -10, 'gross_sales': 200},
        {'stat_month': '2026-09', 'group_code': 'US', 'amount': 0, 'gross_profit': None,
         'platform_income': None, 'refund_amount': None, 'gross_sales': None},
    ])
    metrics = service.report('amz', '2026-08', '2026-09', [])['metrics']
    for key, expected in [('net_sales', '90'), ('gross_profit', '20'), ('platform_income', '100'),
                          ('gross_margin', '0.2'), ('refund_amount', '-10'), ('refund_rate', '-0.05')]:
        assert metrics[key]['summary']['value'] == expected
        assert metrics[key]['summary']['reason'] is None
    assert metrics['gross_profit']['total'] == ['20', None]
    assert metrics['gross_margin']['total'] == ['0.2', None]


def test_empty_range_has_no_fabricated_zero(monkeypatch):
    connection(monkeypatch)
    for metric in service.report('amz', '2026-08', '2026-09', [])['metrics'].values():
        assert metric['summary']['value'] is None
        assert metric['summary']['reason'] == 'missing_data'
        assert metric['total'] == [None, None]


def test_partial_groups_and_independent_field_sums(monkeypatch):
    connection(monkeypatch, [
        {'stat_month': '2026-08', 'group_code': 'US', 'gross_profit': 20, 'platform_income': 100},
        {'stat_month': '2026-08', 'group_code': 'DE', 'gross_profit': None, 'platform_income': 300},
    ])
    metrics = service.report('amz', '2026-08', '2026-09', ['US', 'DE'])['metrics']
    assert metrics['gross_profit']['total'] == ['20', None]
    assert metrics['platform_income']['summary']['value'] == '400'
    assert metrics['gross_margin']['summary']['value'] == '0.05'


@pytest.mark.parametrize('platform', ['ebay', 'amz'])
def test_details_include_every_platform_field_and_filtered_scope(monkeypatch, platform):
    import json
    fields = service.FIELDS[platform]
    rows = [dict.fromkeys(fields.values(), Decimal('12.5')) for _ in range(2)]
    for i, row in enumerate(rows):
        row.update(stat_month='2026-08', group_code=['DE', 'US'][i], amount=Decimal('12.5'),
                   gross_profit=Decimal(10), platform_income=Decimal(100), gross_margin=Decimal('0.99'),
                   raw_data=json.dumps([
                       {'header': '备注', 'value': '<b>原样文本</b>', 'formula': None},
                       {'header': '毛利（RMB）', 'value': 10},
                       {'header': '日期', 'value': 202608},
                       {'header': None, 'value': 'ignored'},
                   ]))
    rows[1]['advertising_fee'] = None
    conn, cur = connection(monkeypatch, rows)
    report = service.report(platform, '2026-08', '2026-09', ['DE', 'US'])
    details = report['details']
    assert [c['key'] for c in details['columns']] == [*fields.values(), 'extra_0']
    assert [c['label'] for c in details['columns']] == [*fields.keys(), '备注']
    assert len(details['rows']) == 2  # Never synthesize missing September rows.
    for row in details['rows']:
        assert row['platform'] == platform
        assert row['gross_margin'] == '0.99'  # Source ratio stays available for checking.
        assert row['extra_0'] == '<b>原样文本</b>'
        assert 'raw_data' not in row
        assert 'source_file' not in row
    assert details['rows'][1]['advertising_fee'] is None
    assert details['summary']['advertising_fee'] == '12.5'
    assert details['summary']['gross_margin'] == report['metrics']['gross_margin']['summary']['value'] == '0.1'
    assert details['summary'][service.METRICS[platform]] == report['metrics']['net_sales']['summary']['value'] == '25.0'
    assert details['summary']['gross_profit'] == '20'
    assert 'raw_data' in cur.execute.call_args.args[0]
    assert cur.execute.call_args.args[1] == ['2026-08', '2026-09', 'DE', 'US']
    conn.commit.assert_not_called()


def test_empty_detail_report_still_returns_platform_headers(monkeypatch):
    connection(monkeypatch)
    details = service.report('amz', '2026-09', '2026-09', [])['details']
    assert len(details['columns']) == len(service.FIELDS['amz'])
    assert details['rows'] == []
    assert all(value is None for value in details['summary'].values())
