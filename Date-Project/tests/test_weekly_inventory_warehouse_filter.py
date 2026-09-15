from copy import deepcopy
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from backend.services import weekly_inventory_export_service as export
from backend.services import weekly_inventory_sync_service as sync


EU_WID = 18677
UK_WID = 19561
OTHER_WID = 99001
NAMES = {
    EU_WID: "CTUAMZ-EU中转仓",
    UK_WID: "CTUAMZ-UK中转仓",
    # A different warehouse sharing an allowed display name must stay excluded.
    OTHER_WID: "CTUAMZ-EU中转仓",
}


def snapshot(warehouse_ids=(EU_WID, UK_WID, OTHER_WID)):
    quantities = {EU_WID: 10, UK_WID: 4, OTHER_WID: 999}
    prices = {EU_WID: "2.5", UK_WID: "3.25", OTHER_WID: "100"}
    buckets = {EU_WID: "0-29天", UK_WID: "30-89天", OTHER_WID: "其他仓独有库龄档"}
    locked = {EU_WID: 2, UK_WID: 1, OTHER_WID: 100}
    inventory, bins = [], []
    for wid in warehouse_ids:
        inventory.append(dict(
            wid=wid, product_id=7, seller_id="seller", sku="SHARED-SKU",
            product_total=quantities[wid], purchase_price=prices[wid],
            product_valid_num=quantities[wid] - locked[wid],
            stock_age_list=[dict(name=buckets[wid], qty=quantities[wid])],
        ))
        bins.append(dict(
            wid=wid, whb_id=1, product_id=7, wh_name=NAMES[wid],
            whb_name=f"BIN-{wid}",
            total=quantities[wid] + 5,
            lockNum=locked[wid], validNum=quantities[wid] - locked[wid],
        ))
    products = [dict(id=7, sku="SHARED-SKU", product_name="共同产品")]
    if OTHER_WID in warehouse_ids:
        inventory.append(dict(
            wid=OTHER_WID, product_id=8, sku="EXCLUDED-SKU",
            product_total=500, purchase_price="99",
        ))
        products.append(dict(id=8, sku="EXCLUDED-SKU", product_name="其他仓独有产品"))
    return sync.normalize(
        inventory, bins, products, date(2026, 9, 14), "filter-test",
        datetime(2026, 9, 14, 12, 0, 0),
    )


def read_export(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = workbook.active.iter_rows(values_only=True)
        headers = list(next(values))
        rows = [dict(zip(headers, row)) for row in values]
        return headers, rows
    finally:
        workbook.close()


def test_export_keeps_only_target_ids_without_merging_shared_sku(tmp_path):
    path = tmp_path / "two-warehouses.xlsx"
    metrics = export.write_export(snapshot(), NAMES, path)
    headers, rows = read_export(path)

    assert metrics["row_count"] == len(rows) == 2
    assert metrics["column_count"] == len(headers) == 26  # 24 fixed + 2 age buckets
    assert {row["仓库名称"] for row in rows} == {NAMES[EU_WID], NAMES[UK_WID]}
    assert [row["SKU"] for row in rows] == ["SHARED-SKU", "SHARED-SKU"]
    by_warehouse = {row["仓库名称"]: row for row in rows}
    eu, uk = by_warehouse[NAMES[EU_WID]], by_warehouse[NAMES[UK_WID]]
    assert (eu["实际库存总量"], eu["采购单价"], eu["库存金额"]) == (10, 2.5, 25)
    assert (uk["实际库存总量"], uk["采购单价"], uk["库存金额"]) == (4, 3.25, 13)
    assert (eu["锁定量(仓位)"], eu["未锁定量(仓位)"]) == (2, 8)
    assert (uk["锁定量(仓位)"], uk["未锁定量(仓位)"]) == (1, 3)
    assert (eu['总量(仓位)'], uk['总量(仓位)']) == (15, 9)
    assert eu['仓位名称'] == f'BIN-{EU_WID}'
    assert uk['仓位名称'] == f'BIN-{UK_WID}'
    assert eu["0-29天"] == 10 and eu["30-89天"] is None
    assert uk["0-29天"] is None and uk["30-89天"] == 4
    assert "其他仓独有库龄档" not in headers


def test_export_uses_id_even_when_target_names_change(tmp_path):
    path = tmp_path / "renamed.xlsx"
    names = {**NAMES, EU_WID: "EU仓新名称", UK_WID: "UK仓新名称"}
    export.write_export(snapshot(), names, path)
    _, rows = read_export(path)
    assert {row["仓库名称"] for row in rows} == {"EU仓新名称", "UK仓新名称"}
    assert len(rows) == 2


def test_multiple_bin_rows_stay_with_their_warehouse_and_product(tmp_path):
    groups = snapshot()
    eu_bin = next(row for row in groups['bins'] if row['wid'] == EU_WID)
    uk_bin = next(row for row in groups['bins'] if row['wid'] == UK_WID)
    groups['bins'].extend([
        {**eu_bin, 'whb_id': 2, 'whb_name': 'EU-SECOND'},
        {**uk_bin, 'whb_id': 2, 'whb_name': 'UK-SECOND'},
        {**eu_bin, 'whb_id': 3, 'product_id': 8, 'whb_name': 'UNRELATED-PRODUCT'},
    ])
    before = deepcopy(groups)
    path = tmp_path / 'two-warehouses-multiple-bins.xlsx'
    metrics = export.write_export(groups, NAMES, path)
    _, rows = read_export(path)
    assert metrics['row_count'] == len(rows) == 4
    for wid, bin_name, quantity, cost, locked, unlocked in [
        (EU_WID, 'EU-SECOND', 10, 25, 2, 8),
        (UK_WID, 'UK-SECOND', 4, 13, 1, 3),
    ]:
        warehouse_rows = [row for row in rows if row['仓库名称'] == NAMES[wid]]
        assert [row.pop('仓位名称') for row in warehouse_rows] == [f'BIN-{wid}', bin_name]
        assert len(warehouse_rows) == 2 and warehouse_rows[0] == warehouse_rows[1]
        row = warehouse_rows[0]
        assert row['SKU'] == 'SHARED-SKU'
        assert (row['实际库存总量'], row['库存金额']) == (quantity, cost)
        assert (row['锁定量(仓位)'], row['未锁定量(仓位)']) == (locked, unlocked)
        assert row['总量(仓位)'] == quantity + 5
        assert row['0-29天'] == (10 if wid == EU_WID else None)
        assert row['30-89天'] == (4 if wid == UK_WID else None)
    assert groups == before


def test_same_bin_name_is_aggregated_only_within_warehouse_and_product(tmp_path):
    groups = snapshot()
    for row in groups['bins']:
        row['whb_name'] = ' SHARED-BIN '
    eu_bin = next(row for row in groups['bins'] if row['wid'] == EU_WID)
    groups['bins'].extend([
        {**eu_bin, 'whb_id': 2, 'store_id': 'another-shop', 'whb_name': 'SHARED-BIN',
         'total': 7, 'lock_num': 0, 'valid_num': 6},
        {**eu_bin, 'whb_id': 3, 'product_id': 8, 'total': 999,
         'lock_num': 999, 'valid_num': 999},
    ])
    before = deepcopy(groups)
    path = tmp_path / 'same-bin-name.xlsx'
    metrics = export.write_export(groups, NAMES, path)
    _, rows = read_export(path)
    assert metrics['row_count'] == len(rows) == 2
    by_warehouse = {row['仓库名称']: row for row in rows}
    for wid, total, locked, valid in [(EU_WID, 22, 2, 14), (UK_WID, 9, 1, 3)]:
        row = by_warehouse[NAMES[wid]]
        assert row['仓位名称'] == 'SHARED-BIN'
        assert (row['总量(仓位)'], row['锁定量(仓位)'], row['未锁定量(仓位)']) == (total, locked, valid)
    assert groups == before


def test_export_does_not_modify_full_snapshot(tmp_path):
    groups = snapshot()
    before = deepcopy(groups)
    names_before = deepcopy(NAMES)
    export.write_export(groups, NAMES, tmp_path / "preserved.xlsx")
    assert groups == before
    assert NAMES == names_before
    assert any(row["wid"] == OTHER_WID for row in groups["inventory"])
    assert any(row["product_id"] == 8 for row in groups["products"])


def test_no_target_warehouse_rejected_without_publishing_file(tmp_path):
    path = tmp_path / "empty.xlsx"
    groups = snapshot((OTHER_WID,))
    with pytest.raises(ValueError, match="没有可导出"):
        export.write_export(groups, NAMES, path)
    assert not path.exists()
    assert not list(tmp_path.glob(".weekly-*"))


@pytest.mark.parametrize("wid,quantity,cost", [(EU_WID, 10, 25), (UK_WID, 4, 13)])
def test_only_one_target_warehouse_is_sufficient(tmp_path, wid, quantity, cost):
    path = tmp_path / f"only-{wid}.xlsx"
    metrics = export.write_export(snapshot((wid, OTHER_WID)), NAMES, path)
    _, rows = read_export(path)
    assert metrics["row_count"] == len(rows) == 1
    assert rows[0]["仓库名称"] == NAMES[wid]
    assert rows[0]["实际库存总量"] == quantity
    assert rows[0]["库存金额"] == cost
