"""Offline eBay search tests; no credentials, database or external requests."""
import json

import pytest
from fastapi import BackgroundTasks, HTTPException, Request
from openpyxl import load_workbook

from backend.ebay_tool import app as tool


def run_immediate(coroutine):
    """Drive endpoints with in-memory bodies without Windows socketpair setup."""
    try:
        coroutine.send(None)
    except StopIteration as completed:
        return completed.value
    raise AssertionError("Unexpected asynchronous I/O in offline test")


class StubUpload:
    filename = "test.xlsx"

    async def read(self):
        return b"dummy"


def raw_items(count=100):
    return [
        {"itemId": str(i), "title": f"Part {i}",
         "price": {"value": str(count - i), "currency": "EUR"},
         "categories": [{"categoryId": "cat"}],
         "itemCreationDate": "2026-10-01T00:00:00Z",
         "itemEndDate": "2026-11-01T00:00:00Z"}
        for i in range(count)
    ]


@pytest.mark.parametrize("selected,expected", list(tool.EBAY_SORTS.items()))
def test_sort_mapping_and_ebay_order(monkeypatch, selected, expected):
    calls = []

    def search(*args, **kwargs):
        calls.append(kwargs)
        return {"itemSummaries": raw_items()}, None

    monkeypatch.setattr(tool, "search_ebay", search)
    monkeypatch.setattr(tool, "get_item_product_type", lambda *args: "Part")
    monkeypatch.setattr(tool, "_pick_dominant_category", lambda *args: ("cat", "Parts", .8, []))
    items, warning, _, strategy, _, _ = tool.scrape_one_oe(
        "OE-1", "de", ebay_sort=selected,
        pickup_country="DE" if selected == "distance" else "",
        pickup_postal_code="10115" if selected == "distance" else "",
        pickup_radius=20 if selected == "distance" else None,
        pickup_unit="km" if selected == "distance" else "",
    )
    assert warning is None and strategy == "category_locked"
    assert len(items) == 100
    assert [item["itemId"] for item in items] == [str(i) for i in range(100)]
    assert [call["sort"] for call in calls] == [expected, expected]
    assert calls[0]["fieldgroups"] == "MATCHING_ITEMS,CATEGORY_REFINEMENTS"
    assert calls[1]["category_id"] == "cat"
    assert all(call["limit"] == 100 for call in calls)


@pytest.mark.parametrize("reason", ["no_category", "search_error", "empty_category"])
def test_category_fallback_requeries_with_selected_sort(monkeypatch, reason):
    calls = []

    def search(*args, **kwargs):
        calls.append(kwargs)
        if len(calls) == 2 and reason == "search_error":
            return None, "category unavailable"
        if len(calls) == 2 and reason == "empty_category":
            return {"itemSummaries": []}, None
        return {"itemSummaries": raw_items()}, None

    monkeypatch.setattr(tool, "search_ebay", search)
    monkeypatch.setattr(tool, "get_item_product_type", lambda *args: "Part")
    cat = "" if reason == "no_category" else "cat"
    monkeypatch.setattr(tool, "_pick_dominant_category", lambda *args: (cat, "Parts", .8, []))
    items, warning, _, strategy, info, _ = tool.scrape_one_oe(
        "OE-1", "de", ebay_sort="newlyListed", price_min="100", price_max="200"
    )
    assert warning is None and strategy == "category_fallback"
    assert info["fallbackReason"]
    assert [item["itemId"] for item in items] == [str(i) for i in range(100)]
    assert len(calls) == (2 if reason == "no_category" else 3)
    assert calls[-1]["sort"] == "newlyListed" and "category_id" not in calls[-1]
    assert all(call["price_min"] == "100" and call["price_max"] == "200" for call in calls)


def test_failed_fallback_does_not_present_probe_as_sorted_results(monkeypatch):
    calls = []

    def search(*args, **kwargs):
        calls.append(kwargs)
        return ({"itemSummaries": raw_items()}, None) if len(calls) == 1 else (None, "sort unavailable")

    monkeypatch.setattr(tool, "search_ebay", search)
    monkeypatch.setattr(tool, "_pick_dominant_category", lambda *args: ("", "", 0, []))
    items, warning, _, strategy, _, _ = tool.scrape_one_oe("OE-1", "de", ebay_sort="endingSoonest")
    assert items == [] and warning == "sort unavailable" and strategy == "error"


def test_verification_checks_first_five_without_moving_suspect(monkeypatch):
    checked = []

    def verify(item_id, marketplace):
        checked.append(item_id)
        return "Other" if item_id == "1" else "Part"

    monkeypatch.setattr(tool, "get_item_product_type", verify)
    items, *_ = tool._finalize_oe(raw_items(130), "OE-1", "category_locked", "EBAY_DE", None, verify=True)
    assert len(items) == 100
    assert checked == ["0", "1", "2", "3", "4"]
    assert [item["itemId"] for item in items[:5]] == checked
    assert items[1]["suspect"] is True
    assert items[0]["listedAt"] and items[0]["endsAt"]


def test_legacy_defaults_and_short_results(monkeypatch):
    calls = []

    def search(*args, **kwargs):
        calls.append(kwargs)
        return {"itemSummaries": raw_items(7)}, None

    monkeypatch.setattr(tool, "search_ebay", search)
    monkeypatch.setattr(tool, "get_item_product_type", lambda *args: "Part")
    monkeypatch.setattr(tool, "_pick_dominant_category", lambda *args: ("cat", "Parts", .8, []))
    assert len(tool.scrape_one_oe("OE-1", "de", search_strategy="best_match")[0]) == 7
    assert calls[-1]["sort"] == ""
    assert len(tool.scrape_one_oe("OE-1", "de", search_strategy="category_lock")[0]) == 7
    assert calls[-1]["sort"] == "price"


def test_batch_records_actual_sort_fallback_and_fetched_count(monkeypatch):
    def scrape(*args, **kwargs):
        info = {"fallbackReason": "无法确定主要类目"}
        items = tool._finalize_oe(raw_items(2), "OE-1", "category_fallback", "EBAY_DE", info)[0]
        return items, None, 0, "category_fallback", info, [2, 1]

    monkeypatch.setattr(tool, "scrape_one_oe", scrape)
    task_id = "offline-test-sort"
    try:
        tool.search_batch(["OE-1"], "de", task_id, ebay_sort="newlyListed")
        result = tool.tasks[task_id]["results"][0]
        assert result["actualSort"] == "newlyListed"
        assert result["categoryFallback"] is True
        assert result["fetchedCount"] == 2
    finally:
        tool.tasks.pop(task_id, None)


def test_filters_and_delivery_context(monkeypatch):
    captured = []
    monkeypatch.setattr(tool, "_api_get", lambda *a, **kw: (captured.append((a, kw)) or {}, None))
    tool.search_ebay("OE-1", "EBAY_DE", sort="distance", pickup_country="DE",
                     pickup_postal_code="10115", pickup_radius=20, pickup_unit="km")
    assert "deliveryOptions:{SELLER_ARRANGED_LOCAL_PICKUP}" in captured[-1][0][1]["filter"]
    assert "pickupCountry:DE" in captured[-1][0][1]["filter"]
    tool.search_ebay("OE-1", "EBAY_DE", sort="price", free_shipping=True,
                     delivery_country="DE", delivery_postal_code="10115")
    assert "maxDeliveryCost:0" in captured[-1][0][1]["filter"]
    assert "deliveryCountry:DE" in captured[-1][0][1]["filter"]
    assert "deliveryPostalCode:10115" in captured[-1][0][1]["filter"]
    assert captured[-1][1]["delivery_country"] == "DE"


@pytest.mark.parametrize("low,high,expected", [
    ("100", "200", "price:[100..200]"),
    ("0", "200", "price:[0..200]"),
    ("100.50", "", "price:[100.50]"),
    ("", "200", "price:[..200]"),
    ("", "", None),
])
def test_query_price_filter_bounds(monkeypatch, low, high, expected):
    captured = []
    monkeypatch.setattr(tool, "_api_get", lambda *a, **kw: (captured.append(a[1]) or {}, None))
    tool.search_ebay("OE-1", "EBAY_US", price_min=low, price_max=high)
    filters = captured[0]["filter"].split(",")
    if expected:
        assert expected in filters and "priceCurrency:USD" in filters
    else:
        assert all(not part.startswith(("price:", "priceCurrency:")) for part in filters)


def test_mixed_site_batch_sends_prices_to_ebay_in_site_currency(monkeypatch):
    captured = []

    def get(url, params, marketplace, **kwargs):
        captured.append((url, params, marketplace))
        return {"itemSummaries": raw_items(1)}, None

    monkeypatch.setattr(tool, "_api_get", get)
    targets = [{"inputKey": f"10144{site.upper()}", "site": site, "oe": "OE-1",
                "groupKey": f"10144{site.upper()}"} for site in ("us", "uk", "de")]
    task_id = "offline-price-range"
    try:
        tool.search_batch([], "auto", task_id, search_strategy="best_match",
                          targets=targets, price_min="100", price_max="200")
        assert len(captured) == 3
        currencies = {"EBAY_US": "USD", "EBAY_GB": "GBP", "EBAY_DE": "EUR"}
        for url, params, marketplace in captured:
            assert url == tool.BROWSE_URL
            assert "price:[100..200]" in params["filter"]
            assert f"priceCurrency:{currencies[marketplace]}" in params["filter"]
        results = tool.tasks[task_id]["results"]
        assert [r["queryPrice"]["currency"] for r in results] == ["USD", "GBP", "EUR"]
    finally:
        tool.tasks.pop(task_id, None)


@pytest.mark.parametrize("us_postal,expected_us", [("08505", "08505"), ("92236", "92236")])
def test_mixed_site_delivery_filters_follow_each_site(monkeypatch, us_postal, expected_us):
    captured = {}

    def get(url, params, marketplace, **kwargs):
        captured[marketplace] = (params["filter"], kwargs)
        return {"itemSummaries": raw_items(1)}, None

    monkeypatch.setattr(tool, "_api_get", get)
    targets = [{"inputKey": f"10144{site.upper()}", "site": site, "oe": "OE-1",
                "groupKey": f"10144{site.upper()}"} for site in ("us", "uk", "de")]
    task_id = "offline-delivery-location"
    try:
        tool.search_batch([], "auto", task_id, search_strategy="best_match",
                          targets=targets, us_postal_code=us_postal)
        assert "deliveryCountry:US" in captured["EBAY_US"][0]
        assert f"deliveryPostalCode:{expected_us}" in captured["EBAY_US"][0]
        assert "deliveryCountry:DE" in captured["EBAY_DE"][0]
        assert "deliveryPostalCode:42489" in captured["EBAY_DE"][0]
        assert "deliveryCountry:GB" in captured["EBAY_GB"][0]
        assert "deliveryPostalCode:" not in captured["EBAY_GB"][0]
        assert captured["EBAY_GB"][1]["delivery_postal_code"] == ""
        locations = {row["site"]: row["queryDelivery"] for row in tool.tasks[task_id]["results"]}
        assert locations == {
            "US": {"country": "US", "postalCode": expected_us},
            "UK": {"country": "GB", "postalCode": ""},
            "DE": {"country": "DE", "postalCode": "42489"},
        }
    finally:
        tool.tasks.pop(task_id, None)


def test_delivery_context_header_is_encoded_and_no_placeholder_affiliate(monkeypatch):
    captured = []
    monkeypatch.setattr(tool, "get_access_token", lambda: "unit-token")

    class Response:
        status_code = 200

        def json(self):
            return {"itemSummaries": []}

    def get(*args, **kwargs):
        captured.append(kwargs["headers"])
        return Response()

    monkeypatch.setattr(tool._requests, "get", get)
    tool._api_get("https://example.invalid", {}, "EBAY_GB", delivery_country="GB",
                  delivery_postal_code="SW1A 1AA")
    assert captured[0]["X-EBAY-C-ENDUSERCTX"] == "contextualLocation=country%3DGB%2Czip%3DSW1A%201AA"
    assert "affiliateCampaignId" not in captured[0]["X-EBAY-C-ENDUSERCTX"]


@pytest.mark.parametrize("fields", [
    {"ebay_sort": "invalid"},
    {"ebay_sort": "distance", "free_shipping": "true", "pickup_country": "DE",
     "pickup_postal_code": "10115", "pickup_radius": "20", "pickup_unit": "km"},
    {"ebay_sort": "distance", "pickup_country": "DE", "pickup_postal_code": "10115",
     "pickup_radius": "0", "pickup_unit": "km"},
    {"price_min": "200", "price_max": "100"},
    {"price_min": "-1"},
    {"price_max": "NaN"},
    {"price_min": "10.123"},
    {"price_max": "100],sellers:{invalid}"},
    {"us_postal_code": "00000"},
])
def test_both_entry_points_reject_invalid_options(monkeypatch, fields):
    monkeypatch.setattr(tool, "_parse_excel_keywords_from_bytes", lambda *a, **kw: (["OE-1"], "oe", 1, 0, "test.xlsx"))
    options = dict(site="de", free_shipping="false", search_strategy="category_lock",
                   ebay_sort=None, delivery_country="", delivery_postal_code="",
                   pickup_country="", pickup_postal_code="", pickup_radius="", pickup_unit="",
                   price_min="", price_max="", us_postal_code="08505")
    options.update(fields)
    with pytest.raises(HTTPException) as text_error:
        run_immediate(tool.api_scrape(nkw="OE-1", input_type="oe", **options))
    with pytest.raises(HTTPException) as excel_error:
        run_immediate(tool.api_scrape_file(
            file=StubUpload(), **options
        ))
    assert text_error.value.status_code == excel_error.value.status_code == 400


def test_text_and_excel_submit_identical_search_options(monkeypatch):
    received = []
    monkeypatch.setattr(tool, "_start_search_task", lambda *a, **kw: (received.append(kw) or {"ok": True}))
    monkeypatch.setattr(tool, "_parse_excel_keywords_from_bytes", lambda *a, **kw: (["OE-1"], "oe", 1, 0, "test.xlsx"))
    fields = dict(site="de", free_shipping="false", search_strategy="category_lock",
                  ebay_sort="price", delivery_country="DE", delivery_postal_code="10115",
                  pickup_country="", pickup_postal_code="", pickup_radius="", pickup_unit="",
                  price_min="100", price_max="200", us_postal_code="08505")
    assert run_immediate(tool.api_scrape(nkw="OE-1", input_type="oe", **fields)) == {"ok": True}
    assert run_immediate(tool.api_scrape_file(
        file=StubUpload(), **fields
    )) == {"ok": True}
    assert received[0] == received[1]
    assert received[0]["price_min"] == "100" and received[0]["price_max"] == "200"


def test_export_preserves_input_order_and_includes_dates():
    items = [
        {"oe": "OE-1", "inputKey": "10144US", "site": "US",
         "price": "30 EUR", "title": "First", "images": [],
         "listedAt": "2026-10-03T00:00:00Z", "endsAt": "2026-11-03T00:00:00Z"},
        {"oe": "OE-1", "price": "10 EUR", "title": "Second", "images": []},
    ]
    payload = json.dumps({"items": items}).encode()

    async def receive():
        return {"type": "http.request", "body": payload, "more_body": False}

    request = Request({"type": "http", "method": "POST", "headers": []}, receive)
    response = run_immediate(tool.api_export(request, BackgroundTasks()))
    try:
        ws = load_workbook(response.path).active
        assert [ws.cell(row=i, column=3).value for i in (2, 3)] == ["First", "Second"]
        assert ws.cell(row=2, column=7).value == "2026-10-03T00:00:00Z"
        assert ws.cell(row=2, column=8).value == "2026-11-03T00:00:00Z"
        assert ws.cell(row=2, column=11).value == "10144US"
        assert ws.cell(row=2, column=12).value == "US"
    finally:
        tool._safe_remove(response.path)
