"""价格导入离线测试；只使用内存生成的工作簿，不读用户文件、不连接数据库。"""
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from backend.services import ebay_inventory_price_parser as parser


def workbook_bytes(rows, *, extra_sheets=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "采购价"
    for row in rows:
        sheet.append(row)
    for title, values in (extra_sheets or {}).items():
        extra = workbook.create_sheet(title)
        for row in values:
            extra.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def standard_rows(*rows):
    return [["产品代码", "单价(默认采购价)"], *rows]


def parse_rows(*rows):
    return parser.parse_prices(workbook_bytes(standard_rows(*rows)), "单价.xlsx")


def test_real_headers_and_sku_normalization_preserve_middle_leading_zero():
    result = parse_rows([" frd-070618-0687 ", "12.3400"], ["BMW-30055-0182", 0])
    assert result["rows"] == [
        {"sku": "FRD-070618-0687", "middle_code": "070618", "unit_price": Decimal("12.34")},
        {"sku": "BMW-30055-0182", "middle_code": "30055", "unit_price": Decimal("0")},
    ]
    assert result["source_rows"] == result["sku_count"] == result["middle_code_count"] == 2
    assert result["skipped_rows"] == result["duplicate_rows"] == result["warning_count"] == 0
    assert result["unmatched_middle_rows"] == 0
    assert result["warnings"] == []
    assert all(isinstance(row["unit_price"], Decimal) for row in result["rows"])


@pytest.mark.parametrize("sku_header", ["SKU", " sku ", "产品代码"])
@pytest.mark.parametrize("price_header", [
    "单价(默认采购价)", "单价（默认采购价）", "单价（含税）", "单价", "采购价",
])
def test_only_supported_header_aliases_are_accepted(sku_header, price_header):
    content = workbook_bytes([[price_header, sku_header], [1.25, "A-001-B"]])
    assert parser.parse_prices(content, "price.XLSX")["rows"][0]["unit_price"] == Decimal("1.25")


def test_dedup_uses_complete_sku_and_exact_price_across_sheets():
    content = workbook_bytes(
        standard_rows(["A-123-X", "10.0000000"], ["A-123-X", "11"]),
        extra_sheets={"补充": standard_rows([" a-123-x ", 10], ["B-123-X", 10])},
    )
    result = parser.parse_prices(content, "price.xlsx")
    assert result["source_rows"] == 4
    assert result["duplicate_rows"] == 1
    assert result["sku_count"] == 2
    assert result["middle_code_count"] == 1
    assert [(row["sku"], row["unit_price"]) for row in result["rows"]] == [
        ("A-123-X", Decimal("10")), ("A-123-X", Decimal("11")), ("B-123-X", Decimal("10")),
    ]


@pytest.mark.parametrize("value,expected", [
    (0, "0"), ("-0.0000000", "0"), ("0.123456", "0.123456"),
    ("12.340000000", "12.34"), ("1e-6", "0.000001"),
    ("999999999999999999.999999", "999999999999999999.999999"),
])
def test_decimal_values_within_storage_precision_are_exact(value, expected):
    assert parse_rows(["A-123-X", value])["rows"][0]["unit_price"] == Decimal(expected)


@pytest.mark.parametrize("value", [
    None, "", " ", -1, "-0.000001", "NaN", "Infinity", "-Infinity", "nonsense",
    True, "1,000", "1_000", "0.1234567", "1e-7", "1000000000000000000", "1e9999",
])
def test_invalid_price_rejects_entire_file_with_row_number(value):
    with pytest.raises(ValueError, match="采购价 第3行.*单价.*整份未导入"):
        parse_rows(["GOOD-123-X", 10], ["BAD-456-X", value])


@pytest.mark.parametrize("sku,price", [("#N/A", 10), ("A-123-X", "#VALUE!"), ("A-123-X", "#DIV/0!")])
def test_excel_error_values_reject_entire_file(sku, price):
    with pytest.raises(ValueError, match="第2行.*Excel错误值"):
        parse_rows([sku, price])


def test_formula_without_cached_price_is_rejected_not_zero():
    with pytest.raises(ValueError, match="第3行.*公式.*缓存"):
        parse_rows(["GOOD-123-X", 10], ["FORMULA-456-X", "=1+1"])


@pytest.mark.parametrize("sku", [None, "", " " * 4, "X" * 256])
def test_missing_or_overlong_sku_is_rejected(sku):
    with pytest.raises(ValueError, match="第2行.*SKU"):
        parse_rows([sku, 10])


@pytest.mark.parametrize("sku", [
    "IDD-LMM-310002-0047", "GM-40031B-0098", "NOHYPHEN", "A--X", "A-１２３-X", "A-" + "1" * 65 + "-X",
])
def test_missing_ascii_middle_keeps_full_sku_price_with_warning(sku):
    result = parse_rows([sku, 10])
    assert result["rows"] == [{"sku": sku, "middle_code": None, "unit_price": Decimal("10")}]
    assert result["source_rows"] == result["sku_count"] == result["unmatched_middle_rows"] == 1
    assert result["warning_count"] == len(result["warnings"]) == 1
    assert "第2行" in result["warnings"][0]
    assert result["middle_code_count"] == result["skipped_rows"] == 0


def test_middle_rule_accepts_second_segment_only_and_trims_it():
    result = parse_rows(["A-001", 1], ["B- 002 -X", 2])
    assert [row["middle_code"] for row in result["rows"]] == ["001", "002"]


def test_warning_details_capped_but_invalid_middle_rows_all_counted():
    result = parse_rows(*[[f"A-CODE-{index}", 10] for index in range(205)])
    assert result["unmatched_middle_rows"] == result["warning_count"] == 205
    assert len(result["warnings"]) == 200
    assert result["sku_count"] == len(result["rows"]) == 205


def test_invalid_middle_duplicates_still_deduplicate_exact_sku_price():
    result = parse_rows(["A-CODE-X", 10], ["a-code-x", "10.0"], ["A-CODE-X", 11])
    assert result["unmatched_middle_rows"] == 3
    assert result["duplicate_rows"] == 1
    assert result["sku_count"] == 1
    assert len(result["rows"]) == 2


def test_explicit_matching_middle_is_validated_without_losing_zeroes():
    content = workbook_bytes([["SKU", "单价", "中间码"], ["A-001-X", 2, "001"], ["A-CODE-X", 3, None]])
    result = parser.parse_prices(content, "price.xlsx")
    assert [row["middle_code"] for row in result["rows"]] == ["001", None]


@pytest.mark.parametrize("sku,middle", [("A-001-X", "1"), ("A-001-X", 1), ("A-001-X", None), ("A-CODE-X", "001")])
def test_explicit_middle_cannot_override_derived_middle(sku, middle):
    content = workbook_bytes([["SKU", "单价", "中间码"], [sku, 2, middle]])
    with pytest.raises(ValueError, match="第2行.*中间码.*不一致"):
        parser.parse_prices(content, "price.xlsx")


def test_header_may_follow_title_and_blank_rows_with_reordered_columns():
    content = workbook_bytes([["价格清单"], [], ["备注", "单价", "SKU"], ["描述", 2, "A-123-X"], []])
    assert parser.parse_prices(content, "price.xlsx")["source_rows"] == 1


@pytest.mark.parametrize("header", [
    ["SKU", "产品代码", "单价"], ["SKU", "单价", "采购价"], ["SKU", "单价", "中间码", "中间码"],
])
def test_duplicate_or_ambiguous_headers_reject_file(header):
    with pytest.raises(ValueError, match="表头.*重复"):
        parser.parse_prices(workbook_bytes([header, ["A-123-X", 10]]), "price.xlsx")


def test_nonempty_sheet_without_valid_header_is_not_silently_ignored():
    content = workbook_bytes(standard_rows(["A-123-X", 1]), extra_sheets={"错误表": [["参考SKU", "建议单价"], ["B-456-X", 2]]})
    with pytest.raises(ValueError, match="错误表.*表头"):
        parser.parse_prices(content, "price.xlsx")


def test_empty_sheets_are_harmless_but_empty_price_file_is_rejected():
    content = workbook_bytes(standard_rows(["A-123-X", 1]), extra_sheets={"空表": []})
    assert parser.parse_prices(content, "price.xlsx")["source_rows"] == 1
    with pytest.raises(ValueError, match="没有有效价格数据"):
        parser.parse_prices(workbook_bytes(standard_rows()), "price.xlsx")


def test_header_must_be_within_first_twenty_physical_rows():
    valid = workbook_bytes([*([[]] * 19), *standard_rows(["A-123-X", 1])])
    assert parser.parse_prices(valid, "price.xlsx")["source_rows"] == 1
    invalid = workbook_bytes([*([[]] * 20), *standard_rows(["A-123-X", 1])])
    with pytest.raises(ValueError, match="前20行"):
        parser.parse_prices(invalid, "price.xlsx")


@pytest.mark.parametrize("column", [65, 66, 100])
def test_data_beyond_column_limit_is_rejected_even_when_column_65_is_empty(column):
    header = ["SKU", "单价"] + [None] * (column - 3) + ["过多列"]
    with pytest.raises(ValueError, match="最多支持64列"):
        parser.parse_prices(workbook_bytes([header, ["A-123-X", 1]]), "price.xlsx")


def test_exact_column_limit_is_allowed():
    header = ["SKU", "单价"] + [None] * 61 + ["备注"]
    assert parser.parse_prices(workbook_bytes([header, ["A-123-X", 1]]), "price.xlsx")["source_rows"] == 1


@pytest.mark.parametrize("filename,content", [
    ("prices.xls", b"anything"), ("prices.csv", b"anything"), ("prices.xlsx", b""), ("prices.xlsx", b"not a zip"),
])
def test_invalid_extension_or_container_rejected(filename, content):
    with pytest.raises(ValueError):
        parser.parse_prices(content, filename)


def test_compressed_limit_checked_before_workbook_load(monkeypatch):
    content = workbook_bytes(standard_rows(["A-123-X", 1]))
    load = MagicMock()
    monkeypatch.setattr(parser, "load_workbook", load)
    monkeypatch.setattr(parser, "MAX_FILE_BYTES", len(content) - 1)
    with pytest.raises(ValueError, match="10MB"):
        parser.parse_prices(content, "price.xlsx")
    load.assert_not_called()


def test_expanded_archive_size_limit(monkeypatch):
    content = workbook_bytes(standard_rows(["A-123-X", 1]))
    monkeypatch.setattr(parser, "MAX_EXPANDED_BYTES", 1)
    with pytest.raises(ValueError, match="解压后过大"):
        parser.parse_prices(content, "price.xlsx")


def test_archive_entry_limit():
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for index in range(2001):
            archive.writestr(f"empty-{index}", b"")
    with pytest.raises(ValueError, match="解压后过大"):
        parser.parse_prices(output.getvalue(), "price.xlsx")


def test_total_rows_limit_includes_duplicates_and_all_sheets(monkeypatch):
    content = workbook_bytes(standard_rows(["A-123-X", 1]), extra_sheets={"补充": standard_rows(["A-123-X", 1], ["B-456-X", 2])})
    monkeypatch.setattr(parser, "MAX_ROWS", 2)
    with pytest.raises(ValueError, match="最多导入2行"):
        parser.parse_prices(content, "price.xlsx")


def test_physical_rows_limit_applies_to_sparse_sheets(monkeypatch):
    content = workbook_bytes([*standard_rows(["A-123-X", 1]), *([[]] * 20), ["B-456-X", 2]])
    monkeypatch.setattr(parser, "MAX_ROWS", 2)
    with pytest.raises(ValueError, match="单张价格表不能超过2行"):
        parser.parse_prices(content, "price.xlsx")


def test_damaged_sheet_xml_is_reported_as_value_error():
    content = workbook_bytes(standard_rows(["A-123-X", 1]))
    output = BytesIO()
    with ZipFile(BytesIO(content)) as original, ZipFile(output, "w", ZIP_DEFLATED) as damaged:
        for item in original.infolist():
            data = b"<broken>" if item.filename == "xl/worksheets/sheet1.xml" else original.read(item.filename)
            damaged.writestr(item.filename, data)
    with pytest.raises(ValueError, match="Excel工作簿|XML内容损坏"):
        parser.parse_prices(output.getvalue(), "price.xlsx")


def test_filename_is_never_used_as_a_write_target(monkeypatch):
    def reject_write(*args, **kwargs):
        raise AssertionError("parser must not write files")
    monkeypatch.setattr(Path, "write_bytes", reject_write)
    monkeypatch.setattr(Path, "write_text", reject_write)
    content = workbook_bytes(standard_rows(["A-123-X", 1]))
    assert parser.parse_prices(content, "../../outside/price.xlsx")["rows"][0]["sku"] == "A-123-X"


def test_workbook_is_closed_after_row_validation_failure(monkeypatch):
    content = workbook_bytes(standard_rows(["A-123-X", -1]))
    original_load = parser.load_workbook
    loaded = []

    def tracked_load(*args, **kwargs):
        assert kwargs == {"read_only": True, "data_only": True, "keep_links": False}
        workbook = original_load(*args, **kwargs)
        workbook.close = MagicMock(wraps=workbook.close)
        loaded.append(workbook)
        return workbook

    monkeypatch.setattr(parser, "load_workbook", tracked_load)
    with pytest.raises(ValueError):
        parser.parse_prices(content, "price.xlsx")
    loaded[0].close.assert_called_once_with()
