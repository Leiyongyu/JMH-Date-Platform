from decimal import Decimal
from datetime import date, datetime
from contextlib import contextmanager
from copy import deepcopy
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.services import weekly_inventory_sync_service as sync
from backend.services import weekly_inventory_export_service as export
from backend.repositories import weekly_inventory_repository as repo
from backend.api.v1 import weekly_inventory as api


INVENTORY_HEADERS = [
    '仓库名称', 'SKU', '本地产品id', '实际库存总量', '可用量',
    '次品量', '待检待上架量', '锁定量',
]
BIN_HEADERS = ['仓位名称', '总量(仓位)', '锁定量(仓位)', '未锁定量(仓位)']
PRODUCT_HEADERS = ['产品名称'] + [
    f'采购-{kind}规格-{axis}(CM)'
    for kind in ('包装', '外箱') for axis in ('长', '宽', '高')
] + ['采购-产品净重(G)', '采购-产品毛重(KG)', '采购-外箱实重(KG)']


def sample(wid=1):
    inventory = [dict(wid=1, product_id=2, seller_id="a", sku="=SKU", product_total="10", purchase_price="1.234567",
                      product_valid_num="9", stock_age_list=[dict(name="0-29天", qty=10)],
                      third_inventory={"third_inventory_data": [dict(name="可用量", local=9, third=8, diff=1)]}),
                 dict(wid=1, product_id=2, seller_id="b", sku="=SKU", product_total="5", purchase_price="99",
                      product_valid_num="20", stock_age_list=[dict(name="0-29天", qty=5)])]
    bins = [dict(wid=1, whb_id=3, product_id=2, whb_name='A-01', total=11, lockNum=1, validNum=3, whb_type=9),
            dict(wid=1, whb_id=4, product_id=2, whb_name='A-01', total=13, lockNum=2, validNum=4)]
    for row in inventory + bins:
        row["wid"] = wid
    return sync.normalize(inventory, bins, [dict(id=2, product_name="产品", model="X"*222)], date(2026,9,9), "batch", datetime.now())


def test_representative_whole_row_and_age_not_double_counted():
    headers, rows = export.build_rows(sample(), {1:"仓库"})
    row = dict(zip(headers, rows[0]))
    assert len(rows) == 1
    assert row["实际库存总量"] == 10
    assert row["可用量"] == 9  # not MAX(20) from a different seller
    assert row["0-29天"] == 10  # not sum(10+5)
    assert row["总量(仓位)"] == 24  # source total, not lock_num + valid_num
    assert row["锁定量(仓位)"] == 3
    assert row["未锁定量(仓位)"] == 7
    assert row["库存金额"] == Decimal("12.345670")
    assert not any(label.startswith("第三方-") for label in headers)
    assert len(headers) == 25  # 24 static + 1 dynamic
    assert headers == INVENTORY_HEADERS + ['0-29天', '采购单价', '库存金额'] + BIN_HEADERS + PRODUCT_HEADERS
    assert len(rows[0]) == len(headers)


@pytest.mark.parametrize('grams,kilograms', [
    (None, None), (Decimal('0'), Decimal('0')),
    (Decimal('1500'), Decimal('1.5')), (Decimal('1000'), Decimal('1')),
    (Decimal('25.5'), Decimal('0.0255')),
])
def test_only_exported_gross_weight_converts_to_kg(grams, kilograms):
    groups = sample()
    product = groups['products'][0]
    product.update(cg_product_gross_weight=grams, cg_product_net_weight=Decimal('900'),
                   cg_box_weight=Decimal('8.5'))
    headers, rows = export.build_rows(groups, {1: '仓库'})
    row = dict(zip(headers, rows[0]))
    assert '采购-产品毛重(G)' not in headers
    assert row['采购-产品毛重(KG)'] == kilograms
    assert row['采购-产品净重(G)'] == Decimal('900')
    assert row['采购-外箱实重(KG)'] == Decimal('8.5')
    assert product['cg_product_gross_weight'] == grams


def test_missing_is_not_zero():
    groups = sample()
    groups["inventory"] = groups["inventory"][:1]
    groups["inventory"][0]["product_bad_num"] = Decimal(0)
    for bin_row in groups['bins']:
        bin_row.update(total=None, lock_num=None, valid_num=None)
    groups["age"] = []
    headers, rows = export.build_rows(groups,{1:"仓库"})
    row=dict(zip(headers,rows[0]))
    assert len(rows) == 1
    assert row["次品量"] == 0
    assert row['仓位名称'] == 'A-01'
    assert all(row[name] is None for name in BIN_HEADERS[1:])
    assert row["采购-产品净重(G)"] is None


@pytest.mark.parametrize("names,expected", [
    ([None, ""], []),
    (["   ", None], []),
    (["\t\r\n", "\u3000\u00a0"], []),
    (["A-01", " A-01 "], [("A-01", 24, 3, 7)]),
    ([" B-02 ", "A-01"], [("A-01", 13, 2, 4), ("B-02", 11, 1, 3)]),
    (["A-01", ""], [("A-01", 11, 1, 3)]),
    ([None, " B-02 "], [("B-02", 13, 2, 4)]),
    (["\u3000A-01\u00a0", "\t"], [("A-01", 11, 1, 3)]),
])
def test_bin_names_expand_unique_names_and_repeat_other_fields(names, expected):
    groups = sample(18677)
    for bin_row, name in zip(groups['bins'], names):
        bin_row['whb_name'] = name
    # Neither another warehouse nor another product can contribute bin names.
    groups['bins'].extend([
        {**groups['bins'][0], 'wid': 19561, 'whb_name': 'OTHER-WAREHOUSE'},
        {**groups['bins'][0], 'product_id': 999, 'whb_name': 'OTHER-PRODUCT'},
    ])
    before = deepcopy(groups)
    named_baseline = deepcopy(groups)
    for bin_row in named_baseline['bins']:
        bin_row['whb_name'] = 'A-01'
    baseline_headers, baseline_rows = export.build_rows(named_baseline, {18677: 'EU'})
    assert len(baseline_rows) == 1
    baseline = dict(zip(baseline_headers, baseline_rows[0]))
    for name in BIN_HEADERS:
        baseline.pop(name)
    headers, rows = export.build_rows(groups, {18677: 'EU'})
    assert headers[headers.index('库存金额') + 1] == '仓位名称'
    assert len(rows) == len(expected)
    named_rows = [dict(zip(headers, values)) for values in rows]
    assert [tuple(row.pop(name) for name in BIN_HEADERS) for row in named_rows] == expected
    assert all(row == baseline for row in named_rows)
    assert baseline['实际库存总量'] == 10
    assert baseline['库存金额'] == Decimal('12.345670')
    assert groups == before


@pytest.mark.parametrize('kind', ['none', 'empty', 'spaces', 'tabs', 'unicode', 'absent', 'no_bins'])
def test_no_named_bins_omits_sku_and_rejects_empty_export_without_file(tmp_path, kind):
    groups = sample(18677)
    for bin_row in groups['bins']:
        if kind == 'absent':
            bin_row.pop('whb_name')
        else:
            bin_row['whb_name'] = {
                'none': None, 'empty': '', 'spaces': '   ', 'tabs': '\t\r\n',
                'unicode': '\u3000\u00a0\u2003', 'no_bins': None,
            }[kind]
    if kind == 'no_bins':
        groups['bins'] = []
    # A named bin for another warehouse/product cannot rescue this SKU.
    groups['bins'].extend([
        dict(wid=19561, product_id=2, whb_name='OTHER-WAREHOUSE', total=999),
        dict(wid=18677, product_id=999, whb_name='OTHER-PRODUCT', total=999),
    ])
    before = deepcopy(groups)
    headers, rows = export.build_rows(groups, {})
    assert rows == []
    assert all(name in headers for name in BIN_HEADERS)
    path = tmp_path / 'not-created' / 'empty.xlsx'
    with pytest.raises(ValueError, match='没有可导出的周报数据（仅导出仓位名称非空的明细）'):
        export.write_export(groups, {}, path)
    assert not path.exists()
    assert not path.parent.exists()
    assert not list(tmp_path.rglob('.weekly-*'))
    assert groups == before


def test_xlsx_filters_unnamed_details_and_whole_skus_without_changing_source(tmp_path):
    groups = sample(18677)
    groups['bins'][1]['whb_name'] = '\u3000 A-01 \t'
    inventory = groups['inventory'][0]
    groups['inventory'].extend([
        {**inventory, 'product_id': 3, 'sku': 'NO-BINS'},
        {**inventory, 'product_id': 4, 'sku': 'EMPTY-NAME'},
        {**inventory, 'wid': 19561},
    ])
    groups['products'].extend([
        {**groups['products'][0], 'product_id': product_id}
        for product_id in (3, 4)
    ])
    bin_row = groups['bins'][0]
    groups['bins'].extend([
        {**bin_row, 'whb_id': 5, 'whb_name': None, 'total': 999, 'lock_num': 999, 'valid_num': 999},
        {**bin_row, 'whb_id': 6, 'whb_name': '', 'total': None, 'lock_num': None, 'valid_num': None},
        {**bin_row, 'product_id': 4, 'whb_name': '\u3000\t'},
        {**bin_row, 'wid': 19561, 'whb_name': ' '},
    ])
    before = deepcopy(groups)
    path = tmp_path / 'named-only.xlsx'
    metrics = export.write_export(groups, {18677: 'EU'}, path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = workbook.active.iter_rows(values_only=True)
        headers = next(values)
        rows = [dict(zip(headers, row)) for row in values]
        assert metrics['row_count'] == len(rows) == 1
        row = rows[0]
        assert tuple(row[name] for name in BIN_HEADERS) == ('A-01', 24, 3, 7)
        assert (row['仓库名称'], row['SKU'], row['本地产品id']) == ('EU', '=SKU', 2)
        assert (row['实际库存总量'], row['可用量'], row['0-29天']) == (10, 9, 10)
        assert (row['采购单价'], row['库存金额']) == (1.234567, 12.34567)
        assert row['产品名称'] == '产品'
    finally:
        workbook.close()
    assert groups == before


@pytest.mark.parametrize('missing_field', ['total', 'lock_num', 'valid_num'])
@pytest.mark.parametrize('missing_kind', ['none', 'absent'])
def test_bin_missing_quantities_are_independent_and_do_not_affect_other_names(missing_field, missing_kind):
    groups = sample()
    for row in groups['bins']:
        row['whb_name'] = 'A-01'
    groups['bins'].append({**groups['bins'][0], 'whb_id': 5, 'whb_name': 'B-02'})
    if missing_kind == 'none':
        groups['bins'][1][missing_field] = None
    else:
        groups['bins'][1].pop(missing_field)
    before = deepcopy(groups)
    headers, rows = export.build_rows(groups, {1: '仓库'})
    by_name = {row['仓位名称']: row for row in (dict(zip(headers, values)) for values in rows)}
    for field, header, total, other in [
        ('total', '总量(仓位)', 24, 11),
        ('lock_num', '锁定量(仓位)', 3, 1),
        ('valid_num', '未锁定量(仓位)', 7, 3),
    ]:
        assert by_name['A-01'][header] == (None if field == missing_field else total)
        assert by_name['B-02'][header] == other
    assert groups == before


def test_bin_zero_is_present_and_same_name_merges_ids_and_sellers():
    groups = sample()
    for index, row in enumerate(groups['bins']):
        row.update(whb_name=' ZERO ' if index else 'ZERO', store_id=f'store-{index}',
                   total=Decimal(0), lock_num=Decimal(0), valid_num=Decimal(0))
    before = deepcopy(groups)
    headers, rows = export.build_rows(groups, {1: '仓库'})
    assert len(rows) == 1
    row = dict(zip(headers, rows[0]))
    assert tuple(row[name] for name in BIN_HEADERS) == ('ZERO', 0, 0, 0)
    assert groups == before


def test_xlsx_expands_bin_names_and_counts_exported_rows(tmp_path):
    groups = sample(18677)
    groups['bins'][0]['whb_name'] = ' B-02 '
    groups['bins'][1]['whb_name'] = 'A-01'
    # Separate source records with the same bin name yield only one display row.
    groups['bins'].append({**groups['bins'][1], 'whb_id': 5, 'store_id': 'other-shop', 'whb_name': ' A-01 '})
    before = deepcopy(groups)
    path = tmp_path / 'bin-rows.xlsx'
    metrics = export.write_export(groups, {18677: 'EU'}, path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = workbook.active.iter_rows(values_only=True)
        headers = next(values)
        rows = [dict(zip(headers, row)) for row in values]
        assert metrics['row_count'] == len(rows) == 2
        assert metrics['column_count'] == len(headers) == 25
        assert [tuple(row.pop(name) for name in BIN_HEADERS) for row in rows] == [
            ('A-01', 26, 4, 8), ('B-02', 11, 1, 3),
        ]
        assert rows[0] == rows[1]
        assert rows[0]['实际库存总量'] == 10
        assert rows[0]['0-29天'] == 10
        assert rows[0]['库存金额'] == 12.34567
    finally:
        workbook.close()
    assert groups == before


def test_exported_bin_name_is_text_not_excel_formula(tmp_path):
    groups = sample(18677)
    for bin_row in groups['bins']:
        bin_row['whb_name'] = '=BIN'
    path = tmp_path / 'bin-name.xlsx'
    export.write_export(groups, {18677: 'EU'}, path)
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        column = headers.index('仓位名称') + 1
        cell = sheet.cell(2, column)
        assert cell.value == '=BIN' and cell.data_type == 's'
        assert sheet.cell(1, column - 1).value == '库存金额'
    finally:
        workbook.close()


def test_xlsx_is_safe_and_archive_never_replaced(tmp_path):
    path=tmp_path/"中文.xlsx"
    metrics=export.write_export(sample(18677),{18677:"仓库"},path)
    original=path.read_bytes()
    wb=load_workbook(path)
    assert wb.active["B2"].value == "=SKU"
    assert wb.active["B2"].data_type == "s"
    assert wb.active.max_column == metrics["column_count"]
    wb.close()
    with pytest.raises(FileExistsError):
        export.write_export(sample(18677),{18677:"仓库"},path)
    assert path.read_bytes()==original
    assert not list(tmp_path.glob(".weekly-*"))


def test_xlsx_numeric_format_and_removed_columns(tmp_path):
    groups = sample(18677)
    groups["inventory"][0]["product_bad_num"] = Decimal("0.000000")
    groups["products"][0].update(cg_product_length=Decimal("75.000000"),
                                cg_product_width=Decimal("25.500000"),
                                cg_product_height=Decimal("-2.000000"),
                                cg_package_length=Decimal("75.000000"),
                                cg_package_width=Decimal("25.500000"),
                                cg_package_height=Decimal("-2.000000"),
                                cg_box_length=Decimal("100.000000"),
                                cg_box_width=Decimal("0.000000"),
                                cg_box_height=Decimal("5.250000"),
                                cg_product_net_weight=Decimal("900.000000"),
                                cg_product_gross_weight=Decimal("1500.000000"),
                                cg_box_weight=Decimal("8.500000"))
    before = deepcopy(groups)
    path = tmp_path / "format.xlsx"
    export.write_export(groups, {18677: "仓库"}, path)
    with_source = groups["inventory"][0]["third_inventory"]
    assert with_source["third_inventory_data"][0]["third"] == 8
    wb = load_workbook(path)
    try:
        sheet = wb.active
        cells = dict(zip([c.value for c in sheet[1]], sheet[2]))
        assert not any(name.startswith("第三方-") for name in cells)
        assert not any(name.startswith('采购-产品规格-') for name in cells)
        assert sheet.max_column == 25
        assert sheet.freeze_panes == "D2"
        for name, expected, fmt in [
            ("实际库存总量", 10, "0"), ("次品量", 0, "0"),
            ("总量(仓位)", 24, "0"),
            ("采购-包装规格-长(CM)", 75, "0"),
            ("采购-包装规格-宽(CM)", 25.5, "0.0#####"),
            ("采购-包装规格-高(CM)", -2, "0"),
            ("采购-外箱规格-长(CM)", 100, "0"),
            ("采购-外箱规格-宽(CM)", 0, "0"),
            ("采购-外箱规格-高(CM)", 5.25, "0.0#####"),
            ("采购-产品净重(G)", 900, "0"),
            ("采购-产品毛重(KG)", 1.5, "0.0#####"),
            ("采购-外箱实重(KG)", 8.5, "0.0#####"),
            ("采购单价", 1.234567, "0.0#####"),
            ("库存金额", 12.34567, "0.0#####"),
        ]:
            assert cells[name].value == expected
            assert cells[name].data_type == "n"
            assert cells[name].number_format == fmt
    finally:
        wb.close()
    assert groups == before


def test_xlsx_source_groups_have_consistent_header_and_body_styles(tmp_path):
    from openpyxl.utils import get_column_letter

    groups = sample(18677)
    groups['age'].append({**groups['age'][0], 'bucket_name': '30-89天', 'bucket_index': 1, 'qty': None})
    groups['bins'][0]['whb_name'] = 'A'
    groups['bins'][1]['whb_name'] = 'B'
    path = tmp_path / 'source-groups.xlsx'
    metrics = export.write_export(groups, {18677: '仓库'}, path)
    source_headers = {
        'inventory': INVENTORY_HEADERS + ['0-29天', '30-89天', '采购单价', '库存金额'],
        'bin': BIN_HEADERS,
        'product': PRODUCT_HEADERS,
    }
    expected_headers = sum(source_headers.values(), [])
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        assert [cell.value for cell in sheet[1]] == expected_headers
        assert sheet.max_column == metrics['column_count'] == 26
        assert sheet.max_row == 3
        assert sheet.freeze_panes == 'D2'
        assert sheet.auto_filter.ref == f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        for layer in ('header', 'body'):
            assert len({export.SOURCE_STYLES[group][layer] for group in source_headers}) == 3
        for group, labels in source_headers.items():
            palette = export.SOURCE_STYLES[group]
            header_rgb = palette['header'].lstrip('#')[-6:].upper()
            body_rgb = palette['body'].lstrip('#')[-6:].upper()
            assert len(header_rgb) == len(body_rgb) == 6
            assert sum(bytes.fromhex(header_rgb)) < sum(bytes.fromhex(body_rgb))
            for label in labels:
                column = expected_headers.index(label) + 1
                header_cell = sheet.cell(1, column)
                assert header_cell.fill.fill_type == 'solid'
                assert header_cell.fill.fgColor.rgb[-6:].upper() == header_rgb
                assert header_cell.alignment.wrap_text
                assert header_cell.font.bold
                assert header_cell.font.name
                assert sheet.column_dimensions[get_column_letter(column)].width >= 10
                for row in (2, 3):
                    cell = sheet.cell(row, column)
                    assert cell.fill.fill_type == 'solid'
                    assert cell.fill.fgColor.rgb[-6:].upper() == body_rgb
    finally:
        workbook.close()


@pytest.mark.parametrize("value", ["NaN", "Infinity", "nonsense", True])
def test_invalid_numbers(value):
    with pytest.raises(ValueError): sync.number(value)


def test_duplicate_primary_key_rejected():
    with pytest.raises(ValueError):
        sync.unique([dict(id=1),dict(id=1)], ("id",))


def test_pagination_complete():
    client=MagicMock()
    client.post_signed_query_auth.side_effect=[dict(code=0,total=3,data=[{},{}]),dict(code=0,total=3,data=[{}])]
    assert len(sync.fetch_pages(client,"inventoryDetails",2)) == 3
    assert client.post_signed_query_auth.call_args.args[1]==dict(offset=2,length=2)


@pytest.mark.parametrize("responses", [
    [dict(code=0,total=3,data=[{}])],
    [dict(code=0,total=3,data=[{},{}]),dict(code=0,total=4,data=[{},{}])],
    [dict(code=0,data=[])], [dict(code=1,total=0,data=[])],
])
def test_incomplete_pages_rejected(responses):
    client=MagicMock(); client.post_signed_query_auth.side_effect=responses
    with pytest.raises(ValueError): sync.fetch_pages(client,"inventoryDetails",2)


def test_snapshot_transaction_rollback(monkeypatch):
    connection=MagicMock(); cursor=connection.cursor.return_value.__enter__.return_value
    cursor.fetchone.return_value=dict(status='RUNNING', snapshot_date=date(2026,9,9))
    cursor.executemany.side_effect=RuntimeError("insert failure")
    @contextmanager
    def connect(): yield connection
    monkeypatch.setattr(repo,"db_connection",connect)
    with pytest.raises(RuntimeError):
        repo.replace_snapshot_and_finish_export(sample(), file_name='new.xlsx',
                                                row_count=1, column_count=25, file_size=123)
    connection.rollback.assert_called_once()
    connection.commit.assert_not_called()


def test_download_traversal_and_missing(monkeypatch,tmp_path):
    monkeypatch.setattr(export,"export_root",lambda:tmp_path)
    with pytest.raises(ValueError):
        export.download_path(dict(status="SUCCESS",file_path=str(tmp_path.parent/"secret.xlsx"),file_name="secret.xlsx"))
    with pytest.raises(FileNotFoundError):
        export.download_path(dict(status="SUCCESS",file_path=str(tmp_path/"missing.xlsx"),file_name="missing.xlsx"))


def test_internal_auth_before_repository(monkeypatch):
    app=FastAPI();app.include_router(api.router)
    lookup=MagicMock(); monkeypatch.setattr(repo,"list_files",lookup)
    client=TestClient(app)
    response=client.get("/api/v1/weekly-inventory/files")
    assert response.status_code in (401,403)
    lookup.assert_not_called()


def test_no_historical_snapshot():
    from backend.services.scheduler_service import run_scheduler_task
    with pytest.raises(ValueError, match="历史"):
        run_scheduler_task(sync.TASK_CODE,stat_month="2026-08")


def test_batch_prunes_only_other_batches_after_inserts(monkeypatch):
    connection=MagicMock();cursor=connection.cursor.return_value.__enter__.return_value
    cursor.fetchone.return_value=dict(status='RUNNING', snapshot_date=date(2026,9,9))
    cursor.rowcount=1
    @contextmanager
    def connect():yield connection
    monkeypatch.setattr(repo,"db_connection",connect)
    repo.replace_snapshot_and_finish_export(sample(), file_name='new.xlsx',
                                            row_count=1, column_count=25, file_size=123)
    assert all(call.args[0].startswith("INSERT INTO") for call in cursor.executemany.call_args_list)
    deletes=[call for call in cursor.execute.call_args_list if call.args[0].startswith('DELETE')]
    assert len(deletes)==len(repo.TABLES)==4
    assert all('WHERE sync_batch_id<>%s' in call.args[0] and call.args[1]==('batch',) for call in deletes)
    assert all('ops_weekly_export_file' not in call.args[0] for call in deletes)
    connection.commit.assert_called_once()


def test_schema_matches_every_normalized_column():
    import re
    from pathlib import Path
    schema=(Path(__file__).parents[1]/'backend/schema.sql').read_text(encoding='utf-8')
    for group, table in repo.TABLES.items():
        definition=re.search(r'CREATE TABLE IF NOT EXISTS `'+table+r'` \((.*?)\) ENGINE',schema,re.S).group(1)
        columns=set(re.findall(r'^  `([^`]+)`',definition,re.M))
        assert set(sample()[group][0]) == columns-{'id','create_time','bin_identity_key'}
        assert 'snapshot_date`,`sync_batch_id`' in definition


@pytest.mark.parametrize('missing_product,registration_failure,capture_failure', [
    (False,False,False),(True,False,False),(False,True,False),(False,False,True),
])
def test_whole_chain_without_real_io(monkeypatch,tmp_path,missing_product,registration_failure,capture_failure):
    import sys
    from types import ModuleType

    monkeypatch.setattr(repo,'require_bin_identity_schema',MagicMock())
    client=MagicMock()
    client.post_signed_query_auth.side_effect=[
        dict(code=0,total=2,data=[
            dict(wid=18677,product_id=2,sku='SKU',product_total=3,purchase_price='1.25',extra={'keep':'raw'}),
            dict(wid=18678,product_id=2,sku='SKU',product_total=99,purchase_price='1.25'),
        ]),
        dict(code=0,total=1,data=[
            dict(wid=18677,whb_id=3,product_id=2,whb_name='A-01',total=3,lockNum=0,validNum=3),
        ]),
        dict(code=0,data=[] if missing_product else [dict(id=2,sku='SKU',product_name='产品')]),
    ]
    monkeypatch.setattr(sync,'LingXingClient',lambda **kwargs:client)
    monkeypatch.setattr(export,'export_root',lambda:tmp_path)
    begin=MagicMock();monkeypatch.setattr(repo,'begin_export',begin)
    stored={}
    def replace(groups, **metrics):
        stored.update(groups)
        assert (tmp_path/metrics['file_name']).is_file()
        if registration_failure:
            raise RuntimeError('registration failure')
        return dict(ods_rows=sum(map(len,groups.values())), deleted_rows=4)
    replace_mock=MagicMock(side_effect=replace)
    monkeypatch.setattr(repo,'replace_snapshot_and_finish_export',replace_mock)
    snapshot=MagicMock(side_effect=AssertionError('new extraction must export normalized rows directly'))
    monkeypatch.setattr(repo,'snapshot',snapshot)
    monkeypatch.setattr(repo,'warehouse_names',lambda:{18677:'仓库'})
    finish=MagicMock()
    monkeypatch.setattr(repo,'finish_export',finish)
    # Only this full-chain test replaces the history boundary. No global test
    # fixture may hide whether production history capture is actually invoked.
    pivot_module=ModuleType('backend.services.ebay_inventory_pivot_service')
    def capture(*,expected_inventory_batch,trigger_type):
        assert stored['inventory'][0]['sync_batch_id']==expected_inventory_batch
        assert len(list(tmp_path.glob('*.xlsx')))==1
        if capture_failure:
            raise RuntimeError('injected historical storage failure')
        return dict(snapshot_id=12,stat_date='2026-09-16',group_count=1,item_count=1)
    capture_mock=MagicMock(side_effect=capture)
    pivot_module.capture_snapshot=capture_mock
    monkeypatch.setitem(sys.modules,pivot_module.__name__,pivot_module)
    if missing_product:
        with pytest.raises(ValueError,match='产品详情'):sync.sync_weekly_inventory()
        replace_mock.assert_not_called()
        capture_mock.assert_not_called()
        assert not list(tmp_path.glob('*.xlsx'))
        assert finish.call_args.kwargs['error']
    elif registration_failure:
        with pytest.raises(RuntimeError,match='registration'):sync.sync_weekly_inventory()
        capture_mock.assert_not_called()
        assert len(list(tmp_path.glob('*.xlsx'))) == 1  # preserve published evidence
        assert stored['inventory']
        assert finish.call_args.kwargs['error']
    elif capture_failure:
        with pytest.raises(ValueError,match='库存和Excel已成功发布，但历史透视保存失败'):
            sync.sync_weekly_inventory('manual')
        replace_mock.assert_called_once()
        capture_mock.assert_called_once_with(
            expected_inventory_batch=stored['inventory'][0]['sync_batch_id'],trigger_type='manual')
        assert len(list(tmp_path.glob('*.xlsx')))==1
        finish.assert_not_called()  # The committed SUCCESS record is untouched.
    else:
        result=sync.sync_weekly_inventory('manual')
        assert result['row_count']==1 and result['column_count']==24
        assert result['deduplicated_inventory_rows'] == 0
        assert {row['wid'] for row in stored['inventory']} == {18677, 18678}
        assert len(list(tmp_path.glob('*.xlsx')))==1
        assert stored['inventory'][0]['raw_json']['extra']=={'keep':'raw'}
        assert stored['inventory'][0]['sync_batch_id']==result['sync_batch_id']
        assert begin.call_args.args[-1]=='manual'
        finish.assert_not_called()
        assert result['deleted_rows']==4
        assert result['pivot_snapshot']['snapshot_id']==12
        capture_mock.assert_called_once_with(expected_inventory_batch=result['sync_batch_id'],trigger_type='manual')
    snapshot.assert_not_called()


def test_scheduler_weekly_is_globally_locked_and_logged(monkeypatch):
    from backend.services import scheduler_service as scheduler
    connection=MagicMock(); logs=[];locks=[]
    @contextmanager
    def connect():yield connection
    @contextmanager
    def locked(name):
        locks.append(name)
        yield True
    monkeypatch.setattr(scheduler.repo,'performance_connection',connect)
    monkeypatch.setattr(scheduler.repo,'insert_scheduler_run',lambda conn,row:logs.append(row))
    monkeypatch.setattr(scheduler.repo,'named_lock',locked)
    runner=MagicMock(return_value=dict(sync_batch_id='b',extract_rows=1,ods_rows=1))
    monkeypatch.setattr(scheduler,'sync_weekly_inventory',runner)
    response=scheduler.run_scheduler_task(sync.TASK_CODE,trigger_type='JOB')
    assert response['status']=='completed'
    assert locks==['inventory:weekly-export']
    assert [row['status'] for row in logs]==['running','completed']
    runner.assert_called_once_with('JOB')
